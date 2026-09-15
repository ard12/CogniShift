"""
Unit and integration tests for DOCX Hybrid RAG in CogniShift:
- Native hierarchical structure extraction (headings, tables, paragraphs)
- Verification that native text chunks omit fake page numbers
- Word COM local rendering and graceful degradation when unavailable
- Format-aware citations: [Doc.docx | Section: X | TEXT] / [Doc.docx | Page Y | VISUAL]
- Evidence fusion and ranking across DOCX text and visual pages
"""
import pytest
from pathlib import Path
from docx import Document
from unittest.mock import patch

from cognishift.core.document_processing.office_renderer import (
    extract_docx_structured_content,
    render_docx_pages
)
from cognishift.core.document_processing.provenance import format_grounded_citation
from cognishift.core.retrieval.evidence_fusion import EvidenceFusion
from cognishift.core.visual_rag.schemas import VisualSearchResult


@pytest.fixture
def sample_docx_file(tmp_path: Path) -> Path:
    """Creates a sample DOCX document with title, headings, paragraphs, and tables."""
    doc = Document()
    doc.add_heading("Refinery Pre-Commissioning Procedure", level=0)
    
    # Section 1
    doc.add_heading("1. Emergency Shutdown Protocols", level=1)
    doc.add_paragraph("In the event of an uncontrolled pressure excursion exceeding 45.0 barg in vessel V-101, immediate trip is mandatory.")
    doc.add_paragraph("Operators must isolate suction valve XV-101 within 30 seconds.")
    
    # Section 2 with Table
    doc.add_heading("2. Transmitter Calibration Thresholds", level=1)
    doc.add_paragraph("The following setpoints must be verified against ASME Section VIII Division 1:")
    
    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Tag"
    hdr_cells[1].text = "Normal Range"
    hdr_cells[2].text = "Trip Limit"
    
    r1 = table.rows[1].cells
    r1[0].text = "PT-101"
    r1[1].text = "35.0 - 40.0 barg"
    r1[2].text = "45.0 barg"
    
    r2 = table.rows[2].cells
    r2[0].text = "TT-201"
    r2[1].text = "180 - 210 degC"
    r2[2].text = "240 degC"

    file_path = tmp_path / "sop_commissioning.docx"
    doc.save(str(file_path))
    return file_path


def test_docx_native_extraction(sample_docx_file: Path):
    """Verify that DOCX native extraction extracts headings, text, and tables without fake page numbers."""
    chunks = extract_docx_structured_content(sample_docx_file)
    assert len(chunks) >= 2

    # Check first section
    sec1 = chunks[0]
    assert sec1["section_heading"] == "1. Emergency Shutdown Protocols"
    assert "excursion exceeding 45.0 barg" in sec1["text"]
    assert "V-101" in sec1["text"]
    # Critical Invariant: native DOCX chunks do not have fake page numbers!
    assert sec1.get("page") is None
    assert sec1["chunk_index"] == 1

    # Check second section containing table
    sec2 = chunks[1]
    assert sec2["section_heading"] == "2. Transmitter Calibration Thresholds"
    assert "PT-101" in sec2["text"]
    assert "45.0 barg" in sec2["text"]
    assert "TT-201" in sec2["text"]
    assert sec2.get("page") is None


def test_docx_renderer_graceful_degradation(tmp_path: Path):
    """Verify that if Word COM automation fails or is absent, renderer degrades truthfully."""
    fake_docx = tmp_path / "test.docx"
    fake_docx.write_bytes(b"dummy")

    with patch("win32com.client.Dispatch", side_effect=Exception("Word COM not installed")):
        pages, prov = render_docx_pages(fake_docx)
        assert pages == []
        assert prov["render_status"] == "RENDERER_UNAVAILABLE"


def test_docx_citation_formatting():
    """Verify authoritative citation formatting for DOCX documents."""
    # 1. Native text chunk with section heading
    meta_text = {
        "filename": "commissioning_sop.docx",
        "document_type": "docx",
        "section_heading": "Emergency Isolation Procedures",
        "channel": "text",
        "extraction_method": "document"
    }
    citation_text = format_grounded_citation(meta_text)
    assert citation_text == "[commissioning_sop.docx | Section: Emergency Isolation Procedures | TEXT]"

    # 2. Rendered visual page
    meta_vis = {
        "filename": "commissioning_sop.docx",
        "document_type": "docx",
        "page": 3,
        "channel": "visual",
        "extraction_method": "visual"
    }
    citation_vis = format_grounded_citation(meta_vis)
    assert citation_vis == "[commissioning_sop.docx | Page 3 | VISUAL]"


def test_docx_evidence_fusion_and_ranking():
    """Verify EvidenceFusion produces format-aware DOCX citations and metadata."""
    fusion = EvidenceFusion(rrf_k=60)

    text_items = [
        {
            "doc": "Trip limit for PT-101 is 45.0 barg as per ASME BPVC Section VIII.",
            "meta": {
                "source_id": 10,
                "processing_version": "v1",
                "filename": "spec.docx",
                "document_type": "docx",
                "section_heading": "Pressure Setpoints",
                "section_index": 1
            },
            "score": 0.92,
            "source_id": 10,
            "filename": "spec.docx"
        }
    ]

    visual_results = [
        VisualSearchResult(
            workspace_id=1,
            source_id=10,
            processing_version="v1",
            page_number=2,
            filename="spec.docx",
            document_type="docx",
            locator_kind="page",
            score=5.1
        )
    ]

    fused = fusion.fuse(text_items, visual_results, query="What is the trip limit for PT-101?")
    assert len(fused) == 2

    # Find text candidate
    txt_cand = next(c for c in fused if c.retrieval_channel == "text")
    assert txt_cand.document_type == "docx"
    assert txt_cand.citation == "[spec.docx | Section: Pressure Setpoints | TEXT]"

    # Find visual candidate
    vis_cand = next(c for c in fused if c.retrieval_channel == "visual")
    assert vis_cand.document_type == "docx"
    assert vis_cand.citation == "[spec.docx | Page 2 | VISUAL]"
