import pytest
from pathlib import Path
from PIL import Image
import docx
import fitz

from cognishift.app.db.database import get_db, init_db
from cognishift.core.artifact_generators import (
    generate_docx_document,
    generate_pdf_document,
    create_and_register_artifact,
    resolve_image_for_embedding,
    SecurityError,
)


@pytest.fixture
def sample_chart_png(tmp_path: Path) -> Path:
    img_path = tmp_path / 'sample_chart.png'
    im = Image.new('RGB', (600, 400), color='#1F4E79')
    im.save(str(img_path))
    return img_path


def test_generate_docx_with_embedded_image(tmp_path: Path, sample_chart_png: Path):
    docx_out = tmp_path / 'test_report.docx'
    sections = [
        {
            'heading': '1. Executive Summary',
            'paragraphs': ['Overview of the quarterly operating regime.']
        },
        {
            'heading': '2. Visual Observations & Telemetry',
            'paragraphs': ['Discharge pressure trends displayed below.'],
            'images': [
                {'path': str(sample_chart_png), 'caption': 'Figure 1: Discharge Pressure Trajectory'}
            ]
        }
    ]

    generate_docx_document(docx_out, 'Compressor Report', sections)
    assert docx_out.exists()

    doc = docx.Document(str(docx_out))
    assert len(doc.paragraphs) >= 3
    # Verify caption and inline shapes
    para_texts = [p.text for p in doc.paragraphs]
    assert any('Figure: Figure 1: Discharge Pressure Trajectory' in t for t in para_texts)
    assert len(doc.inline_shapes) >= 1


def test_generate_pdf_with_embedded_image(tmp_path: Path, sample_chart_png: Path):
    pdf_out = tmp_path / 'test_report.pdf'
    sections = [
        {
            'heading': '1. Executive Summary',
            'paragraphs': ['Refinery operations telemetry summary.']
        },
        {
            'heading': '2. Visual Observations & Telemetry',
            'paragraphs': ['Discharge pressure chart below.'],
            'images': [
                {'path': str(sample_chart_png), 'caption': 'Discharge Pressure Trend'}
            ]
        }
    ]

    generate_pdf_document(pdf_out, 'Operational Report', sections)
    assert pdf_out.exists()

    pdf_doc = fitz.open(str(pdf_out))
    assert len(pdf_doc) >= 1
    # Check that image is embedded in the PDF
    page = pdf_doc[0]
    img_list = page.get_images()
    assert len(img_list) >= 1
    pdf_doc.close()


@pytest.mark.asyncio
async def test_artifact_registration_and_backlinking(tmp_path: Path, sample_chart_png: Path):
    await init_db()
    ws_id = 999
    async with get_db() as db:
        await db.execute('INSERT OR IGNORE INTO workspaces (id, name, description) VALUES (?, ?, ?)', (ws_id, 'Test WS', 'Desc'))
        await db.commit()

    # 1. Register a PNG chart artifact
    chart_artifact = await create_and_register_artifact(
        workspace_id=ws_id,
        filename='chart_p101.png',
        artifact_type='png',
        generator_fn=lambda p: Image.open(str(sample_chart_png)).save(str(p)),
        title='P-101 Pressure Chart',
        description='Pressure telemetry graph'
    )
    chart_id = chart_artifact['id']

    # 2. Resolve image reference securely
    res_path, res_cap, res_id = await resolve_image_for_embedding(ws_id, chart_id)
    assert res_path is not None
    assert res_id == chart_id

    # Cross-workspace security check
    with pytest.raises(SecurityError):
        await resolve_image_for_embedding(workspace_id=888, image_ref=chart_id)

    # 3. Register PDF with embedded chart metadata
    sections = [
        {
            'heading': 'Section 1',
            'paragraphs': ['Analysis text'],
            'images': [{'path': str(res_path), 'caption': res_cap}]
        }
    ]
    doc_artifact = await create_and_register_artifact(
        workspace_id=ws_id,
        filename='report_final.pdf',
        artifact_type='pdf',
        generator_fn=lambda p: generate_pdf_document(p, 'Official Deliverable', sections),
        title='Official Deliverable',
        description='Final report with chart',
        metadata={'embedded_visualization_artifact_ids': [chart_id]}
    )
    doc_id = doc_artifact['id']

    # 4. Verify backlink in chart artifact metadata
    import json
    async with get_db() as db:
        c = await db.execute('SELECT metadata FROM workspace_artifacts WHERE id = ?', (chart_id,))
        row = await c.fetchone()
        meta = json.loads(row['metadata'])
        assert doc_id in meta.get('embedded_in_artifact_ids', [])
