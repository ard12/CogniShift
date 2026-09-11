"""
Document Type Detection & Triage Inspector.
Deterministic inspection of file types, page bounds, image dimensions, and native extraction quality.
"""
from pathlib import Path
from typing import Tuple
import pymupdf  # PyMuPDF
from PIL import Image

from cognishift.app.config import settings
from cognishift.core.document_processing.schemas import (
    DocumentType,
    DocumentInspectionResult,
    UnsupportedFileError,
    CorruptedDocumentError,
    ResourceLimitExceededError
)

import zipfile
import docx
import openpyxl

# Magic bytes
PDF_MAGIC = b"%PDF-"
PNG_MAGIC = b"\x89PNG\r\n\x1a\n"
JPEG_MAGIC = b"\xff\xd8\xff"
ZIP_MAGIC = b"PK\x03\x04"
OLE2_MAGIC = b"\xd0\xcf\x11\xe0"  # Legacy .doc / .xls


def detect_file_type(file_path: Path) -> DocumentType:
    """Detect file type via magic byte header inspection and OOXML inner structure."""
    if not file_path.exists():
        raise UnsupportedFileError(f"File does not exist: {file_path}")

    with open(file_path, "rb") as f:
        header = f.read(16)

    # Reject legacy binary office formats (.doc, .xls)
    if header.startswith(OLE2_MAGIC):
        raise UnsupportedFileError(
            f"Legacy binary office format (.doc/.xls) detected in '{file_path.name}'. "
            "Please convert the file to modern open formats (.docx, .xlsx) or PDF."
        )

    if header.startswith(PDF_MAGIC):
        return DocumentType.PDF
    elif header.startswith(PNG_MAGIC):
        return DocumentType.PNG
    elif header.startswith(JPEG_MAGIC):
        return DocumentType.JPEG
    elif header.startswith(ZIP_MAGIC) or file_path.suffix.lower() in [".docx", ".xlsx"]:
        try:
            with zipfile.ZipFile(file_path) as z:
                names = z.namelist()
                if any(n.startswith("word/") for n in names) or file_path.suffix.lower() == ".docx":
                    return DocumentType.DOCX
                if any(n.startswith("xl/") for n in names) or file_path.suffix.lower() == ".xlsx":
                    return DocumentType.XLSX
        except Exception:
            pass

    # Check by extension as fallback
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return DocumentType.PDF
    elif ext in [".png"]:
        return DocumentType.PNG
    elif ext in [".jpg", ".jpeg"]:
        return DocumentType.JPEG
    elif ext == ".docx":
        return DocumentType.DOCX
    elif ext in [".xlsx", ".xlsm"]:
        return DocumentType.XLSX
    elif ext == ".csv":
        return DocumentType.CSV

    return DocumentType.UNSUPPORTED


def inspect_document(file_path: Path) -> DocumentInspectionResult:
    """
    Deterministically inspects the document structure, page count, and resource bounds.
    Fails closed if corrupt, unsupported, or exceeding resource quotas.
    """
    doc_type = detect_file_type(file_path)
    if doc_type == DocumentType.UNSUPPORTED:
        raise UnsupportedFileError(
            f"Unsupported file format for '{file_path.name}'. Permitted formats: PDF, DOCX, XLSX, CSV, PNG, JPEG."
        )

    file_size = file_path.stat().st_size
    max_bytes = settings.max_upload_size_mb * 1024 * 1024
    if file_size > max_bytes:
        raise ResourceLimitExceededError(
            f"Document size {file_size} bytes exceeds configured limit of {max_bytes} bytes ({settings.max_upload_size_mb}MB)."
        )

    if doc_type == DocumentType.PDF:
        try:
            doc = pymupdf.open(str(file_path))
            page_count = len(doc)
            doc.close()
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to parse PDF document '{file_path.name}': {e}")

        if page_count > settings.max_pdf_pages:
            raise ResourceLimitExceededError(
                f"PDF document has {page_count} pages, which exceeds the limit of {settings.max_pdf_pages} pages."
            )

        return DocumentInspectionResult(
            document_type=doc_type,
            is_valid=True,
            page_count=page_count,
            file_size_bytes=file_size,
            mime_type="application/pdf"
        )

    elif doc_type == DocumentType.DOCX:
        try:
            doc = docx.Document(str(file_path))
            paras = len(doc.paragraphs)
            tables = len(doc.tables)
            approx_pages = max(1, (paras + tables * 3) // 4)
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to parse DOCX document '{file_path.name}': {e}")

        return DocumentInspectionResult(
            document_type=doc_type,
            is_valid=True,
            page_count=approx_pages,
            file_size_bytes=file_size,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    elif doc_type == DocumentType.XLSX:
        try:
            wb = openpyxl.load_workbook(str(file_path), read_only=True)
            sheet_count = len(wb.sheetnames)
            wb.close()
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to parse XLSX document '{file_path.name}': {e}")

        return DocumentInspectionResult(
            document_type=doc_type,
            is_valid=True,
            page_count=sheet_count,
            file_size_bytes=file_size,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    elif doc_type == DocumentType.CSV:
        return DocumentInspectionResult(
            document_type=doc_type,
            is_valid=True,
            page_count=1,
            file_size_bytes=file_size,
            mime_type="text/csv"
        )

    else:
        # Image inspection (PNG / JPEG)
        try:
            # Set decompression bomb limit
            Image.MAX_IMAGE_PIXELS = settings.max_input_image_pixels
            with Image.open(file_path) as img:
                width, height = img.size
                mime = Image.MIME.get(img.format, "image/png" if doc_type == DocumentType.PNG else "image/jpeg")
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to parse image '{file_path.name}': {e}")

        if width > settings.max_input_image_dimension or height > settings.max_input_image_dimension:
            raise ResourceLimitExceededError(
                f"Image dimensions ({width}x{height}) exceed maximum allowed dimension of {settings.max_input_image_dimension}px."
            )

        if width * height > settings.max_input_image_pixels:
            raise ResourceLimitExceededError(
                f"Image pixel count ({width * height}) exceeds maximum allowed pixels of {settings.max_input_image_pixels}."
            )

        return DocumentInspectionResult(
            document_type=doc_type,
            is_valid=True,
            page_count=1,
            file_size_bytes=file_size,
            mime_type=mime
        )


def assess_native_page_quality(page_text: str, image_count: int = 0) -> bool:
    """
    Deterministic heuristics to determine whether native PDF text extraction is sufficient.
    Returns True if native text is usable; False if targeted OCR fallback is required.
    """
    cleaned = page_text.strip()
    if len(cleaned) < 50:
        return False

    printable_count = sum(1 for c in cleaned if c.isprintable() or c in ('\n', '\r', '\t', ' '))
    if printable_count / max(len(cleaned), 1) < 0.85:
        return False

    # Check for excessive replacement or unmapped characters
    replacement_count = cleaned.count('\ufffd') + cleaned.count('?')
    if replacement_count / max(len(cleaned), 1) > 0.05:
        return False

    # If page has very little text but contains embedded images, it is likely a scanned form or diagram
    if len(cleaned) < 150 and image_count > 0:
        return False

    return True
