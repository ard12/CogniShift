"""Exhaustive tests for tool parameter validation resilience, format normalization, and plan durability."""
import pytest
from pathlib import Path
import tempfile
import openpyxl
import pandas as pd
from PIL import Image

from cognishift.core.tool_schemas import (
    validate_proposed_tool_call,
    bounded_repair_tool_parameters,
    validate_safe_relative_path,
    parse_agent_action,
    FinalAnswer,
    StepObservation,
    RenderDocumentPageArgs,
    GenerateDocxArgs,
    GenerateXlsxArgs,
    GeneratePptxArgs,
    GeneratePdfArgs,
    ExecuteCodeArgs,
    DocxSection,
    XlsxSheet,
    XlsxRow,
    PptxSlide
)
from cognishift.core.artifact_generators import render_document_page_to_image
from cognishift.core.semantic_router import get_semantic_router, SemanticIntent


def test_render_document_page_uppercase_format_and_extension():
    """Verify that uppercase format 'PNG' and missing/uppercase extensions are normalized."""
    # Test Run #100473 exact failure payload
    result = validate_proposed_tool_call(
        tool_name="render_document_page",
        raw_parameters={
            "source_path_or_id": "MRPL_3Year_Financial_and_Operational_Audit.xlsx",
            "page_number": 1,
            "format": "PNG",
            "output_filename": "grm_trend_figure.png"
        },
        allowed_tools=["render_document_page"]
    )
    assert result.valid is True
    assert result.validated_parameters["format"] == "png"
    assert result.validated_parameters["output_filename"] == "grm_trend_figure.png"


def test_render_document_page_format_variants():
    """Verify that different format variations (.jpg, image/jpeg, JPG) are normalized."""
    for raw_fmt, expected_fmt in [
        ("JPG", "jpg"),
        ("JPEG", "jpeg"),
        (".png", "png"),
        ("image/png", "png"),
        ("image/jpeg", "jpeg"),
        ("INVALID", "png")  # fallback default
    ]:
        obj = RenderDocumentPageArgs(
            source_path_or_id="doc.pdf",
            format=raw_fmt,
            output_filename="output"  # missing extension
        )
        assert obj.format == expected_fmt
        assert obj.output_filename == "output.png"


def test_generate_docx_casing_and_extension_repair():
    """Verify that GenerateDocxArgs handles uppercase .DOCX and missing extensions."""
    args1 = GenerateDocxArgs(
        filename="Operational_Report.DOCX",
        title="Audit Report",
        sections=[DocxSection(heading="Executive Summary", level=1, paragraphs=["All nominal."])]
    )
    assert args1.filename == "Operational_Report.docx"

    args2 = GenerateDocxArgs(
        filename="Operational_Report",
        title="Audit Report",
        sections=[DocxSection(heading="Executive Summary", level=1, paragraphs=["All nominal."])]
    )
    assert args2.filename == "Operational_Report.docx"


def test_generate_xlsx_casing_and_extension_repair():
    """Verify that GenerateXlsxArgs handles uppercase .XLSX and missing extensions."""
    args1 = GenerateXlsxArgs(
        filename="Telemetry_Audit.XLSX",
        title="Telemetry Data",
        sheets=[XlsxSheet(name="Unit1", headers=["Time", "PSI"], rows=[XlsxRow(cells=["12:00", 14.5])])]
    )
    assert args1.filename == "Telemetry_Audit.xlsx"

    args2 = GenerateXlsxArgs(
        filename="Telemetry_Audit",
        title="Telemetry Data",
        sheets=[XlsxSheet(name="Unit1", headers=["Time", "PSI"], rows=[XlsxRow(cells=["12:00", 14.5])])]
    )
    assert args2.filename == "Telemetry_Audit.xlsx"


def test_generate_pptx_casing_and_extension_repair():
    """Verify that GeneratePptxArgs handles uppercase .PPTX and missing extensions."""
    args1 = GeneratePptxArgs(
        filename="Deck.PPTX",
        title="Briefing",
        slides=[PptxSlide(title="Overview", bullet_points=["Point 1"])]
    )
    assert args1.filename == "Deck.pptx"

    args2 = GeneratePptxArgs(
        filename="Deck",
        title="Briefing",
        slides=[PptxSlide(title="Overview", bullet_points=["Point 1"])]
    )
    assert args2.filename == "Deck.pptx"


def test_generate_pdf_casing_and_extension_repair():
    """Verify that GeneratePdfArgs handles uppercase .PDF and missing extensions."""
    args1 = GeneratePdfArgs(
        filename="Doc.PDF",
        title="Manual",
        sections=[DocxSection(heading="Overview", level=1, paragraphs=["Content"])]
    )
    assert args1.filename == "Doc.pdf"

    args2 = GeneratePdfArgs(
        filename="Doc",
        title="Manual",
        sections=[DocxSection(heading="Overview", level=1, paragraphs=["Content"])]
    )
    assert args2.filename == "Doc.pdf"


def test_execute_code_casing_and_extension_repair():
    """Verify that ExecuteCodeArgs handles uppercase .PY and bounded repair handles missing extensions."""
    args1 = ExecuteCodeArgs(
        code="print('Hello')",
        entrypoint="SCRIPT.PY"
    )
    assert args1.entrypoint == "SCRIPT.py"

    repaired, _ = bounded_repair_tool_parameters(
        "execute_code",
        {"code": "print('Hello')", "entrypoint": "runner"}
    )
    assert repaired["entrypoint"] == "runner.py"



def test_bounded_repair_universal_hygiene():
    """Verify bounded repair normalizes without requiring explicit references object."""
    repaired, src = bounded_repair_tool_parameters(
        "render_document_page",
        {"source_path_or_id": "1", "format": "PNG", "output_filename": "chart"}
    )
    assert repaired["format"] == "png"
    assert repaired["output_filename"] == "chart.png"

    # Equipment ID whitespace and dot trimming
    repaired_eq, _ = bounded_repair_tool_parameters(
        "check_pressure",
        {"sensor_id": " PT-101. "}
    )
    assert repaired_eq["sensor_id"] == "PT-101"

    # File path leading ./ stripping
    repaired_fp, _ = bounded_repair_tool_parameters(
        "file_read",
        {"file_path": "./documents/sop.pdf"}
    )
    assert repaired_fp["file_path"] == "documents/sop.pdf"


def test_validate_safe_relative_path_leading_dotslash():
    """Verify that safe leading ./ is sanitized while path traversal is still blocked."""
    assert validate_safe_relative_path("./documents/report.pdf") == "documents/report.pdf"
    assert validate_safe_relative_path("documents/report.pdf") == "documents/report.pdf"

    with pytest.raises(ValueError, match="Path traversal"):
        validate_safe_relative_path("../etc/passwd")

    with pytest.raises(ValueError, match="Leading slash forbidden"):
        validate_safe_relative_path("/etc/passwd")


def test_render_document_page_to_image_xlsx_and_csv():
    """Verify that render_document_page_to_image renders XLSX and CSV tables to images without error."""
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)

        # 1. Create a dummy XLSX
        xlsx_file = tmp_path / "test_data.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Financials"
        ws.append(["Year", "Revenue ($M)", "GRM ($/bbl)"])
        ws.append(["FY23", 1200, 8.45])
        ws.append(["FY24", 1350, 10.20])
        ws.append(["FY25", 1520, 11.80])
        wb.save(xlsx_file)
        wb.close()

        out_img1 = tmp_path / "xlsx_rendered.png"
        render_document_page_to_image(xlsx_file, out_img1, page_number=1, image_format="png")
        assert out_img1.exists()
        assert out_img1.stat().st_size > 500
        with Image.open(out_img1) as im:
            assert im.size[0] > 100

        # 2. Create a dummy CSV
        csv_file = tmp_path / "test_readings.csv"
        df = pd.DataFrame({
            "Timestamp": ["2026-09-01 10:00", "2026-09-01 10:05"],
            "Pressure_PSI": [102.5, 103.1]
        })
        df.to_csv(csv_file, index=False)

        out_img2 = tmp_path / "csv_rendered.png"
        render_document_page_to_image(csv_file, out_img2, page_number=1, image_format="png")
        assert out_img2.exists()
        assert out_img2.stat().st_size > 500


def test_semantic_router_grm_trend_query():
    """Verify that GRM and refining margin trend queries route cleanly to KNOWLEDGE_QUERY."""
    router = get_semantic_router()
    result = router.route("what was MRPL's gross refining margin trend over the last 3 years? break it down simply")
    assert result.intent == SemanticIntent.KNOWLEDGE_QUERY
    assert result.abstained is False


def test_parse_agent_action_slm_pseudo_action_tags():
    """Verify parse_agent_action handles Run #100474 pseudo-action tags and prose without crashing."""
    # 1. Exact Run #100474 failure case: prose ending with [action: 'final_answer']
    raw_output_1 = (
        "Total Revenue for FY 2023-24 was INR 1,05,214 Crores. "
        "Raw material expenses increased by 14.2% [MRPL_Annual_Report_FY24.pdf | Page 38].\n\n"
        "[action: 'final_answer']"
    )
    res_1 = parse_agent_action(raw_output_1, strict=False)
    assert isinstance(res_1, FinalAnswer)
    assert "1,05,214 Crores" in res_1.content
    assert "[action: 'final_answer']" not in res_1.content
    assert any("MRPL_Annual_Report_FY24.pdf" in c for c in res_1.citations)

    # 2. Case variation: [Action: final_answer]
    raw_output_2 = "Operational telemetry is nominal across all 4 pumps.\n[Action: final_answer]"
    res_2 = parse_agent_action(raw_output_2, strict=False)
    assert isinstance(res_2, FinalAnswer)
    assert "nominal across all 4 pumps" in res_2.content

    # 3. Case variation: Action: final_answer
    raw_output_3 = "Content: The compressor speed is within API-617 boundaries.\nAction: final_answer"
    res_3 = parse_agent_action(raw_output_3, strict=False)
    assert isinstance(res_3, FinalAnswer)
    assert "within API-617 boundaries" in res_3.content

    # 4. Step observation pseudo tag: [action: 'step_observation']
    raw_output_4 = "Extracted 4 rows from sensor telemetry table.\n[action: 'step_observation']"
    res_4 = parse_agent_action(raw_output_4, strict=False)
    assert isinstance(res_4, StepObservation)
    assert "Extracted 4 rows" in res_4.content

    # 5. Non-JSON prose containing bracketed text [Note: ...] should fall through to FinalAnswer
    raw_output_5 = "[Note: Audited values] Gross Refining Margin averaged $8.40/bbl over the 3-year period."
    res_5 = parse_agent_action(raw_output_5, strict=False)
    assert isinstance(res_5, FinalAnswer)
    assert "Gross Refining Margin averaged $8.40/bbl" in res_5.content


def test_semantic_router_capability_presentation_questions():
    """Verify Run #100481 capability inquiries route to CONVERSATION, while imperative actions route to CODE_EXECUTION."""
    router = get_semantic_router()

    # Capability queries without concrete target file
    for q in [
        "can you convert a pdf file to a ppt?",
        "can you convert a pdf to a ppt?",
        "can you convert files to presentations?",
        "can you convert documents to powerpoint?",
        "do you support pptx export?"
    ]:
        res = router.route(q)
        assert res.intent == SemanticIntent.CONVERSATION, f"Query '{q}' routed to {res.intent} instead of CONVERSATION"

    # Concrete imperative command with existing file
    cmd_res = router.route("convert P-101A_Inspection_Report.pdf to pptx")
    assert cmd_res.intent == SemanticIntent.CODE_EXECUTION


@pytest.mark.asyncio
async def test_generate_pptx_artifact_execution(tmp_path):
    """Verify that generate_pptx creates a valid PowerPoint presentation file on disk."""
    from cognishift.core.artifact_generators import generate_pptx_presentation
    import pptx

    pptx_path = tmp_path / "test_briefing.pptx"
    slides = [
        {"title": "Slide 1", "bullet_points": ["First point", "Second point"]},
        {"title": "Slide 2", "bullet_points": ["Point A", "Point B", "Point C"]}
    ]
    generate_pptx_presentation(pptx_path, title="Executive Test Briefing", subtitle="CogniShift Test", slides=slides)
    assert pptx_path.exists()
    assert pptx_path.stat().st_size > 1000

    prs = pptx.Presentation(str(pptx_path))
    assert len(prs.slides) == 3  # Title slide + 2 bullet slides

