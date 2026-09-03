"""
Permanent Test Suite for Phase 3: Workspace Filesystem, Artifact Registry, Generators, and Tamper Detection.
Phase 3 Verification for CogniShift.
"""

import pytest
import os
import json
from pathlib import Path
from unittest.mock import patch

from cognishift.app.config import settings
settings.operating_mode = "simulated"

from cognishift.core.security import (
    ensure_workspace_layout,
    get_workspace_root,
    resolve_workspace_path,
    SecurityError
)
from cognishift.core.tools import execute_tool
from cognishift.core.artifact_generators import (
    create_and_register_artifact,
    validate_artifact_structure,
    generate_docx_document,
    generate_xlsx_workbook,
    generate_pptx_presentation,
    compute_sha256
)
from cognishift.app.db.database import get_db, init_db
from fastapi.testclient import TestClient
from cognishift.app.main import app
from tests.conftest import (
    TEST_OPERATOR_TOKEN,
    TEST_SUPERVISOR_TOKEN,
    TEST_TENANT2_TOKEN
)

import docx
import openpyxl
import pptx


@pytest.fixture(autouse=True)
async def setup_artifacts_db():
    await init_db()
    async with get_db() as db:
        await db.execute("INSERT OR IGNORE INTO workspaces (id, name, description) VALUES (1, 'Refinery-1', 'MRPL')")
        await db.execute("INSERT OR IGNORE INTO workspaces (id, name, description) VALUES (2, 'Refinery-2', 'Mangalore')")
        await db.execute(
            "INSERT OR REPLACE INTO agent_runs (id, workspace_id, agent_id, status, user_id) VALUES (10, 1, 1, 'completed', 'operator_sam')"
        )
        await db.execute(
            "INSERT OR REPLACE INTO agent_runs (id, workspace_id, agent_id, status, user_id) VALUES (20, 2, 1, 'completed', 'operator_t2')"
        )
        await db.execute("DELETE FROM workspace_artifacts WHERE workspace_id IN (1, 2)")
        await db.commit()

    import shutil
    for ws_id in [1, 2]:
        ws_root = get_workspace_root(ws_id)
        shutil.rmtree(ws_root / "generated", ignore_errors=True)
        shutil.rmtree(ws_root / "temporary", ignore_errors=True)
        ensure_workspace_layout(ws_id)


# -----------------------------------------------------------------------------
# 1. STANDARDIZED WORKSPACE LAYOUT
# -----------------------------------------------------------------------------
def test_standardized_workspace_layout_creation():
    ws_root = ensure_workspace_layout(1)
    expected_dirs = ["uploads", "documents", "generated", "code", "temporary", "knowledge"]
    for d in expected_dirs:
        sub = ws_root / d
        assert sub.exists() and sub.is_dir(), f"Expected directory {d} missing from layout!"


# -----------------------------------------------------------------------------
# 2. SAFE FILE TOOLS & PATH TRAVERSAL RESISTANCE
# -----------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_safe_file_write_and_read():
    ensure_workspace_layout(1)
    p = resolve_workspace_path(1, "documents/sop_notes.txt", purpose="write")
    if p.exists():
        p.unlink()
    # 1. Write file
    write_res = await execute_tool(
        "file_write",
        {"file_path": "documents/sop_notes.txt", "content": "Suction pressure threshold: 120 PSI", "overwrite": False},
        workspace_id=1
    )
    assert "Successfully wrote" in write_res

    # 2. Overwrite default False: second write fails
    write_fail = await execute_tool(
        "file_write",
        {"file_path": "documents/sop_notes.txt", "content": "Modified content", "overwrite": False},
        workspace_id=1
    )
    assert "already exists and overwrite is set to False" in write_fail

    # 3. Read file
    read_res = await execute_tool(
        "file_read",
        {"file_path": "documents/sop_notes.txt"},
        workspace_id=1
    )
    assert read_res == "Suction pressure threshold: 120 PSI"


@pytest.mark.asyncio
async def test_file_tools_path_traversal_strictly_rejected():
    traversal_targets = [
        "../../outside.txt",
        "..\\windows.txt",
        "/etc/passwd",
        "C:\\Windows\\System32\\cmd.exe",
        "temporary/../../evil.txt"
    ]
    for target in traversal_targets:
        res_write = await execute_tool(
            "file_write",
            {"file_path": target, "content": "malicious write", "overwrite": True},
            workspace_id=1
        )
        assert "Error" in res_write or "Security" in res_write, f"Traversal write not blocked: {target}"

        res_read = await execute_tool(
            "file_read",
            {"file_path": target},
            workspace_id=1
        )
        assert "Error" in res_read or "Security" in res_read, f"Traversal read not blocked: {target}"


# -----------------------------------------------------------------------------
# 3. REGISTERED ARTIFACT IMMUTABILITY
# -----------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_generic_file_write_cannot_overwrite_registered_artifact():
    # 1. Create a valid DOCX artifact
    docx_res = await execute_tool(
        "generate_docx",
        {
            "filename": "inspection_report.docx",
            "title": "MRPL Unit 1 Inspection",
            "sections": [{"heading": "Visual Check", "level": 1, "paragraphs": ["All piping nominal."]}]
        },
        workspace_id=1,
        run_id=10
    )
    assert "Successfully generated and registered DOCX artifact" in docx_res

    # 2. Attempt to overwrite the generated file via generic file_write
    overwrite_res = await execute_tool(
        "file_write",
        {
            "file_path": "generated/run_10/inspection_report.docx",
            "content": "Corrupted content",
            "overwrite": True
        },
        workspace_id=1
    )
    assert "Cannot overwrite registered immutable artifact" in overwrite_res


# -----------------------------------------------------------------------------
# 4. TYPED ARTIFACT GENERATORS (DOCX, XLSX, PPTX)
# -----------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_generate_and_validate_docx():
    res = await execute_tool(
        "generate_docx",
        {
            "filename": "unit_debrief.docx",
            "title": "Shift Handover Debrief",
            "sections": [
                {
                    "heading": "Telemetry Overview",
                    "level": 1,
                    "paragraphs": ["Turbine speed nominal.", "Bearing temperature stabilized."],
                    "table": {
                        "headers": ["Tag", "Reading", "Unit"],
                        "rows": [["PT-101", "105.2", "PSI"], ["TT-204", "78.4", "C"]]
                    }
                }
            ]
        },
        workspace_id=1,
        run_id=10
    )
    assert "Successfully generated" in res

    # Reopen and inspect package
    doc_path = resolve_workspace_path(1, "generated/run_10/unit_debrief.docx", purpose="read")
    doc = docx.Document(str(doc_path))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Telemetry Overview" in headings


@pytest.mark.asyncio
async def test_generate_and_validate_xlsx_with_formula_sanitization():
    res = await execute_tool(
        "generate_xlsx",
        {
            "filename": "telemetry_dump.xlsx",
            "title": "SCADA Telemetry Dump",
            "sheets": [
                {
                    "name": "LiveSensors",
                    "headers": ["Sensor", "Value", "RawStatus"],
                    "rows": [
                        {"cells": ["PT-101", 105.2, "NOMINAL"]},
                        {"cells": ["TT-204", 78.4, "=SUM(A1:A10)"]},  # Attempted formula injection
                        {"cells": ["SV-402", 0, "@MALICIOUS_CMD"]}      # Attempted formula injection
                    ]
                }
            ]
        },
        workspace_id=1,
        run_id=10
    )
    assert "Successfully generated" in res

    # Verify formula sanitization in openpyxl
    wb_path = resolve_workspace_path(1, "generated/run_10/telemetry_dump.xlsx", purpose="read")
    wb = openpyxl.load_workbook(str(wb_path))
    ws = wb["LiveSensors"]
    # Row 6 is the second data row (=SUM...)
    val_row6 = ws.cell(row=6, column=3).value
    val_row7 = ws.cell(row=7, column=3).value
    assert val_row6.startswith("'"), "Formula '=' was not escaped!"
    assert val_row7.startswith("'"), "Formula '@' was not escaped!"


@pytest.mark.asyncio
async def test_generate_and_validate_pptx():
    res = await execute_tool(
        "generate_pptx",
        {
            "filename": "operations_briefing.pptx",
            "title": "MRPL Unit 1 Briefing",
            "subtitle": "Prepared by Autonomous Industrial Agent",
            "slides": [
                {
                    "title": "Executive Summary",
                    "bullet_points": ["Hydrotreater V-102 operating within design envelope.", "Zero high-priority alarms active."]
                }
            ]
        },
        workspace_id=1,
        run_id=10
    )
    assert "Successfully generated" in res

    prs_path = resolve_workspace_path(1, "generated/run_10/operations_briefing.pptx", purpose="read")
    prs = pptx.Presentation(str(prs_path))
    assert len(prs.slides) == 2  # Title slide + 1 content slide


# -----------------------------------------------------------------------------
# 5. ATOMIC ARTIFACT LIFECYCLE FAILURES
# -----------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_structural_validation_failure_leaves_no_orphan_or_db():
    temp_called = False
    def broken_generator(dest_path: Path):
        nonlocal temp_called
        temp_called = True
        # Write corrupted 0-byte file
        with open(dest_path, "wb") as f:
            pass

    with pytest.raises(RuntimeError, match="Structural artifact validation failed"):
        await create_and_register_artifact(
            workspace_id=1,
            filename="corrupt.docx",
            artifact_type="docx",
            generator_fn=broken_generator,
            title="Corrupted Report"
        )

    # Verify 0 files in generated/ and 0 records in DB
    ws_root = get_workspace_root(1)
    gen_dir = ws_root / "generated"
    corrupt_matches = list(gen_dir.glob("**/corrupt.docx"))
    assert len(corrupt_matches) == 0

    async with get_db() as db:
        c = await db.execute("SELECT COUNT(*) as cnt FROM workspace_artifacts WHERE filename = 'corrupt.docx'")
        assert (await c.fetchone())["cnt"] == 0


# -----------------------------------------------------------------------------
# 6. ARTIFACT VERSIONING (NO OVERWRITE ACROSS RUNS)
# -----------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_two_runs_same_filename_produce_independent_artifacts():
    # Run 10 requests Approval_Note.docx
    res1 = await execute_tool(
        "generate_docx",
        {"filename": "Approval_Note.docx", "title": "Run 10 Note", "sections": [{"heading": "H1", "level": 1, "paragraphs": ["P1"]}]},
        workspace_id=1,
        run_id=10
    )
    # Insert another run 11 in workspace 1
    async with get_db() as db:
        await db.execute("INSERT OR IGNORE INTO agent_runs (id, workspace_id, agent_id, status) VALUES (11, 1, 1, 'completed')")
        await db.commit()

    # Run 11 requests identical filename Approval_Note.docx
    res2 = await execute_tool(
        "generate_docx",
        {"filename": "Approval_Note.docx", "title": "Run 11 Note", "sections": [{"heading": "H2", "level": 1, "paragraphs": ["P2"]}]},
        workspace_id=1,
        run_id=11
    )
    assert "Successfully generated" in res1
    assert "Successfully generated" in res2

    async with get_db() as db:
        c = await db.execute("SELECT id, run_id, relative_path FROM workspace_artifacts WHERE workspace_id = 1 AND filename = 'Approval_Note.docx'")
        rows = await c.fetchall()
        assert len(rows) == 2
        paths = [r["relative_path"] for r in rows]
        assert "generated/run_10/Approval_Note.docx" in paths
        assert "generated/run_11/Approval_Note.docx" in paths


# -----------------------------------------------------------------------------
# 7. RUN / WORKSPACE OWNERSHIP INTEGRITY
# -----------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_register_artifact_with_cross_workspace_run_id_rejected():
    # Run 20 belongs to Workspace 2, but attempt registration in Workspace 1
    with pytest.raises(SecurityError, match="Cross-workspace run ownership violation"):
        await create_and_register_artifact(
            workspace_id=1,
            filename="cross_leak.docx",
            artifact_type="docx",
            generator_fn=lambda p: generate_docx_document(p, "Title", [{"heading": "H", "level": 1}]),
            title="Cross Run Leak",
            run_id=20  # Belongs to WS 2!
        )


# -----------------------------------------------------------------------------
# 8. TAMPER DETECTION ON REST DOWNLOAD
# -----------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_tamper_detection_refuses_modified_download():
    client = TestClient(app)
    headers = {"Authorization": f"Bearer {TEST_SUPERVISOR_TOKEN}"}

    # 1. Register valid artifact
    art = await create_and_register_artifact(
        workspace_id=1,
        filename="tamper_test.docx",
        artifact_type="docx",
        generator_fn=lambda p: generate_docx_document(p, "Tamper Test", [{"heading": "Original", "level": 1}]),
        title="Tamper Test"
    )
    art_id = art["id"]

    # 2. Verify normal download succeeds
    dl_res = client.get(f"/api/v1/workspaces/1/artifacts/{art_id}/download", headers=headers)
    assert dl_res.status_code == 200
    assert dl_res.headers.get("X-Checksum-SHA256") == art["sha256_hash"]

    # 3. Manually tamper with file bytes on disk
    file_path = resolve_workspace_path(1, art["relative_path"], purpose="write")
    with open(file_path, "ab") as f:
        f.write(b"TAMPERED_MALICIOUS_BYTES")

    # 4. Request download again -> Expect 409 Conflict
    dl_tampered = client.get(f"/api/v1/workspaces/1/artifacts/{art_id}/download", headers=headers)
    assert dl_tampered.status_code == 409
    assert "Tamper detection alert" in dl_tampered.json()["detail"]


# -----------------------------------------------------------------------------
# 9. CROSS-TENANT IDOR REST PROTECTION
# -----------------------------------------------------------------------------
def test_cross_tenant_idor_blocked_on_artifacts():
    client = TestClient(app)
    # Tenant 2 operator token (only authorized for WS 2)
    headers_t2 = {"Authorization": f"Bearer {TEST_TENANT2_TOKEN}"}

    # Tenant 2 attempts to list Workspace 1 artifacts
    res_list = client.get("/api/v1/workspaces/1/artifacts", headers=headers_t2)
    assert res_list.status_code == 403

    # Tenant 2 attempts to download Workspace 1 artifact
    res_dl = client.get("/api/v1/workspaces/1/artifacts/1/download", headers=headers_t2)
    assert res_dl.status_code == 403
