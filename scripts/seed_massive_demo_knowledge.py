"""
Seed Massive Industrial & Operational Demo Datasets into CogniShift Workspace 1.

Includes:
- 10-Year Audited Enterprise Financial & Operational History (12 Sheets, 5,000+ data rows)
- 50-Page API 610 / OISD-118/240 Petrochemical Engineering Master Manual
- 5,000-Row Master P&ID Instrumentation & Piping Registry
- 2,550-Row 5-Year Enterprise SAP PM Maintenance Work Order Log
- 12-Page Process Safety HAZOP & Risk Audit Report
"""

import os
import sys
import json
import shutil
import hashlib
import sqlite3
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR / "src"))

from cognishift.app.config import settings
from cognishift.core.retriever import chroma_client, embedding_model
import pymupdf
import openpyxl
import csv

DB_PATH = settings.database_path
DEMO_DIR = ROOT_DIR / "data" / "demo"
WS1_UPLOADS = ROOT_DIR / "data" / "workspaces" / "1" / "uploads"
WS1_DOCS = ROOT_DIR / "data" / "workspaces" / "1" / "documents"
WS1_UPLOADS.mkdir(parents=True, exist_ok=True)
WS1_DOCS.mkdir(parents=True, exist_ok=True)


def seed_massive_knowledge():
    print("=== SEEDING MASSIVE ENTERPRISE INDUSTRIAL ASSETS INTO WORKSPACE 1 ===")
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    collection = chroma_client.get_or_create_collection(name="workspace_1")

    target_files = [
        ("MRPL_Enterprise_Financial_and_Operational_History_10Y.xlsx", "spreadsheet"),
        ("MRPL_Centrifugal_Pumps_and_Hydrotreater_Operations_Manual_API610.pdf", "pdf"),
        ("MRPL_Master_PID_Instrumentation_and_Piping_Registry_5000.csv", "spreadsheet"),
        ("MRPL_Enterprise_Plant_Asset_and_Maintenance_Log_5Y.csv", "spreadsheet"),
        ("MRPL_Unit01_Hydrotreater_HAZOP_and_Risk_Audit.pdf", "pdf")
    ]

    for fname, stype in target_files:
        src_file = DEMO_DIR / fname
        if not src_file.exists():
            print(f"[SKIP] Source file missing: {src_file}")
            continue

        dest_upload = WS1_UPLOADS / fname
        dest_doc = WS1_DOCS / fname
        shutil.copy2(src_file, dest_upload)
        shutil.copy2(src_file, dest_doc)

        with open(src_file, "rb") as f:
            file_bytes = f.read()
        checksum = hashlib.sha256(file_bytes).hexdigest()

        c.execute("SELECT id FROM knowledge_sources WHERE workspace_id = 1 AND name = ?", (fname,))
        existing = c.fetchone()
        if existing:
            source_id = existing["id"]
            print(f"[UPDATE] Re-indexing '{fname}' (ID #{source_id})...")
            collection.delete(where={"source_id": source_id})
            c.execute("DELETE FROM document_pages WHERE source_id = ?", (source_id,))
        else:
            c.execute("""
                INSERT INTO knowledge_sources (workspace_id, name, source_type, original_filename, local_path, processing_status, checksum)
                VALUES (1, ?, ?, ?, ?, 'processing', ?)
            """, (fname, stype, fname, str(dest_upload), checksum))
            source_id = c.lastrowid
            print(f"[INSERT] Registered '{fname}' as Knowledge Source #{source_id}")

        chunks = []
        metadatas = []
        ids = []

        if stype == "pdf":
            doc = pymupdf.open(str(src_file))
            for pno in range(len(doc)):
                page = doc[pno]
                page_text = page.get_text()
                if not page_text.strip():
                    continue
                
                chunk_str = f"[{fname} | Page {pno + 1}]\n{page_text.strip()}"
                chunk_id = f"src_{source_id}_p{pno + 1}"
                chunks.append(chunk_str)
                metadatas.append({
                    "source_id": source_id,
                    "document_name": fname,
                    "filename": fname,
                    "page": pno + 1,
                    "workspace_id": 1,
                    "checksum": checksum,
                    "processing_version": "v1",
                    "extraction_method": "native_pdf"
                })
                ids.append(chunk_id)

                c.execute("""
                    INSERT INTO document_pages (source_id, workspace_id, processing_version, page_number, text_content, extraction_method)
                    VALUES (?, 1, 'v1', ?, ?, 'native_pdf')
                """, (source_id, pno + 1, page_text.strip()))
            doc.close()

        elif fname.endswith(".xlsx"):
            wb = openpyxl.load_workbook(str(src_file), data_only=True)
            chunk_counter = 1
            for sname in wb.sheetnames:
                ws = wb[sname]
                raw_rows = list(ws.iter_rows(values_only=True))
                if not raw_rows:
                    continue

                header_row = [str(x) if x is not None else "" for x in raw_rows[0]]
                h_text = " | ".join(header_row)
                batch_size = 40 if len(raw_rows) > 500 else 15

                for r_start in range(1, len(raw_rows), batch_size):
                    batch = raw_rows[r_start:r_start + batch_size]
                    b_lines = []
                    for r in batch:
                        r_str = " | ".join(str(x) if x is not None else "" for x in r)
                        if any(c.strip() for c in r_str.split("|")):
                            b_lines.append(r_str)
                    if not b_lines:
                        continue

                    chunk_text = f"=== SPREADSHEET: {fname} | SHEET: {sname} (Rows #{r_start+1} to #{r_start+len(batch)}) ===\nHeaders: {h_text}\n" + "\n".join(b_lines)
                    chunk_id = f"src_{source_id}_s{sname}_r{r_start}"
                    chunks.append(chunk_text)
                    metadatas.append({
                        "source_id": source_id,
                        "document_name": fname,
                        "filename": fname,
                        "sheet_name": sname,
                        "workspace_id": 1,
                        "checksum": checksum,
                        "processing_version": "v1",
                        "extraction_method": "spreadsheet"
                    })
                    ids.append(chunk_id)

                    c.execute("""
                        INSERT INTO document_pages (source_id, workspace_id, processing_version, page_number, text_content, extraction_method)
                        VALUES (?, 1, 'v1', ?, ?, 'spreadsheet')
                    """, (source_id, chunk_counter, chunk_text))
                    chunk_counter += 1

        elif fname.endswith(".csv"):
            with open(src_file, "r", encoding="utf-8", errors="replace") as f:
                reader = list(csv.reader(f))
            if reader:
                headers = " | ".join(reader[0])
                batch_size = 50 if len(reader) > 2000 else 30
                chunk_counter = 1
                for r_start in range(1, len(reader), batch_size):
                    batch = reader[r_start:r_start + batch_size]
                    b_lines = [" | ".join(r) for r in batch if any(r)]
                    if not b_lines:
                        continue

                    chunk_text = f"=== SPREADSHEET: {fname} (Rows #{r_start} to #{r_start+len(batch)}) ===\nHeaders: {headers}\n" + "\n".join(b_lines)
                    chunk_id = f"src_{source_id}_csv_r{r_start}"
                    chunks.append(chunk_text)
                    metadatas.append({
                        "source_id": source_id,
                        "document_name": fname,
                        "filename": fname,
                        "workspace_id": 1,
                        "checksum": checksum,
                        "processing_version": "v1",
                        "extraction_method": "spreadsheet"
                    })
                    ids.append(chunk_id)

                    c.execute("""
                        INSERT INTO document_pages (source_id, workspace_id, processing_version, page_number, text_content, extraction_method)
                        VALUES (?, 1, 'v1', ?, ?, 'spreadsheet')
                    """, (source_id, chunk_counter, chunk_text))
                    chunk_counter += 1

        print(f"  Embedding and indexing {len(chunks)} chunks into ChromaDB...")
        batch_emb_size = 64
        for i in range(0, len(chunks), batch_emb_size):
            sub_chunks = chunks[i:i+batch_emb_size]
            sub_metas = metadatas[i:i+batch_emb_size]
            sub_ids = ids[i:i+batch_emb_size]
            gen = embedding_model.embed(sub_chunks)
            embs = [e.tolist() if hasattr(e, "tolist") else [float(x) for x in e] for e in gen]
            collection.upsert(documents=sub_chunks, embeddings=embs, metadatas=sub_metas, ids=sub_ids)

        c.execute("""
            UPDATE knowledge_sources
            SET processing_status = 'completed', chunk_count = ?, active_processing_version = 'v1'
            WHERE id = ?
        """, (len(chunks), source_id))
        print(f"  [OK] Completed indexing for '{fname}' ({len(chunks)} chunks).")

    c.execute("SELECT id FROM knowledge_sources WHERE workspace_id = 1")
    all_ids = [r["id"] for r in c.fetchall()]
    c.execute("UPDATE agent_definitions SET knowledge_source_ids = ? WHERE id = 1", (json.dumps(all_ids),))

    conn.commit()
    conn.close()
    print(f"\n[SUCCESS] Completed Seeding! Total Knowledge Sources in Workspace 1: {len(all_ids)}")


if __name__ == "__main__":
    seed_massive_knowledge()
