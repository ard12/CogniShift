import pytest
from pathlib import Path
import docx

from cognishift.core.document_processing.schemas import DocumentType, UnsupportedFileError
from cognishift.core.document_processing.inspector import detect_file_type, inspect_document
from cognishift.core.document_processing.office_renderer import extract_docx_structured_content


def test_detect_docx_and_reject_legacy_doc(tmp_path: Path):
    docx_path = tmp_path / 'operating_manual.docx'
    doc = docx.Document()
    doc.add_heading('Compressor Operating Procedure', level=1)
    doc.add_paragraph('Ensure suction valve XV-101 is locked open before startup.')
    doc.save(str(docx_path))

    doc_type = detect_file_type(docx_path)
    assert doc_type == DocumentType.DOCX

    inspection = inspect_document(docx_path)
    assert inspection.is_valid is True
    assert inspection.document_type == DocumentType.DOCX
    assert inspection.mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert inspection.page_count >= 1

    doc_legacy_path = tmp_path / 'legacy_manual.doc'
    with open(doc_legacy_path, 'wb') as f:
        f.write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 64)

    with pytest.raises(UnsupportedFileError) as exc_info:
        detect_file_type(doc_legacy_path)
    assert 'Legacy binary office format' in str(exc_info.value)


def test_extract_docx_structured_content(tmp_path: Path):
    docx_path = tmp_path / 'procedure.docx'
    doc = docx.Document()
    doc.add_heading('Section 1: Initial Checks', level=1)
    doc.add_paragraph('Check lube oil pressure on PI-101.')
    doc.add_paragraph('Verify cooling water return temperature is below 45C.')

    doc.add_heading('Section 2: Valve Lineup Table', level=1)
    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Tag'
    hdr_cells[1].text = 'Description'
    hdr_cells[2].text = 'Required State'

    r1_cells = table.rows[1].cells
    r1_cells[0].text = 'FV-302'
    r1_cells[1].text = 'Feed Flow Control Valve'
    r1_cells[2].text = 'Auto / 50%'

    r2_cells = table.rows[2].cells
    r2_cells[0].text = 'XV-101'
    r2_cells[1].text = 'Suction Block Valve'
    r2_cells[2].text = 'Open'

    doc.save(str(docx_path))

    chunks = extract_docx_structured_content(docx_path)
    assert len(chunks) >= 2

    headings = [c['section_heading'] for c in chunks]
    assert any('Section 1' in h for h in headings)
    assert any('Section 2' in h for h in headings)

    table_chunk = next(c for c in chunks if '[Table' in c['text'])
    assert 'FV-302' in table_chunk['text']
    assert 'Feed Flow Control Valve' in table_chunk['text']
    assert 'XV-101' in table_chunk['text']
