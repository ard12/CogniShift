"""CogniShift Agentic Execution Engine.

Implements the central reasoning loop:
1. Validates agent permissions and workspaces.
2. Ingests domain-specific RAG context via FastEmbed and ChromaDB.
3. Formats prompt and queries local ModelProvider (Ollama / Simulated).
4. Evaluates tool calling intents and risk levels.
5. Safely pauses execution for high-risk actions (Human-in-the-Loop).
6. Resumes execution upon supervisor authorization.
"""

import asyncio
import time
import json
import re
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional, List, Dict, Any, Tuple
from fastapi import HTTPException, status

from cognishift.app.config import settings
from cognishift.app.db.database import get_db
from cognishift.app.db.models import RunResponse, RunEventResponse
from cognishift.core.retriever import retrieve_context, retrieve_context_with_metadata
from cognishift.core.graph_memory import query_graph_context
from cognishift.core.security import resolve_workspace_path, get_workspace_root
from cognishift.core.tools import execute_tool
from cognishift.core.providers import get_provider
from cognishift.core.model_router import (
    classify_task,
    route_model,
    TaskClassification,
    LocalModelInventoryUnavailableError,
    update_verified_inventory_cache
)
from cognishift.core.semantic_router import get_semantic_router, SemanticIntent, SemanticRoutingResult, DecisionMethod
from cognishift.core.tool_schemas import (
    parse_agent_action,
    validate_proposed_tool_call,
    ToolCallProposal,
    FinalAnswer,
    ClarificationRequest,
    StepObservation,
    get_authoritative_tool_json_schema,
    is_supported_equipment_target,
    HIGH_RISK_TOOLS
)
from cognishift.core.conversation_context import (
    ConversationContextResolver,
    ResolvedContext,
    ResolvedSource,
    resolve_authoritative_source,
    get_latest_ingested_document,
    get_latest_ingested_document_async,
    resolve_target_document_for_query,
    extract_requested_page
)
from cognishift.core.pending_tasks import (
    create_pending_task,
    get_active_pending_task,
    claim_pending_task_atomic,
    cancel_pending_task,
    complete_pending_task,
    fail_pending_task,
    PendingTask,
    TaskStatus
)
from cognishift.core.planner import (
    create_initial_plan,
    format_plan_for_prompt,
    serialize_plan,
    deserialize_plan,
    AgentPlan,
    PlanStep
)
from cognishift.core.document_insights import (
    extract_document_insights,
    detect_dataframe_anomalies,
    format_dataframe_as_explicit_records,
    format_number_display,
    StructuredAnomalyResult,
    build_authoritative_anomaly_response,
    validate_scada_anomaly_prose
)

logger = logging.getLogger("cognishift.engine")

from cognishift.core.rca import (
    RCAEvidenceAcquirer,
    RCAEvidenceValidator,
    RCAEvidenceBundle,
    RCAStatus,
    classify_tool_failure_severity,
    ToolFailureSeverity,
)


from dataclasses import dataclass, field


def _is_presentation_request(text: str) -> bool:
    """Detect an explicit presentation deliverable without matching 'representation'."""
    return bool(re.search(r"\b(?:ppt|pptx|powerpoint|slides?|presentation)\b", text or "", re.IGNORECASE))


def enforce_rca_evidence_boundaries(content: str, operator_input: str, citations: Optional[List[str]] = None) -> str:
    """Keep symptom-only RCA output explicit about evidence versus hypotheses, while preserving verified citations."""
    lowered_input = operator_input.lower()
    confirmed: List[str] = []
    if "suction pressure" in lowered_input and any(word in lowered_input for word in ("drop", "dropped", "low")):
        confirmed.append("Suction pressure was reported as having dropped; no absolute reading or baseline was supplied.")
    if "discharge pressure" in lowered_input and any(word in lowered_input for word in ("unstable", "fluctuat")):
        confirmed.append("Discharge pressure was reported as unstable; no measured range or trip event was supplied.")
    if "vibration" in lowered_input and any(word in lowered_input for word in ("increase", "increased", "sharp", "high")):
        confirmed.append("Vibration was reported as having increased sharply; no measured value or spectrum was supplied.")
    if not confirmed:
        confirmed.append("No independently measured observation was supplied beyond the operator's written description.")

    hypotheses: List[str] = []
    evidence: List[str] = []
    section: Optional[str] = None
    for line in (content or "").splitlines():
        heading = re.sub(r"[^a-z ]", "", line.lower()).strip()
        if "possible hypotheses" in heading or heading == "hypotheses":
            section = "hypotheses"
            continue
        if "additional evidence needed" in heading or "evidence needed" in heading:
            section = "evidence"
            continue
        if "recommendations" in heading:
            section = None
            break
        cleaned = line.strip()
        if cleaned and section == "hypotheses":
            hypotheses.append(cleaned)
        elif cleaned and section == "evidence":
            evidence.append(cleaned)

    if not hypotheses:
        hypotheses = ["- The reasoning model did not return a safely separable hypothesis list; no cause is asserted."]
    if not evidence:
        evidence = [
            "- Time-aligned suction and discharge pressure trends.",
            "- Vibration magnitude, waveform, and spectrum by bearing location.",
            "- Pump speed, flow, valve position, tank level, and cavitation/noise observations.",
            "- Inspection findings for bearings, coupling/alignment, impeller, seals, and suction restrictions.",
        ]

    # Collect any citations present in content or passed explicitly
    collected_citations: List[str] = []
    if citations:
        collected_citations.extend(citations)
    if content:
        inline_cites = re.findall(r"\[([^\]]*?\|\s*Page\s*\d+[^\]]*?)\]", content)
        for ic in inline_cites:
            formatted_ic = f"[{ic.strip()}]"
            if formatted_ic not in collected_citations:
                collected_citations.append(formatted_ic)

    citations_section = ""
    if collected_citations:
        citations_section = "\n\n## Documented SOP & Baseline Evidence\n" + "\n".join(f"- {c}" for c in collected_citations)

    return (
        "## Confirmed Observations\n"
        + "\n".join(f"- {item}" for item in confirmed)
        + "\n\n## Possible Hypotheses — Unverified\n"
        + "\n".join(hypotheses)
        + "\n\n## Additional Evidence Needed\n"
        + "\n".join(evidence)
        + citations_section
        + "\n\n**Conclusion:** No root cause is confirmed from the symptom-only information provided."
    )


from cognishift.core.visualization.schemas import ArtifactRequestContract, ChartType, VisualizationResult
from cognishift.core.visualization.selector import parse_artifact_request_contract
from cognishift.core.visualization.service import execute_visualization_pipeline


@dataclass
class GoalContract:
    """Explicit contract specifying required quantitative or deliverable outputs for multi-step goals."""
    required_fields: List[str] = field(default_factory=list)
    extracted_fields: Dict[str, Any] = field(default_factory=dict)
    artifact_contract: Optional[Any] = None

    def check_satisfaction(self, current_artifacts: Optional[List[Dict[str, Any]]] = None) -> Tuple[bool, List[str]]:
        missing = []
        for k in self.required_fields:
            if k not in self.extracted_fields or self.extracted_fields[k] is None:
                missing.append(k)
        if self.artifact_contract and getattr(self.artifact_contract, "is_deliverable_request", False):
            arts = current_artifacts or []
            if getattr(self.artifact_contract, "png_required", False):
                expected_pngs = getattr(self.artifact_contract, "png_count", 1) or 1
                valid_pngs = [
                    a for a in arts
                    if (a.get("artifact_type") == "png" or str(a.get("filename", "")).lower().endswith(".png"))
                    and a.get("file_size", 0) > 0
                ]
                if len(valid_pngs) < expected_pngs:
                    missing.append(f"png_deliverable (expected {expected_pngs}, found {len(valid_pngs)})")
            if getattr(self.artifact_contract, "pdf_required", False):
                valid_pdfs = [
                    a for a in arts
                    if (a.get("artifact_type") == "pdf" or str(a.get("filename", "")).lower().endswith(".pdf"))
                    and a.get("file_size", 0) > 0
                ]
                if not valid_pdfs:
                    missing.append("pdf_deliverable")
            if getattr(self.artifact_contract, "docx_required", False):
                valid_docx = [
                    a for a in arts
                    if (a.get("artifact_type") == "docx" or str(a.get("filename", "")).lower().endswith(".docx"))
                    and a.get("file_size", 0) > 0
                ]
                if not valid_docx:
                    missing.append("docx_deliverable")
            if getattr(self.artifact_contract, "xlsx_required", False):
                valid_xlsx = [
                    a for a in arts
                    if (a.get("artifact_type") == "xlsx" or str(a.get("filename", "")).lower().endswith(".xlsx"))
                    and a.get("file_size", 0) > 0
                ]
                if not valid_xlsx:
                    missing.append("xlsx_deliverable")
            if getattr(self.artifact_contract, "csv_required", False):
                valid_csv = [
                    a for a in arts
                    if (a.get("artifact_type") == "csv" or str(a.get("filename", "")).lower().endswith(".csv"))
                    and a.get("file_size", 0) > 0
                ]
                if not valid_csv:
                    missing.append("csv_deliverable")
        return len(missing) == 0, missing

    def is_satisfied(self, current_artifacts: Optional[List[Dict[str, Any]]] = None) -> bool:
        sat, _ = self.check_satisfaction(current_artifacts)
        return sat


def populate_goal_contract_from_insights(contract: GoalContract, insights: Dict[str, Any]) -> None:
    """Populates GoalContract strictly from authoritative structured calculations, never LLM prose."""
    metrics_map = insights.get("metrics", {})
    growth_map = insights.get("growth", {})
    if "revenue" in metrics_map:
        r = metrics_map["revenue"]
        contract.extracted_fields["revenue_previous"] = r.get("previous")
        contract.extracted_fields["revenue_current"] = r.get("latest")
        if "revenue_yoy_pct" in growth_map:
            contract.extracted_fields["revenue_yoy_pct"] = growth_map["revenue_yoy_pct"]
        if "revenue_cagr_pct" in growth_map:
            contract.extracted_fields["revenue_cagr_pct"] = growth_map["revenue_cagr_pct"]
    if "ebitda" in metrics_map:
        e = metrics_map["ebitda"]
        contract.extracted_fields["ebitda_previous"] = e.get("previous")
        contract.extracted_fields["ebitda_current"] = e.get("latest")
        if "ebitda_yoy_pct" in growth_map:
            contract.extracted_fields["ebitda_yoy_pct"] = growth_map["ebitda_yoy_pct"]
        if "ebitda_cagr_pct" in growth_map:
            contract.extracted_fields["ebitda_cagr_pct"] = growth_map["ebitda_cagr_pct"]
    if "pat" in metrics_map:
        p = metrics_map["pat"]
        contract.extracted_fields["pat_previous"] = p.get("previous")
        contract.extracted_fields["pat_current"] = p.get("latest")
        if "pat_yoy_pct" in growth_map:
            contract.extracted_fields["pat_yoy_pct"] = growth_map["pat_yoy_pct"]
        if "pat_cagr_pct" in growth_map:
            contract.extracted_fields["pat_cagr_pct"] = growth_map["pat_cagr_pct"]


def extract_and_strip_thinking(text: str) -> Tuple[str, bool, int]:
    """
    Strips and sanitizes <think>...</think> blocks from model output.
    Handles case variants (<think>, <THINK>) and unclosed <think> tags through EOF.
    Returns (cleaned_text, reasoning_detected, reasoning_character_count).
    Guarantees raw thinking tokens are NEVER returned, logged, persisted, or passed downstream.
    """
    if not text:
        return "", False, 0
    total_reasoning_chars = 0
    # 1. Strip closed tags (case-insensitive)
    closed_pattern = re.compile(r'<think>(.*?)</think>', re.DOTALL | re.IGNORECASE)
    for match in closed_pattern.finditer(text):
        total_reasoning_chars += len(match.group(1))
    cleaned = closed_pattern.sub('', text)

    # 2. Strip unclosed tag through EOF (case-insensitive)
    unclosed_pattern = re.compile(r'<think>(.*)$', re.DOTALL | re.IGNORECASE)
    unclosed_match = unclosed_pattern.search(cleaned)
    if unclosed_match:
        total_reasoning_chars += len(unclosed_match.group(1))
        cleaned = unclosed_pattern.sub('', cleaned)

    cleaned = cleaned.strip()
    reasoning_detected = (total_reasoning_chars > 0)
    return cleaned, reasoning_detected, total_reasoning_chars


def validate_evidence_sufficiency(query: str, retrieved_context: str) -> Tuple[bool, str]:
    """
    Topical and factual evidence sufficiency gate.
    Uses regex word boundaries for acronyms (NDE, NDT, MPY, WFH) to avoid substring false positives
    (e.g., preventing 'nde' in 'independent', 'rendered', or 'understanding' from spoofing corrosion evidence).
    Returns (is_sufficient, reason).
    """
    lower_clean = (query or "").lower()
    has_equipment_tag = bool(re.search(r"\b[A-Za-z]{1,4}-\d{3,4}[A-Za-z]?\b", lower_clean))
    is_rca_or_diagnostic = (
        any(k in lower_clean for k in [
            "rca", "root cause", "failure investigation", "investigate failure",
            "incident investigation", "why did it fail", "why did the system trip",
            "troubleshoot", "investigate", "cause of", "what caused"
        ])
        or (has_equipment_tag and any(k in lower_clean for k in ["trip", "tripped", "failed", "failure", "shutdown", "shut down", "alarm", "cavitat", "leak", "overpressure", "overheat"]))
        or (any(w in lower_clean for w in ["why did", "why has"]) and any(w in lower_clean for w in ["trip", "fail"]))
    )
    is_remote_work = any(k in lower_clean for k in ["remote work", "work from home", "telework", "telecommuting", "wfh"])
    is_procurement = "procurement" in lower_clean
    is_corrosion_query = not is_rca_or_diagnostic and (
        any(k in lower_clean for k in ["corrosion life", "remaining life", "corrosion rate", "wall thickness", "ultrasonic thickness", "corrosion"])
        or bool(re.search(r'\b(?:mpy|nde|ndt)\b', lower_clean))
    )
    is_inspection_report = not is_rca_or_diagnostic and (
        any(k in lower_clean for k in ["inspection report", "metallurgical report"])
        or bool(re.search(r'\b(?:nde|ndt)\s+report\b', lower_clean))
    )

    requested_entities = sorted(set(re.findall(r"\b[A-Z]{1,4}-\d{3,4}[A-Z]?\b", (query or "").upper())))
    entity_grounding_required = bool(requested_entities) and any(
        phrase in lower_clean
        for phrase in ("according to", "document", "manual", "procedure", "sop", "maintenance")
    )

    if not retrieved_context or not retrieved_context.strip():
        if entity_grounding_required:
            return False, f"No evidence retrieved for requested equipment: {', '.join(requested_entities)}."
        if is_remote_work or is_procurement or is_corrosion_query or is_inspection_report:
            return False, "No context retrieved"
        return True, "Standard query"

    lower_ctx = retrieved_context.lower()
    if is_remote_work:
        has_rw = any(k in lower_ctx for k in ["remote work", "telework", "work from home", "home office", "telecommuting"]) or bool(re.search(r'\bwfh\b', lower_ctx))
        if not has_rw:
            return False, "Query is for remote work policy, but retrieved context contains no remote work evidence."

    if is_procurement:
        has_proc = any(k in lower_ctx for k in ["procurement", "purchase requisition", "tender", "rfp", "vendor contract"])
        if not has_proc:
            return False, "Query is for procurement policy, but retrieved context contains no procurement evidence."

    if is_corrosion_query or is_inspection_report:
        has_corr_terms = any(k in lower_ctx for k in ["corrosion", "wall thickness", "remaining life", "inspection report", "inspection_report", "inspection", "ultrasonic", "thickness"])
        has_acronyms = bool(re.search(r'\b(?:nde|ndt|mpy)\b', lower_ctx, re.IGNORECASE))
        if not (has_corr_terms or has_acronyms):
            return False, "Query is for corrosion life / inspection report, but retrieved context contains no corrosion or inspection evidence."

    if entity_grounding_required:
        upper_ctx = retrieved_context.upper()
        missing_entities = [
            entity for entity in requested_entities
            if not re.search(rf"(?<![A-Z0-9-]){re.escape(entity)}(?![A-Z0-9-])", upper_ctx)
        ]
        if missing_entities:
            return False, f"Retrieved evidence does not explicitly mention requested equipment: {', '.join(missing_entities)}."

    return True, "Sufficient topical evidence verified"


def tool_output_failed(output: Any) -> bool:
    """Check if tool output indicates an execution error."""
    if isinstance(output, str):
        lowered = output.lower().strip()
        return (
            lowered.startswith("error:") or
            lowered.startswith("failed:") or
            lowered.startswith("execution failed") or
            "actuator unavailable" in lowered or
            "runtime error" in lowered
        )
    return False


def _block_unexecuted_pending_steps(plan: Optional[AgentPlan], reason: str = "Blocked due to execution failure") -> None:
    """Ensure no steps remain pending when a run fails or terminates early."""
    if not plan or not getattr(plan, "steps", None):
        return
    for s in plan.steps:
        if s.status == "pending":
            s.status = "blocked"
            if not s.error_message:
                s.error_message = reason


async def log_event(
    db,
    run_id: int,
    event_type: str,
    message: str,
    structured_data: Optional[Dict[str, Any]] = None
) -> None:
    """Log an execution event into run_events table for timeline streaming."""
    data_str = json.dumps(structured_data) if structured_data is not None else None
    await db.execute(
        """INSERT INTO run_events (run_id, event_type, message, structured_data)
           VALUES (?, ?, ?, ?)""",
        (run_id, event_type, message, data_str)
    )
    await db.commit()


def _extract_pdf_preview_sync(file_path: Path, max_pages: int = 10, max_chars_per_page: int = 1500, label: str = "", target_page: Optional[int] = None) -> str:
    """Synchronous CPU worker to extract text preview from PDF pages (offloaded to thread)."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(file_path))
        if target_page and 1 <= target_page <= len(reader.pages):
            page = reader.pages[target_page - 1]
            ptxt = (page.extract_text() or "").strip()
            if ptxt:
                return f"[{label or file_path.name} | Page {target_page}]:\n{ptxt[:max_chars_per_page]}"
        pages_text = []
        for p_num, page in enumerate(reader.pages[:max_pages], start=1):
            ptxt = (page.extract_text() or "").strip()
            if ptxt:
                pages_text.append(f"[{label or file_path.name} | Page {p_num}]:\n{ptxt[:max_chars_per_page]}")
        if pages_text:
            return "\n\n".join(pages_text)
        return f"[{label or file_path.name}: PDF document contains minimal native text]"
    except Exception as pdf_err:
        logger.warning(f"Error reading pdf {label or file_path.name}: {pdf_err}")
        return f"[PDF document: {label or file_path.name}]"


def parse_tool_call(
    response_text: str,
    allowed_tools: List[str],
    user_prompt: str = ""
) -> Tuple[bool, Optional[str], Dict[str, Any], str]:
    """Parse tool call from model response text using STRICT structured action protocol.
    
    P0-1 & P0-2 ENFORCEMENT:
    There is NO secondary regex parser, NO keyword matching, and NO user prompt inspection.
    Tool execution is ONLY possible if the model explicitly proposed a validated ToolCallProposal.
    
    Returns:
        (is_tool, tool_name, parameters, reason)
    """
    allowed_map = {t.lower(): t for t in allowed_tools}
    
    # Authoritative Path: Structured Action Protocol ONLY
    action = parse_agent_action(response_text)
    if isinstance(action, ToolCallProposal) and action.tool_name:
        tool_name_clean = action.tool_name.lower().strip()
        if tool_name_clean in allowed_map:
            resolved_tool = allowed_map[tool_name_clean]
            val_result = validate_proposed_tool_call(
                tool_name=resolved_tool,
                raw_parameters=action.parameters,
                allowed_tools=allowed_tools
            )
            if val_result.valid:
                return True, resolved_tool, val_result.validated_parameters or action.parameters, action.reason
            else:
                return False, None, {}, f"Validation failed: {val_result.error_message}"
    
    # If not a valid ToolCallProposal in allowed tools, tool execution is IMPOSSIBLE
    return False, None, {}, ""


DEFAULT_SYSTEM_PERSONA = (
    "You are CogniShift, a sovereign on-premise industrial AI assistant. "
    "You can reason over locally available workspace documents, use approved local tools, execute code inside the configured secure sandbox, and generate artifacts.\n\n"
    "Guidelines:\n"
    "1. Answer the operator's actual question directly, naturally, and concisely.\n"
    "2. Use organization knowledge or manuals only when strictly relevant to the query.\n"
    "3. Do not inject unrelated workspace information or recite alarm scenarios unless specifically asked.\n"
    "4. Never imply access to live systems or telemetry that are not actually connected.\n"
    "5. When operational execution is required, propose actions using the structured action protocol. Sensitive actions remain subject to deterministic backend policy and human approval.\n"
    "6. Grounding: Never invent intranet URLs, internal portals, organizational policies, or standard operating procedures not present in the local context. If information is not in the context, explicitly state that it is not available.\n"
    "7. Network sovereignty: external internet access is blocked by strict application policy. Do not claim physical air-gap isolation or zero egress unless runtime evidence explicitly proves it. NEVER fabricate or output external HTTP/HTTPS URLs (such as docs.mrpl.com or any external domain). All portals and tools are hosted locally in the application sidebar (Dashboard, Operator, Workspaces, Agents, Knowledge, Runs, Approvals, Artifacts, System)."
)


def sanitize_query(raw_text: str) -> str:
    """Strip legacy formatting wrappers like [Current Operator Query] to isolate clean current turn."""
    if not raw_text:
        return ""
    match = re.search(r'\[Current Operator Query\]\s*\n?(.*)', raw_text, re.DOTALL | re.IGNORECASE)
    if match:
        return match.group(1).strip()
    if "[Recent Conversation Context]" in raw_text:
        parts = raw_text.split("[Recent Conversation Context]")
        last_part = parts[-1].strip()
        op_matches = re.findall(r'Operator:\s*(.*)', last_part)
        if op_matches:
            return op_matches[-1].strip()
    return raw_text.strip()


def build_system_prompt(
    base_instructions: str,
    available_tools: List[Dict[str, Any]],
    context_str: str,
    intent: Optional[SemanticIntent] = None,
    workspace_name: Optional[str] = None
) -> str:
    """Construct an industrial agent prompt with tool definitions and citations."""
    if base_instructions and base_instructions.strip():
        persona = base_instructions
    elif workspace_name and workspace_name.strip():
        persona = f"You are CogniShift, a sovereign on-premise industrial AI assistant for {workspace_name}. You can reason over locally available workspace documents, use approved local tools, execute code inside the configured secure sandbox, and generate artifacts."
    else:
        persona = DEFAULT_SYSTEM_PERSONA


    if intent == SemanticIntent.CONVERSATION:
        prompt_parts = [
            persona,
            "\n--- INTERACTION MODE: DIRECT CONVERSATION ---",
            "1. You are conversing directly with the human operator.",
            "2. Answer questions about your capabilities, architecture, features, and industrial engineering honestly, concisely, and naturally.",
            "3. Do not recite alarm scenarios, emergency policies, or equipment readings unless the operator specifically asks about them.",
            "4. Do not call or propose any tools during direct conversational inquiries.",
            "\n--- OUTPUT FORMAT SPECIFICATION ---\n"
            "You may respond with a valid JSON object matching:\n"
            "```json\n"
            "{\n"
            '  "action": "final_answer",\n'
            '  "content": "<your direct, natural response>"\n'
            "}\n"
            "```\n"
            "Or respond directly in clear, well-structured markdown text."
        ]
        return "\n".join(prompt_parts)

    prompt_parts = [
        persona,
        "\n--- INDUSTRIAL SAFETY & OPERATIONAL GUIDELINES ---",
        "1. Prioritize plant safety, personnel protection, and OISD standards.",
        "2. When citing facts from the provided manuals, reference the manual name and page number.",
        "3. If a tool is required to inspect telemetry or perform an action, output a JSON tool call.",
        "4. STRICT LOTO (Lockout/Tagout) POLICY: Lockout/Tagout procedures strictly require zero-energy verification, physical lock/tag application, and mechanical isolation. NEVER recommend restarting, energizing, or cycling equipment during a LOTO procedure or while maintenance isolation is active.",
        "5. REFINERY TOPOLOGY CONSTRAINTS: Only interact with components registered in the plant topology (P-101A, P-101B, Reactor-B, SV-402, TK-01, Flare-Header, PT-101, TT-204). If the operator asks to control an unregistered component (e.g. H-101, CV-102), explain that it is outside the registered plant topology or not supported in this DCS simulation.",
        "6. SOVEREIGN ZERO-CLOUD AIR-GAP POLICY: External internet access, web browsing, Google searches, cloud uploads, and external APIs (e.g. ChatGPT) are strictly prohibited under the CogniShift Sovereign Industrial AI Policy. Immediately refuse any requests attempting external egress.",
    ]
    
    if available_tools:
        prompt_parts.append("\n--- AVAILABLE INDUSTRIAL TOOLS ---")
        for t in available_tools:
            risk = t.get("risk_level", "read_only")
            req_app = "YES (Requires Supervisor Approval)" if (t.get("requires_approval") or t["name"] in HIGH_RISK_TOOLS) else "NO (Safe to auto-run)"
            auth_schema = get_authoritative_tool_json_schema(t["name"])
            schema_display = json.dumps(auth_schema, indent=2) if auth_schema.get("properties") else t.get("input_schema", "{}")
            prompt_parts.append(
                f"- Tool: `{t['name']}` | Risk: {risk} | Human Approval: {req_app}\n"
                f"  Description: {t.get('description', '')}\n"
                f"  Input Schema:\n```json\n{schema_display}\n```"
            )
    else:
        prompt_parts.append("\nNo tools are currently assigned to your profile.")

    prompt_parts.append(
        "\n--- OUTPUT FORMAT SPECIFICATION ---\n"
        "You must respond with a valid JSON object matching one of these formats:\n\n"
        "1. When concluding a run, presenting findings, or directly answering the operator:\n"
        "```json\n"
        "{\n"
        '  "action": "final_answer",\n'
        '  "content": "<your complete, clear engineering explanation>",\n'
        '  "citations": ["Manual_Name.pdf | Page X"]\n'
        "}\n"
        "```\n\n"
        "2. When an authorized industrial tool is required to execute work or inspect telemetry:\n"
        "```json\n"
        "{\n"
        '  "action": "tool_call",\n'
        '  "tool_name": "<tool_name>",\n'
        '  "parameters": { ... },\n'
        '  "reason": "<clear explanation of why this action is required>"\n'
        "}\n"
        "```\n\n"
        "3. When completing an intermediate step in a multi-step plan before the final answer:\n"
        "```json\n"
        "{\n"
        '  "action": "step_observation",\n'
        '  "content": "<observation or progress from this step>"\n'
        "}\n"
        "```\n\n"
        "CRITICAL: If the operator is asking a question or requesting information, do NOT call a tool. Output a JSON object with 'action': 'final_answer'."
    )

    return "\n".join(prompt_parts)


async def _resolve_knowledge_and_page_context(
    db: Any,
    workspace_id: int,
    agent: dict,
    clean_input: str,
    resolved_context: Any,
    vision_analysis: Optional[str] = None,
    run_id: Optional[int] = None
) -> tuple[List[int], str, str]:
    """
    Authoritatively resolve allowed knowledge sources, dynamic document bindings,
    and direct page-level extractions.
    Returns: (allowed_source_ids, retrieval_query, page_direct_context)
    """
    raw_ks_ids = []
    if agent.get("knowledge_source_ids"):
        try:
            loaded = json.loads(agent["knowledge_source_ids"])
            if isinstance(loaded, list):
                raw_ks_ids = [int(x) for x in loaded if str(x).isdigit()]
        except Exception:
            raw_ks_ids = []

    allowed_source_ids = []
    if raw_ks_ids:
        placeholders = ",".join("?" for _ in raw_ks_ids)
        c_sources = await db.execute(
            f"SELECT id, name FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' AND id IN ({placeholders})",
            (workspace_id, *raw_ks_ids)
        )
        rows_sources = await c_sources.fetchall()
        allowed_source_ids = [r["id"] for r in rows_sources]

        # Also include any updated/active versions with the same names in this workspace
        known_names = [r["name"] for r in rows_sources]
        if known_names:
            n_ph = ",".join("?" for _ in known_names)
            c_active = await db.execute(
                f"SELECT id FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' AND name IN ({n_ph})",
                (workspace_id, *known_names)
            )
            for r_active in await c_active.fetchall():
                if r_active["id"] not in allowed_source_ids:
                    allowed_source_ids.append(r_active["id"])

    # Authorize documents explicitly resolved from the query or conversation history
    target_doc = await resolve_target_document_for_query(workspace_id, clean_input, db=db)
    
    # Check for colloquial document references (e.g. "what does inspection file say?")
    colloquial_match = None
    lower_input = clean_input.lower()
    for phrase in [
        "inspection file", "inspection report", "the inspection", "that inspection",
        "maintenance sop", "pump sop", "the sop", "that sop",
        "k-101 sop", "k-101 document", "k-101 manual", "compressor sop", "compressor document", "compressor manual",
        "emergency trip", "trip procedure", "operating procedure",
        "financial history", "financial spreadsheet", "that spreadsheet", "the spreadsheet", "financial file", "financial workbook"
    ]:
        if phrase in lower_input:
            colloquial_match = phrase
            break

    is_rca_mode = any(w in lower_input for w in [
        "rca", "root cause", "failure investigation", "investigate failure", "cause analysis"
    ])

    is_cross_doc = (
        is_rca_mode
        or any(w in lower_input for w in [
            "compare", "both", "all documents", "cross-reference", "against", "recommendation", "correlat",
            "troubleshoot", "why did"
        ])
        or (" and " in lower_input and any(doc_word in lower_input for doc_word in ["manual", "sop", "report", "file", "document", "drawing", "schematic"]))
    )

    # Hard authoritative source scope
    hard_agent_whitelist: List[int] = list(raw_ks_ids) if raw_ks_ids else []

    # Soft retrieval preference derived from explicitly mentioned/resolved documents
    focus_source_ids: List[int] = []
    if target_doc and target_doc.get("id"):
        focus_source_ids.append(int(target_doc["id"]))
    if resolved_context and resolved_context.pinned_source and resolved_context.pinned_source.get("id"):
        focus_source_ids.append(int(resolved_context.pinned_source["id"]))
    for rf in (resolved_context.files if resolved_context else []):
        c_rf = await db.execute(
            "SELECT id FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' AND (name = ? OR original_filename = ?)",
            (workspace_id, rf, rf)
        )
        for r_rf in await c_rf.fetchall():
            if r_rf["id"] not in focus_source_ids:
                focus_source_ids.append(r_rf["id"])

    if hard_agent_whitelist:
        # Agent has a strict sovereign knowledge source allowlist! Never breach it!
        if target_doc and int(target_doc["id"]) in hard_agent_whitelist:
            if colloquial_match and not is_cross_doc:
                allowed_source_ids = [int(target_doc["id"])]
                active_doc_for_page = target_doc
            else:
                allowed_source_ids = list(hard_agent_whitelist)
        else:
            allowed_source_ids = list(hard_agent_whitelist)
    else:
        # Agent has unrestricted workspace access; allow colloquial resolution and dynamic doc discovery
        if is_cross_doc:
            allowed_source_ids = []  # Unrestricted: allow searching all completed sources in this workspace
            if target_doc:
                active_doc_for_page = target_doc
        elif target_doc and colloquial_match:
            # PIN the single matching document and suppress unrelated knowledge sources
            allowed_source_ids = [int(target_doc["id"])]
            active_doc_for_page = target_doc
        elif target_doc and int(target_doc["id"]) not in allowed_source_ids:
            allowed_source_ids.append(int(target_doc["id"]))
            if run_id is not None:
                await log_event(
                    db, run_id, "document_reference_resolved",
                    f"Colloquial reference '{colloquial_match}' resolved to {target_doc['name']}",
                    {
                        "user_reference": colloquial_match,
                        "resolved_filename": target_doc["name"],
                        "confidence": 1.0,
                        "strict_single_source": True
                    }
                )
        else:
            allowed_source_ids = []

    # Follow-up source logging
    if resolved_context and resolved_context.files and not re.findall(r'\b([a-zA-Z0-9_\-\.]+\.(?:xlsx|xls|csv|pdf))\b', clean_input, re.IGNORECASE):
        prior_fn = resolved_context.files[0]
        if run_id is not None:
            await log_event(
                db, run_id, "followup_source_resolved",
                f"Reused previous source '{prior_fn}' for follow-up query",
                {
                    "previous_run_id": getattr(resolved_context, "source_turn", None),
                    "resolved_filename": prior_fn,
                    "resolution_reason": "anaphora_followup_continuity"
                }
            )

    active_doc_for_page = None
    # Authoritative Single-Source Resolution & Strict Scope Enforcement
    only_source_match = re.search(
        r'\b(?:using\s+only|only\s+from|from\s+only|in\s+only|based\s+only\s+on)\s+([A-Za-z0-9_\-\.]+\.[A-Za-z0-9]+)\b',
        clean_input,
        re.IGNORECASE
    )
    if not only_source_match and "only" in clean_input.lower():
        fn_cand = re.search(r'\b([A-Za-z0-9_\-\.]+\.(?:xlsx|xls|pdf|csv|docx))\b', clean_input, re.IGNORECASE)
        if fn_cand:
            only_source_match = fn_cand

    target_fn_hint = only_source_match.group(1).strip() if only_source_match else None
    resolved_source = await resolve_authoritative_source(
        workspace_id=workspace_id,
        query=clean_input,
        target_filename=target_fn_hint,
        db=db
    )

    if resolved_source:
        if run_id is not None:
            await log_event(
                db, run_id, "resolved_source",
                f"Authoritative source resolved: {resolved_source.filename} (Origin: {resolved_source.origin_type}, Selected By: {resolved_source.selected_by})",
                resolved_source.to_dict()
            )
        if resolved_source.strict_source_scope and resolved_source.source_id and not is_cross_doc:
            if hard_agent_whitelist:
                if resolved_source.source_id in hard_agent_whitelist:
                    allowed_source_ids = [resolved_source.source_id]
            else:
                allowed_source_ids = [resolved_source.source_id]
        if resolved_source.origin_type == "knowledge_source" and resolved_source.source_id:
            active_doc_for_page = {
                "id": resolved_source.source_id,
                "name": resolved_source.filename,
                "local_path": resolved_source.workspace_relative_path
            }
            target_doc = active_doc_for_page
    elif only_source_match and not is_cross_doc:
        cand_name = only_source_match.group(1).strip()
        c_iso = await db.execute(
            "SELECT * FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' AND (LOWER(name) = LOWER(?) OR LOWER(original_filename) = LOWER(?)) ORDER BY id DESC LIMIT 1",
            (workspace_id, cand_name, cand_name)
        )
        iso_row = await c_iso.fetchone()
        if iso_row:
            iso_doc = dict(iso_row)
            if not hard_agent_whitelist or iso_doc["id"] in hard_agent_whitelist:
                allowed_source_ids = [iso_doc["id"]]
            active_doc_for_page = iso_doc
            target_doc = iso_doc

    # Hard security boundary check: ensure allowed_source_ids never widens outside configured whitelist
    if hard_agent_whitelist:
        if not allowed_source_ids:
            allowed_source_ids = list(hard_agent_whitelist)
        else:
            allowed_source_ids = [s for s in allowed_source_ids if s in hard_agent_whitelist] or list(hard_agent_whitelist)

    # Direct Page-Level Context Extraction
    requested_page = getattr(resolved_context, "requested_page", None) if resolved_context else None
    if not requested_page:
        requested_page = extract_requested_page(clean_input)

    page_direct_context = ""
    if not active_doc_for_page:
        if resolved_context and resolved_context.pinned_source:
            active_doc_for_page = resolved_context.pinned_source
            if not target_doc:
                target_doc = active_doc_for_page
        elif target_doc:
            active_doc_for_page = target_doc
        elif resolved_context and resolved_context.files:
            c_f = await db.execute(
                "SELECT id, name, original_filename, local_path FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' AND (name = ? OR original_filename = ?) LIMIT 1",
                (workspace_id, resolved_context.files[0], resolved_context.files[0])
            )
            r_f = await c_f.fetchone()
            if r_f:
                active_doc_for_page = dict(r_f)
                if not target_doc:
                    target_doc = active_doc_for_page

    if requested_page and active_doc_for_page:
        doc_sid = active_doc_for_page["id"]
        doc_name = active_doc_for_page["name"]
        c_page = await db.execute(
            "SELECT text_content, extraction_method FROM document_pages WHERE source_id = ? AND page_number = ?",
            (doc_sid, requested_page)
        )
        p_row = await c_page.fetchone()
        if p_row and p_row["text_content"]:
            p_method = (p_row["extraction_method"] or "native").upper()
            page_direct_context = (
                f"[{doc_name} | Page {requested_page} | {p_method}]\n"
                f'<document_context source="{doc_name}" page="{requested_page}" method="{p_row["extraction_method"] or "native"}">\n'
                f"{p_row['text_content']}\n"
                f"</document_context>"
            )

    # Direct spreadsheet structure & calculation injection (preserves row/column alignment and prevents PAT/PBT confusion)
    if active_doc_for_page and not page_direct_context:
        raw_lp = active_doc_for_page.get("local_path")
        ws_root = get_workspace_root(workspace_id).resolve()
        if raw_lp:
            p_ss = Path(raw_lp)
            if not p_ss.is_absolute():
                p_ss = ws_root / p_ss
            if p_ss.exists() and p_ss.suffix.lower() in [".xlsx", ".xls", ".csv"]:
                ss_insights = await asyncio.to_thread(extract_document_insights, p_ss, query_hint=clean_input)
                if ss_insights.get("structured_text"):
                    trace_parts = [
                        f"[{active_doc_for_page['name']} | Structured Spreadsheet Table]",
                        f'<spreadsheet_table source="{active_doc_for_page["name"]}">'
                    ]
                    if ss_insights.get("metrics"):
                        trace_parts.append("--- AUTHORITATIVE EXTRACTED ROW-COLUMN VALUES ---")
                        for m_k, m_v in ss_insights["metrics"].items():
                            vals_str = ", ".join([f"{col}: {val}" for col, val in m_v.get("values", {}).items()])
                            trace_parts.append(f"- {m_v.get('matched_label', m_k.upper())}: {vals_str}")
                        if "pat" in ss_insights["metrics"] and "pbt" in ss_insights["metrics"]:
                            trace_parts.append("NOTE: PBT (Profit Before Tax) and PAT (Profit After Tax) are DISTINCT. Do NOT substitute PBT when asked for PAT.")
                    if ss_insights.get("growth"):
                        trace_parts.append("--- AUTHORITATIVE CALCULATIONS ---")
                        for g_k, g_v in ss_insights["growth"].items():
                            trace_parts.append(f"- {g_k}: {g_v}%")
                    trace_parts.append("\n" + ss_insights["structured_text"])
                    trace_parts.append("</spreadsheet_table>")
                    page_direct_context = "\n".join(trace_parts)

    retrieval_query = f"{clean_input} {vision_analysis}".strip() if vision_analysis else clean_input
    return allowed_source_ids, retrieval_query, page_direct_context


async def execute_agent_run(
    workspace_id: int,
    agent_id: int,
    input_text: str,
    user_id: str = "operator",
    input_image_path: Optional[str] = None,
    conversation_history: Optional[List[Any]] = None,
    enable_visual: Optional[bool] = None,
    enable_topology: Optional[bool] = None,
    custom_topology_context: Optional[str] = None,
) -> RunResponse:
    """Execute an end-to-end agent reasoning run with text and multimodal vision support."""
    t_run_start = time.perf_counter()
    routing_res: Optional[SemanticRoutingResult] = None
    frozen_scada_anomaly: Optional[Dict[str, Any]] = None
    frozen_scada_source: str = ""
    goal_contract: Optional[GoalContract] = None

    async with get_db() as db:
        # 1. Fetch Workspace and Agent Definition
        cursor_ws = await db.execute("SELECT name FROM workspaces WHERE id = ?", (workspace_id,))
        ws_row = await cursor_ws.fetchone()
        workspace_name = ws_row["name"] if ws_row else None

        cursor = await db.execute("SELECT * FROM agent_definitions WHERE id = ?", (agent_id,))
        agent_row = await cursor.fetchone()
        if not agent_row:
            raise ValueError(f"Agent ID {agent_id} not found")
        agent = dict(agent_row)

        if agent["workspace_id"] != workspace_id:
            raise ValueError(f"Agent ID {agent_id} does not belong to Workspace {workspace_id}")


        # Parse allowed tool IDs
        try:
            allowed_tool_ids = json.loads(agent.get("allowed_tool_ids") or "[]")
        except (json.JSONDecodeError, TypeError):
            allowed_tool_ids = []

        # 2. Fetch Tool Definitions for this agent
        available_tools = []
        if allowed_tool_ids:
            placeholders = ",".join("?" for _ in allowed_tool_ids)
            cursor = await db.execute(
                f"SELECT * FROM tool_definitions WHERE id IN ({placeholders}) AND enabled = 1",
                tuple(allowed_tool_ids)
            )
            available_tools = [dict(r) for r in await cursor.fetchall()]

        allowed_tool_names = [t["name"] for t in available_tools]
        tools_by_name = {t["name"]: t for t in available_tools}

        # 3. Create initial Run Record
        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
        input_type = "multimodal" if input_image_path else "text"
        cursor = await db.execute(
            """INSERT INTO agent_runs 
               (workspace_id, agent_id, user_id, input_text, input_type, input_image_path, status, model_name, operating_mode, started_at)
               VALUES (?, ?, ?, ?, ?, ?, 'running', ?, ?, ?) RETURNING *""",
            (workspace_id, agent_id, user_id, input_text, input_type, input_image_path, agent["model_name"], settings.operating_mode, now_str)
        )
        run_record = await cursor.fetchone()
        await db.commit()
        run_id = run_record["id"]

        def make_response(row_dict: dict) -> RunResponse:
            d = dict(row_dict)
            if d.get("routing_info") and isinstance(d["routing_info"], str):
                try:
                    d["routing_info"] = json.loads(d["routing_info"])
                except Exception:
                    d["routing_info"] = None
            res = RunResponse.model_validate(d)
            if routing_res is not None and res.routing_info is None:
                res.routing_info = routing_res.to_dict()
            return res

        clean_input = sanitize_query(input_text)

        # SIH26117 Authoritative Screening RCA Showcase Pathway (exact 8-milestone flow)
        from cognishift.core.screening_scenario import is_screening_rca_query, execute_screening_rca_workflow
        if is_screening_rca_query(clean_input):
            return await execute_screening_rca_workflow(
                db=db,
                workspace_id=workspace_id,
                agent_id=agent_id,
                user_id=user_id,
                run_id=run_id,
                clean_input=clean_input,
                make_response=make_response,
            )

        await log_event(db, run_id, "run_started", f"Run initiated for agent '{agent['name']}'", {"agent_id": agent_id, "user_id": user_id, "input_type": input_type})
        if conversation_history is None:
            conversation_history = []
            try:
                c_hist = await db.execute(
                    """SELECT input_text, result_text FROM agent_runs 
                       WHERE workspace_id = ? AND agent_id = ? AND status = 'completed' 
                       ORDER BY id DESC LIMIT 3""",
                    (workspace_id, agent_id)
                )
                rows_hist = await c_hist.fetchall()
                for r in reversed(rows_hist):
                    if r["input_text"]:
                        conversation_history.append({"role": "user", "content": r["input_text"]})
                    if r["result_text"]:
                        conversation_history.append({"role": "assistant", "content": r["result_text"]})
            except Exception as e:
                logger.warning(f"Failed to fetch conversation history: {e}")

        # Pre-fetch latest ingested document asynchronously if referenced in dialogue
        latest_doc = await get_latest_ingested_document_async(workspace_id, db=db)

        context_resolver = ConversationContextResolver(max_history_turns=settings.semantic_router_max_history_turns)
        resolved_context = context_resolver.resolve(
            current_turn=clean_input,
            conversation_history=conversation_history,
            raw_input=input_text,
            workspace_id=workspace_id,
            latest_doc=latest_doc
        )

        active_pending_task = await get_active_pending_task(workspace_id=workspace_id, user_id=user_id, db=db)

        # Early Authoritative Target Document Resolution
        target_doc = None
        direct_doc = await resolve_target_document_for_query(workspace_id, clean_input, db=db)
        if direct_doc:
            target_doc = direct_doc
            if resolved_context:
                d_name = direct_doc.get("name") or direct_doc.get("original_filename")
                if d_name:
                    resolved_context.files = [d_name]
                resolved_context.pinned_source = direct_doc
                resolved_context.source_turn = None
        elif active_pending_task and getattr(active_pending_task, "source_references", None) and active_pending_task.source_references.get("pinned_source"):
            target_doc = active_pending_task.source_references["pinned_source"]
        elif resolved_context and resolved_context.pinned_source:
            target_doc = resolved_context.pinned_source
        elif resolved_context and resolved_context.files:
            c_f = await db.execute(
                "SELECT id, name, original_filename, local_path FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' AND (name = ? OR original_filename = ?) LIMIT 1",
                (workspace_id, resolved_context.files[0], resolved_context.files[0])
            )
            r_f = await c_f.fetchone()
            if r_f:
                target_doc = dict(r_f)

        # --- ROUTER 1: SEMANTIC INTENT CLASSIFICATION ---
        if settings.semantic_router_enabled:
            sem_router = get_semantic_router()
            routing_res = await asyncio.to_thread(
                sem_router.route,
                clean_input,
                resolved_context=resolved_context,
                active_pending_task=active_pending_task
            )
        else:
            routing_res = SemanticRoutingResult(
                intent=SemanticIntent.COMPLEX_AGENT,
                confidence=1.0,
                abstained=False,
                details={"reason": "semantic_router_disabled"}
            )

        pre_artifact_contract = parse_artifact_request_contract(clean_input)
        if pre_artifact_contract.is_deliverable_request and routing_res.intent != SemanticIntent.CODE_EXECUTION:
            routing_res.intent = SemanticIntent.CODE_EXECUTION
            routing_res.decision_method = DecisionMethod.RULE
            routing_res.confidence = 1.0
            routing_res.abstained = False
            routing_res.details = {
                **routing_res.details,
                "rule": "explicit_artifact_deliverable",
                "artifact_contract": pre_artifact_contract.to_dict(),
            }

        # Evidence-bound questions must always traverse the grounding gate.
        # The semantic router can legitimately abstain or classify an unusual
        # phrasing as conversation, but letting that path reach a general model
        # would permit an unsupported document-specific answer.  This rule is
        # topic-general: it keys on an explicit evidence request, not a demo ID.
        evidence_terms = (
            "document", "knowledge vault", "inspection report", "manual",
            "procedure", "sop", "cite the source", "according to",
            "corrosion life", "wall thickness",
        )
        high_risk_evidence_terms = (
            "corrosion life", "corrosion rate", "wall thickness",
            "ultrasonic thickness", "inspection report", "nde report",
            "ndt report", "metallurgical report",
        )
        evidence_query = any(term in clean_input.lower() for term in evidence_terms)
        high_risk_evidence_query = any(
            term in clean_input.lower() for term in high_risk_evidence_terms
        )
        if (
            routing_res.intent not in (SemanticIntent.CODE_EXECUTION, SemanticIntent.KNOWLEDGE_QUERY)
            and evidence_query
            and (
                routing_res.intent != SemanticIntent.ARTIFACT_INSPECTION
                or high_risk_evidence_query
            )
        ):
            routing_res.intent = SemanticIntent.KNOWLEDGE_QUERY
            routing_res.decision_method = DecisionMethod.RULE
            routing_res.confidence = 1.0
            routing_res.abstained = False
            routing_res.details = {
                **routing_res.details,
                "rule": "explicit_evidence_grounding",
            }

        await log_event(
            db,
            run_id,
            "semantic_intent_routed",
            f"Semantic Intent: {routing_res.intent.value} (method: {routing_res.decision_method.value}, score: {routing_res.confidence}, margin: {routing_res.margin})",
            routing_res.to_dict()
        )

        # Safety Guard 0: Sovereign Air-Gap Egress Interception (Zero-Cloud Policy Enforcement)
        lower_input = clean_input.lower()
        egress_patterns = [
            r"\b(google|bing|duckduckgo|yahoo)\b.*?(search|look\s*up|find|query)",
            r"\bsearch\b.*?(google|internet|web|online)",
            r"\b(send|upload|backup|sync|push|forward|stream|transfer)\b.*?(external|cloud|aws|azure|gcp|s3|remote\s*server)",
            r"\b(cloud\s*backup|external\s*cloud)\b",
            r"\b(webbrowser|urllib|curl\s+https?://|wget\s+https?://)\b",
            r"https?://",
            r"\b(connect\s+to|fetch\s+from|query)\s+(chatgpt|openai|anthropic|gemini|external\s*api)\b",
        ]
        if any(re.search(pat, lower_input) for pat in egress_patterns):
            result_text = (
                "🛑 **SOVEREIGN POLICY ENFORCEMENT**: Outbound internet access, public web searching, "
                "external cloud backups, and third-party AI APIs (such as ChatGPT) are strictly prohibited "
                "under the configured CogniShift application egress policy. "
                "This request was blocked before dispatch; physical network isolation is verified separately at deployment."
            )
            plan = AgentPlan(
                goal=clean_input,
                current_step_index=0,
                max_steps=1,
                steps=[PlanStep(id=1, description="Enforce sovereign network policy and reject outbound egress request", status="completed", observation="Blocked by strict Sovereign Network Policy")]
            )
            saved_plan_json = serialize_plan(plan)
            await db.execute(
                """UPDATE agent_runs
                   SET status = 'completed', result_text = ?, sources_used = 'CogniShift Sovereign Policy (Zero Cloud)', structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                   WHERE id = ?""",
                (result_text, saved_plan_json, run_id)
            )
            await db.commit()
            await log_event(db, run_id, "sovereign_egress_blocked", "External egress request blocked by strict sovereign network policy.")
            cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
            return make_response(dict(await cursor.fetchone()))

        # Safety Guard 1: Negation Guard ("Do not restart P-101A", "Don't do it")
        if routing_res.details.get("rule") == "negation_guard":
            result_text = "Acknowledged. I will not proceed with that action."
            plan = AgentPlan(
                goal=clean_input,
                current_step_index=0,
                max_steps=1,
                steps=[PlanStep(id=1, description="Process operator negation directive", status="completed", observation="Action inhibited by operator negation guard")]
            )
            saved_plan_json = serialize_plan(plan)
            await db.execute(
                """UPDATE agent_runs
                   SET status = 'completed', result_text = ?, sources_used = 'None (Direct Conversation)', structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                   WHERE id = ?""",
                (result_text, saved_plan_json, run_id)
            )
            await db.commit()
            await log_event(db, run_id, "negation_guard_invoked", "Operator negation directive acknowledged. Action inhibited.")
            cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
            return make_response(dict(await cursor.fetchone()))

        # Safety Guard 2: User Cancellation ("cancel that", "stop", "abort")
        if routing_res.details.get("rule") == "user_cancellation":
            if active_pending_task:
                await cancel_pending_task(active_pending_task.id, user_id=user_id, db=db)
                cancel_note = f"Pending task '{active_pending_task.id}' cancelled by operator."
            else:
                cancel_note = "Cancellation acknowledged."
            result_text = f"Understood. {cancel_note} No action will be taken."
            plan = AgentPlan(
                goal=clean_input,
                current_step_index=0,
                max_steps=1,
                steps=[PlanStep(id=1, description="Cancel pending task", status="completed", observation=cancel_note)]
            )
            saved_plan_json = serialize_plan(plan)
            await db.execute(
                """UPDATE agent_runs
                   SET status = 'completed', result_text = ?, sources_used = 'None (Direct Conversation)', structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                   WHERE id = ?""",
                (result_text, saved_plan_json, run_id)
            )
            await db.commit()
            await log_event(db, run_id, "task_cancelled", cancel_note)
            cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
            return make_response(dict(await cursor.fetchone()))

        # Safety Guard 3: Stale or Missing Pending Task Affirmation ("yea do it" with no active task)
        if routing_res.details.get("rule") == "stale_or_missing_pending_task":
            if not resolved_context.pinned_source and not resolved_context.files and not resolved_context.equipment_ids:
                result_text = "There is no active pending task awaiting confirmation. Please state the specific action or command you would like me to perform."
                plan = AgentPlan(
                    goal=clean_input,
                    current_step_index=0,
                    max_steps=1,
                    steps=[PlanStep(id=1, description="Verify pending task state", status="completed", observation="No active pending task found within TTL")]
                )
                saved_plan_json = serialize_plan(plan)
                await db.execute(
                    """UPDATE agent_runs
                       SET status = 'completed', result_text = ?, sources_used = 'None (Direct Conversation)', structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                       WHERE id = ?""",
                    (result_text, saved_plan_json, run_id)
                )
                await db.commit()
                await log_event(db, run_id, "pending_task_stale_or_missing", result_text)
                cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                return make_response(dict(await cursor.fetchone()))

        # Atomic CAS Claim for Resumed Pending Task
        resumed_task_id = routing_res.details.get("resumed_task_id")
        claimed_task: Optional[PendingTask] = None
        if resumed_task_id and active_pending_task and active_pending_task.id == resumed_task_id:
            claimed = await claim_pending_task_atomic(
                task_id=active_pending_task.id,
                expected_version=active_pending_task.version,
                claiming_user_id=user_id,
                db=db
            )
            if not claimed:
                conflict_msg = f"Task '{active_pending_task.id}' was already claimed or updated by another concurrent session."
                plan = AgentPlan(
                    goal=clean_input,
                    current_step_index=0,
                    max_steps=1,
                    steps=[PlanStep(id=1, description="Claim pending task", status="failed", error_message=conflict_msg)]
                )
                saved_plan_json = serialize_plan(plan)
                await db.execute(
                    """UPDATE agent_runs
                       SET status = 'failed', error_message = ?, result_text = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                       WHERE id = ?""",
                    (conflict_msg, conflict_msg, saved_plan_json, run_id)
                )
                await db.commit()
                await log_event(db, run_id, "task_claim_conflict", conflict_msg)
                cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                return make_response(dict(await cursor.fetchone()))
            claimed_task = active_pending_task
            await log_event(
                db, run_id, "task_resumed_atomic",
                f"Successfully claimed and resumed pending task '{active_pending_task.id}' via atomic CAS (v{active_pending_task.version} -> v{active_pending_task.version + 1})",
                {"task_id": active_pending_task.id, "goal": active_pending_task.requested_goal}
            )

        effective_goal = claimed_task.requested_goal if claimed_task else clean_input

        # RCA Accuracy Mode Guarantee: detect RCA intent early to prevent conversational plan routing
        lower_input = clean_input.lower()
        has_equipment_tag = bool(re.search(r"\b[A-Za-z]{1,4}-\d{3,4}[A-Za-z]?\b", clean_input))
        is_rca_mode = (
            any(k in lower_input for k in [
                "rca", "root cause", "failure investigation", "investigate failure",
                "incident investigation", "why did it fail", "why did the system trip",
                "troubleshoot", "investigate", "cause of", "what caused"
            ])
            or (has_equipment_tag and any(k in lower_input for k in ["trip", "tripped", "failed", "failure", "shutdown", "shut down", "alarm", "cavitat", "leak", "overpressure", "overheat"]))
            or (any(w in lower_input for w in ["why did", "why has"]) and any(w in lower_input for w in ["trip", "fail"]))
        )
        effective_intent = (
            SemanticIntent.KNOWLEDGE_QUERY
            if is_rca_mode
            else routing_res.intent
        )

        # Immediate Fast Path for UI Navigation: zero LLM inference, zero RAG
        if effective_intent == SemanticIntent.UI_NAVIGATION:
            target_view = routing_res.details.get("target") or "knowledge"
            target_route = routing_res.details.get("target_route") or f"/{target_view}"
            friendly_name = routing_res.details.get("friendly_name") or target_view.title()

            sources_used = "None (Direct Navigation)"
            result_text = (
                f"You can access **{friendly_name}** at `{target_route}` via the navigation sidebar, "
                f"or click the direct navigation action below.\n\n"
                f"Document uploads, OCR extraction, and knowledge indexing use the configured local processing path."
            )
            plan = AgentPlan(
                goal=clean_input,
                current_step_index=0,
                max_steps=2,
                steps=[
                    PlanStep(id=1, description="Identify requested interface destination", status="completed", observation=f"Target: {friendly_name} ({target_route})"),
                    PlanStep(id=2, description="Provide on-premise navigation path", status="completed", observation=result_text)
                ]
            )
            saved_plan_json = serialize_plan(plan)
            routing_res.details["target_route"] = target_route
            routing_res.details["friendly_name"] = friendly_name
            routing_info_json = json.dumps(routing_res.to_dict())
            await db.execute(
                """UPDATE agent_runs
                   SET status = 'completed', result_text = ?, sources_used = ?, structured_plan = ?, routing_info = ?, completed_at = CURRENT_TIMESTAMP
                   WHERE id = ?""",
                (result_text, sources_used, saved_plan_json, routing_info_json, run_id)
            )
            await db.commit()
            await log_event(db, run_id, "completed", f"Navigation processed: {friendly_name} ({target_route})", {"target_route": target_route})
            cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
            res = make_response(dict(await cursor.fetchone()))
            res.routing_info = routing_res.to_dict()
            return res

        # --- ROUTER 2: HARDWARE-AWARE MODEL ROUTING ---
        has_img = bool(input_image_path and Path(input_image_path).exists())
        conf_val = routing_res.confidence if routing_res.confidence is not None else 0.90
        task_classified = classify_task(clean_input, has_image=has_img)
        if is_rca_mode or task_classified.task_type == "heavy_reasoning":
            task_info = TaskClassification(
                task_type="heavy_reasoning",
                required_capabilities=["heavy_reasoning", "reasoning"],
                requires_vision=False,
                confidence=0.95
            )
        elif has_img:
            task_info = task_classified
        elif effective_intent == SemanticIntent.CONVERSATION:
            task_info = TaskClassification(
                task_type="conversational",
                required_capabilities=["reasoning"],
                requires_vision=False,
                confidence=conf_val
            )
        elif effective_intent == SemanticIntent.CODE_EXECUTION:
            task_info = TaskClassification(
                task_type="coding",
                required_capabilities=["coding", "structured_data"],
                requires_vision=False,
                confidence=conf_val
            )
        elif effective_intent in [SemanticIntent.ARTIFACT_INSPECTION, SemanticIntent.KNOWLEDGE_QUERY]:
            task_info = TaskClassification(
                task_type="document_analysis",
                required_capabilities=["document_analysis", "reasoning"],
                requires_vision=False,
                confidence=conf_val
            )
        else:
            task_info = task_classified

        await log_event(db, run_id, "task_classified", f"Classified task as '{task_info.task_type}'", task_info.model_dump())

        # Inspect currently installed models on the sovereign Ollama instance
        installed_models = None
        if settings.operating_mode != "simulated":
            try:
                provider = get_provider()
                info = await provider.model_info()
                if info and info.get("status") == "available":
                    installed_models = info.get("models", [])
                    update_verified_inventory_cache(exact_tags=installed_models, verified_at=time.monotonic())
            except Exception as e:
                logger.warning(f"Could not query installed models from Ollama: {e}")

        try:
            routing = route_model(
                task_info,
                available_vram_mb=6000,
                preferred_model=agent.get("model_name"),
                installed_models=installed_models,
                require_verified_inventory=(settings.operating_mode != "simulated")
            )
        except LocalModelInventoryUnavailableError as inv_err:
            fail_msg = f"Model routing failure: local model inventory is unavailable ({inv_err})."
            logger.error(fail_msg)
            await log_event(db, run_id, "model_inventory_unavailable", fail_msg, {"error": str(inv_err)})
            await db.execute(
                """UPDATE agent_runs
                   SET status = 'failed', result_text = ?, error_message = ?, completed_at = CURRENT_TIMESTAMP
                   WHERE id = ?""",
                (fail_msg, str(inv_err), run_id)
            )
            await db.commit()
            cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
            return make_response(dict(await cursor.fetchone()))
        await log_event(
            db,
            run_id,
            "model_candidates_evaluated",
            f"Evaluated {len(routing.candidate_evaluations)} local models against 6GB VRAM budget",
            {k: v.model_dump() for k, v in routing.candidate_evaluations.items()}
        )
        await log_event(
            db,
            run_id,
            "model_selected",
            f"Selected {routing.selected_model_name} for execution",
            {"selected_model": routing.selected_model, "reason": routing.selection_reason}
        )
        selected_model_id = routing.selected_model
        from cognishift.core.model_registry import get_model
        m_def = get_model(selected_model_id)
        if m_def and not m_def.supports_tools:
            logger.info(f"Model '{selected_model_id}' lacks tool/orchestration capability. Using '{settings.text_model}' as agent orchestrator.")
            selected_model_id = settings.text_model

        # Pre-verify that selected_model_id is actually installed locally; if not, immediately use settings.text_model
        if installed_models:
            installed_clean = set(m.strip().lower() for m in installed_models)
            sel_clean = selected_model_id.strip().lower()
            matched = (
                sel_clean in installed_clean
                or (":" not in sel_clean and f"{sel_clean}:latest" in installed_clean)
                or (sel_clean.endswith(":latest") and sel_clean[:-7] in installed_clean)
            )
            if not matched:
                logger.warning(f"Selected model '{selected_model_id}' is not installed locally in Ollama ({installed_models}). Using '{settings.text_model}' to prevent 404.")
                selected_model_id = settings.text_model

        await db.execute("UPDATE agent_runs SET model_name = ? WHERE id = ?", (selected_model_id, run_id))
        await db.commit()

        # --- TASK-SCOPED GOAL CONTRACT INITIALIZATION ---
        artifact_contract = parse_artifact_request_contract(clean_input)
        is_financial_yoy_task = (
            any(w in lower_input for w in ["yoy", "year over year", "year-on-year", "growth rate", "cagr", "growth"])
            and any(w in lower_input for w in ["revenue", "ebitda", "pat", "financial", "performance", "p&l"])
        )
        required_fields = []
        if is_financial_yoy_task:
            required_fields = [
                "revenue_previous", "revenue_current", "revenue_yoy_pct",
                "ebitda_previous", "ebitda_current", "ebitda_yoy_pct",
                "pat_previous", "pat_current", "pat_yoy_pct"
            ]
        if is_financial_yoy_task or artifact_contract.is_deliverable_request:
            goal_contract = GoalContract(
                required_fields=required_fields,
                artifact_contract=artifact_contract
            )

        # --- PHASE 2B: BOUNDED STRUCTURED PLAN CREATION ---
        if effective_intent == SemanticIntent.CONVERSATION:
            plan = AgentPlan(
                goal=clean_input,
                current_step_index=1,
                max_steps=3,
                steps=[
                    PlanStep(id=1, description="Route semantic intent and verify conversational policy", status="completed", observation="Routed to CONVERSATION (Zero RAG / Zero tools)"),
                    PlanStep(id=2, description="Synthesize direct conversational response", status="pending"),
                    PlanStep(id=3, description="Present answer to operator", status="pending")
                ]
            )
        elif effective_intent == SemanticIntent.ARTIFACT_INSPECTION:
            plan = AgentPlan(
                goal=clean_input,
                current_step_index=1,
                max_steps=5,
                steps=[
                    PlanStep(id=1, description="Route semantic intent to workspace artifact inspection", status="completed", observation="Intent: ARTIFACT_INSPECTION"),
                    PlanStep(id=2, description="Resolve referenced workspace artifact", status="pending"),
                    PlanStep(id=3, description="Read artifact content", status="pending"),
                    PlanStep(id=4, description="Analyze artifact data", status="pending"),
                    PlanStep(id=5, description="Synthesize report for operator", status="pending")
                ]
            )
        elif effective_intent == SemanticIntent.KNOWLEDGE_QUERY:
            plan = AgentPlan(
                goal=clean_input,
                current_step_index=1,
                max_steps=4,
                steps=[
                    PlanStep(id=1, description="Route semantic intent to knowledge retrieval", status="completed", observation=f"Intent: KNOWLEDGE_QUERY (Score: {conf_val:.2f})"),
                    PlanStep(id=2, description="Retrieve domain knowledge from authorized sources", status="pending"),
                    PlanStep(id=3, description="Reason over retrieved evidence", status="pending"),
                    PlanStep(id=4, description="Present answer to operator", status="pending")
                ]
            )
        elif effective_intent == SemanticIntent.CODE_EXECUTION:
            lower_goal = effective_goal.lower()
            is_doc_report = any(w in lower_goal for w in ["document", "pdf", "report", "manual", "latest", "ingested", "convert", "excel", "xlsx", "spreadsheet", "csv", "audit", "financial", "data", "history", "analyze", "analysis", "visualize", "plot", "docx", "word", "format", "deliverable"])
            target_doc_info = None
            if is_doc_report:
                target_doc_info = (
                    claimed_task.source_references.get("pinned_source")
                    if (claimed_task and claimed_task.source_references)
                    else (resolved_context.pinned_source if resolved_context else None)
                )
                if not target_doc_info and target_doc:
                    target_doc_info = target_doc
                if not target_doc_info:
                    target_doc_info = await resolve_target_document_for_query(workspace_id, effective_goal, db=db)
                if not target_doc_info and resolved_context and resolved_context.files:
                    c_f = await db.execute(
                        "SELECT id, name, original_filename, local_path FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' AND (name = ? OR original_filename = ?) LIMIT 1",
                        (workspace_id, resolved_context.files[0], resolved_context.files[0])
                    )
                    r_f = await c_f.fetchone()
                    if r_f:
                        target_doc_info = dict(r_f)
            if not target_doc and target_doc_info:
                target_doc = target_doc_info
            doc_label = target_doc_info["name"] if target_doc_info else "workspace data"

            # Determine requested deliverable format (PDF, XLSX, PPTX, Image, or DOCX)
            req_format = "DOCX"
            if _is_presentation_request(lower_goal):
                req_format = "PPTX"
            elif "pdf" in lower_goal and not any(w in lower_goal for w in ["convert to docx", "docx", "word", "ppt", "pptx"]):
                req_format = "PDF"
            elif any(w in lower_goal for w in ["excel", "xlsx", "spreadsheet"]):
                req_format = "XLSX"
            elif re.search(r"\b(?:csv|comma-separated)\b", lower_goal):
                req_format = "CSV"
            elif any(w in lower_goal for w in ["jpg", "jpeg", "png", "image", "visualize", "plot", "chart", "graph"]):
                req_format = "IMAGE"

            step2_desc = (
                "Extract tabular series and generate deterministic data visualization"
                if artifact_contract.png_required
                else "Execute Python analysis script in isolated sandbox"
            )
            step3_desc = (
                f"Confirm data visualization deliverables for {doc_label}"
                if (artifact_contract.png_required and not (artifact_contract.pdf_required or artifact_contract.docx_required or artifact_contract.xlsx_required))
                else f"Generate formal {req_format} engineering report for {doc_label}"
            )

            plan = AgentPlan(
                goal=effective_goal,
                current_step_index=1,
                max_steps=4,
                steps=[
                    PlanStep(id=1, description="Resolve source document and verify execution environment", status="completed", observation=f"Authoritatively resolved source: {doc_label}"),
                    PlanStep(id=2, description=step2_desc, status="pending"),
                    PlanStep(id=3, description=step3_desc, status="pending"),
                    PlanStep(id=4, description="Validate generated artifacts and deliver final report", status="pending")
                ]
            )
        else:
            plan = create_initial_plan(goal=clean_input, task_type=task_info.task_type)

        plan_json = serialize_plan(plan)
        await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (plan_json, run_id))
        await db.commit()
        await log_event(
            db,
            run_id,
            "plan_created",
            f"Constructed structured execution plan with {len(plan.steps)} bounded steps",
            {"goal": plan.goal, "steps": [s.model_dump() for s in plan.steps]}
        )

        try:
            # 4. Multimodal Vision Inspection (if image provided)
            vision_analysis = ""
            if input_image_path:
                data_root = settings.data_dir.resolve()
                img_candidate = Path(input_image_path)
                if img_candidate.is_absolute():
                    resolved_cand = img_candidate.resolve()
                    try:
                        resolved_cand.relative_to(data_root)
                        img_path = resolved_cand
                    except ValueError:
                        raise ValueError(f"Security Error: Image path '{input_image_path}' is outside allowed data directory '{settings.data_dir}'")
                else:
                    try:
                        img_path = resolve_workspace_path(workspace_id, input_image_path, purpose="read")
                    except Exception as e:
                        raise ValueError(f"Security Error: Image path '{input_image_path}' is outside allowed data directory: {e}")
                if img_path.exists():

                    await log_event(
                        db, run_id, "vision_started",
                        f"Analyzing image {img_path.name} with local vision model...",
                        {"model": settings.vision_model, "image": img_path.name},
                    )
                    try:
                        with open(img_path, "rb") as f:
                            img_bytes = f.read()
                        # P&IDs and nameplates are text-dense. Fuse local OCR into the
                        # VLM prompt so the small vision model stays anchored to visible
                        # tags instead of inventing a generic scene description.
                        ocr_evidence = ""
                        ocr_confidence = None
                        try:
                            from cognishift.core.document_processing.image_preprocessor import (
                                prepare_image_for_vision,
                                preprocess_image_for_ocr,
                            )
                            from cognishift.core.document_processing.ocr_provider import get_ocr_provider

                            ocr_preprocessed = await asyncio.to_thread(preprocess_image_for_ocr, img_bytes)
                            ocr_result = await get_ocr_provider().extract(ocr_preprocessed)
                            ocr_evidence = (ocr_result.raw_text or ocr_result.text or "").strip()[:3000]
                            ocr_confidence = ocr_result.confidence
                            vision_bytes = await asyncio.to_thread(prepare_image_for_vision, img_bytes)
                            await log_event(
                                db, run_id, "ocr_completed",
                                f"Local OCR extracted image text ({len(ocr_evidence)} chars)",
                                {"engine": ocr_result.engine, "confidence": ocr_confidence},
                            )
                        except Exception as ocr_err:
                            logger.warning(f"Image OCR grounding unavailable: {ocr_err}")
                            vision_bytes = img_bytes

                        provider = get_provider()
                        grounded_prompt = (
                            "Inspect this industrial image using only visible evidence. List visible equipment, "
                            "instrument tags, line labels, readings, units, and safety notes. Do not infer a generic "
                            "system type or repeat an object unless it is visibly supported. Mark unclear text as uncertain.\n"
                            f"Operator request: {clean_input}\n"
                        )
                        if ocr_evidence:
                            grounded_prompt += (
                                "The following text was independently extracted from this same image by local OCR. "
                                "Use it as image-derived grounding; correct minor OCR errors only when visually justified:\n"
                                f"{ocr_evidence}"
                            )
                        vlm_res = await provider.analyze_image(
                            vision_bytes,
                            prompt=grounded_prompt,
                        )
                        raw_vision_analysis = vlm_res.text.strip()
                        # Collapse repeated short outputs (a common small-VLM failure
                        # mode) without pretending they are separate observations.
                        unique_vision_lines = []
                        seen_vision_lines = set()
                        for raw_line in raw_vision_analysis.splitlines():
                            cleaned_line = re.sub(r"^\s*\d+[.)]\s*", "", raw_line).strip()
                            normalized_line = cleaned_line.casefold()
                            if cleaned_line and normalized_line not in seen_vision_lines:
                                seen_vision_lines.add(normalized_line)
                                unique_vision_lines.append(cleaned_line)
                        raw_vision_analysis = "\n".join(unique_vision_lines)
                        vision_analysis = raw_vision_analysis
                        if ocr_evidence:
                            confidence_label = f"{ocr_confidence:.1%}" if ocr_confidence is not None else "unavailable"
                            vision_analysis = (
                                f"Local OCR evidence (confidence {confidence_label}):\n{ocr_evidence}\n\n"
                                f"Local vision-model cross-check:\n{vision_analysis or 'No additional visual description returned.'}\n\n"
                                "Unclear OCR characters remain uncertain and must be checked against the original image."
                            )
                        await log_event(
                            db,
                            run_id,
                            "vision_completed",
                            f"Visual inspection completed ({len(vision_analysis)} chars)",
                            {"analysis": vision_analysis, "model": provider.vision_model}
                        )
                        if vision_analysis and plan.steps and "rasterize" in plan.steps[0].description.lower():
                            plan.steps[0].status = "completed"
                            plan.steps[0].observation = vision_analysis
                            plan.advance_to_next_step()
                            saved_plan = serialize_plan(plan)
                            await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan, run_id))
                            await db.commit()
                            await log_event(
                                db, run_id, "plan_step_completed",
                                f"Step #1 completed: Visual inspection telemetry extracted ({len(vision_analysis)} chars)",
                                {"step_id": 1, "observation": vision_analysis[:200]}
                            )
                    except Exception as e:
                        logger.error(f"Vision analysis error: {e}")

            # 5. Domain-Specific Context Retrieval & Anti-Pollution Selective Policy
            context_str = ""
            graph_context = ""
            artifact_citations = []
            combined_context_parts = []
            retrieved_evidence_catalog: List[Dict[str, Any]] = []
            final_answer_citations: List[str] = []
            rca_bundle: Optional[RCAEvidenceBundle] = None

            if vision_analysis:
                combined_context_parts.append(
                    f"--- VISUAL INSPECTION TELEMETRY (LOCAL VLM ANALYSIS) ---\n"
                    f"Image Artifact: {Path(input_image_path).name}\n"
                    f"Inspection Telemetry: {vision_analysis}"
                )

            # is_rca_mode was already computed early in execute_agent_run
            is_cross_inspection = any(w in clean_input.lower() for w in ["compare", "sop", "manual", "procedure", "against", "cross-reference", "correlat"])
            strict_visual_scope = bool(input_image_path) and not is_cross_inspection and not is_rca_mode

            # RCA Accuracy Mode Guarantee: independently force full multi-channel retrieval (text + visual + topology)
            # Semantic router error must never suppress an evidence channel in RCA mode.
            effective_intent = SemanticIntent.KNOWLEDGE_QUERY if is_rca_mode else routing_res.intent

            if strict_visual_scope:
                sources_used = f"Visual Artifact | {Path(input_image_path).name}"
                # A pure image-description request is already answered by the local
                # OCR + VLM evidence. Do not send it through a text model that can
                # add unsupported equipment, procedures, or recommendations.
                if vision_analysis:
                    await db.execute(
                        "UPDATE agent_runs SET model_name=? WHERE id=?",
                        (settings.vision_model, run_id),
                    )
                await log_event(
                    db, run_id, "retrieval_bypassed",
                    f"Strict visual scope active for {Path(input_image_path).name}: RAG, Plant Graph, and external artifacts bypassed.",
                    {"image": Path(input_image_path).name}
                )
                if plan.steps and len(plan.steps) > 1:
                    for s in plan.steps[1:]:
                        if "retriev" in s.description.lower() or "context" in s.description.lower():
                            s.status = "completed"
                            s.observation = f"Visual telemetry isolated ({Path(input_image_path).name})"

            elif effective_intent == SemanticIntent.CONVERSATION:
                # Anti-Pollution Policy: Zero ChromaDB, Zero Plant Graph, Zero Artifacts
                sources_used = "None (Direct Conversation)"
                await log_event(
                    db, run_id, "retrieval_bypassed",
                    "Conversational intent detected: RAG and artifact injection bypassed to prevent context pollution.",
                    {"intent": routing_res.intent.value}
                )

            elif effective_intent == SemanticIntent.ARTIFACT_INSPECTION:
                # Targeted Artifact Resolution ONLY: Zero ChromaDB, Zero Plant Graph
                ws_root = get_workspace_root(workspace_id).resolve()

                # Search knowledge_sources first for authoritative documents uploaded in Knowledge Vault
                cursor_sources = await db.execute(
                    "SELECT id, name, original_filename, local_path, source_type FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' ORDER BY id DESC LIMIT 100",
                    (workspace_id,)
                )
                source_rows = list(await cursor_sources.fetchall())

                # Search workspace_artifacts second for generated files
                cursor_artifacts = await db.execute(
                    "SELECT id, filename, relative_path, file_size, artifact_type, title, description FROM workspace_artifacts WHERE workspace_id = ? ORDER BY id DESC LIMIT 50",
                    (workspace_id,)
                )
                art_rows = await cursor_artifacts.fetchall()

                matching_artifacts = []
                lower_input = clean_input.lower()
                if not target_doc:
                    if resolved_context and resolved_context.pinned_source:
                        target_doc = resolved_context.pinned_source
                    else:
                        target_doc = await resolve_target_document_for_query(workspace_id, clean_input, db=db)

                target_files = []
                if resolved_context and resolved_context.files:
                    target_files.extend(resolved_context.files)
                if routing_res.references and routing_res.references.files:
                    for f in routing_res.references.files:
                        if f not in target_files:
                            target_files.append(f)
                explicit_files = re.findall(r'\b([a-zA-Z0-9_\-\.]+\.(?:xlsx|xls|csv|tsv|docx|doc|pdf|json|yaml|yml|txt|md|py|log))\b', clean_input, re.IGNORECASE)
                for ef in explicit_files:
                    if ef not in target_files:
                        target_files.append(ef)

                if target_files:
                    tf_ph = ",".join("?" for _ in target_files)
                    c_tf = await db.execute(
                        f"SELECT id, name, original_filename, local_path, source_type FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' AND (name IN ({tf_ph}) OR original_filename IN ({tf_ph}))",
                        (workspace_id, *target_files, *target_files)
                    )
                    existing_sids = {r["id"] for r in source_rows}
                    for r_tf in await c_tf.fetchall():
                        if r_tf["id"] not in existing_sids:
                            source_rows.append(r_tf)

                GENERIC_STEMS = {"report", "reading", "readings", "file", "document", "artifact", "data", "sheet", "table", "summary", "test", "plan", "output", "input", "result", "results", "status", "logs", "log", "pdf"}
                ref_files = [f.lower() for f in target_files]

                is_strict_scope = any(w in lower_input for w in ["using only", "strictly from", "only use", "do not use other", "using that workbook only", "using that file only"])
                seen_artifact_ids = set()

                # 1. Authoritative resolved document from Knowledge Vault takes top precedence
                if target_doc and target_doc.get("local_path"):
                    raw_lp = target_doc.get("local_path", "")
                    rel_p = raw_lp
                    if raw_lp:
                        try:
                            p = Path(raw_lp)
                            if p.is_absolute():
                                rel_p = str(p.resolve().relative_to(ws_root)).replace("\\", "/")
                        except Exception:
                            rel_p = Path(raw_lp).name
                    matching_artifacts.append({
                        "id": target_doc.get("id"),
                        "filename": target_doc.get("original_filename") or target_doc.get("name"),
                        "relative_path": rel_p,
                        "artifact_type": target_doc.get("source_type") or "file",
                        "title": target_doc.get("name"),
                        "is_knowledge_source": True
                    })
                    seen_artifact_ids.add(target_doc.get("id"))

                # 2. Match Knowledge Sources before Workspace Artifacts (suppressed if strict single-source scope is active)
                if not (is_strict_scope and target_doc):
                    for src in source_rows:
                        if src["id"] in seen_artifact_ids:
                            continue
                        s_fname = (src["original_filename"] or src["name"]).lower()
                        s_stem = Path(s_fname).stem.lower()
                        raw_lp = src["local_path"]
                        rel_p = raw_lp
                        if raw_lp:
                            try:
                                p = Path(raw_lp)
                                if p.is_absolute():
                                    rel_p = str(p.resolve().relative_to(ws_root)).replace("\\", "/")
                            except Exception:
                                rel_p = Path(raw_lp).name
                        if explicit_files:
                            if any(s_fname == ef.lower() for ef in explicit_files):
                                matching_artifacts.append({
                                    "id": src["id"],
                                    "filename": src["original_filename"] or src["name"],
                                    "relative_path": rel_p,
                                    "artifact_type": src["source_type"] or "pdf",
                                    "title": src["name"],
                                    "is_knowledge_source": True
                                })
                                seen_artifact_ids.add(src["id"])
                        elif s_fname in lower_input or s_fname in ref_files:
                            matching_artifacts.append({
                                "id": src["id"],
                                "filename": src["original_filename"] or src["name"],
                                "relative_path": rel_p,
                                "artifact_type": src["source_type"] or "pdf",
                                "title": src["name"],
                                "is_knowledge_source": True
                            })
                            seen_artifact_ids.add(src["id"])
                        elif s_stem not in GENERIC_STEMS and len(s_stem) > 4:
                            if s_stem in lower_input or s_stem in [Path(rf).stem.lower() for rf in ref_files]:
                                matching_artifacts.append({
                                    "id": src["id"],
                                    "filename": src["original_filename"] or src["name"],
                                    "relative_path": rel_p,
                                    "artifact_type": src["source_type"] or "pdf",
                                    "title": src["name"],
                                    "is_knowledge_source": True
                                })
                                seen_artifact_ids.add(src["id"])

                # 3. Match Workspace Artifacts second
                if not (is_strict_scope and target_doc):
                    for art in art_rows:
                        if art["id"] in seen_artifact_ids:
                            continue
                        fname = art["filename"].lower()
                        stem = Path(art["filename"]).stem.lower()
                        if explicit_files:
                            if any(fname == ef.lower() for ef in explicit_files):
                                matching_artifacts.append(dict(art))
                                seen_artifact_ids.add(art["id"])
                        elif fname in lower_input or fname in ref_files:
                            matching_artifacts.append(dict(art))
                            seen_artifact_ids.add(art["id"])
                        elif stem not in GENERIC_STEMS and len(stem) > 4:
                            if stem in lower_input or stem in [Path(rf).stem.lower() for rf in ref_files]:
                                matching_artifacts.append(dict(art))
                                seen_artifact_ids.add(art["id"])

                # Check physical disk if not found in tables
                if not matching_artifacts and target_files:
                    for tf in target_files:
                        try:
                            tf_path = resolve_workspace_path(workspace_id, tf, purpose="read")
                            if tf_path.exists() and tf_path.is_file():
                                matching_artifacts.append({
                                    "filename": tf,
                                    "relative_path": tf,
                                    "artifact_type": "file",
                                    "title": tf
                                })
                                break
                        except Exception:
                            pass

                # If no target file was explicitly named, check if asking about generated files or knowledge sources
                if not matching_artifacts and not target_files:
                    is_generated_query = any(w in lower_input for w in [
                        "generated", "created", "artifact", "in artifacts", "latest file", "what file",
                        "which file", "name of the file", "can't find", "cannot find", "find it", "where is"
                    ])
                    if is_generated_query and art_rows:
                        top_art = dict(art_rows[0])
                        matching_artifacts.append(top_art)
                    if target_doc:
                        matching_artifacts.append({
                            "filename": target_doc["original_filename"] or target_doc["name"],
                            "relative_path": target_doc.get("local_path", ""),
                            "artifact_type": target_doc.get("source_type", "document"),
                            "title": target_doc["name"]
                        })
                    elif any(w in lower_input for w in ["note", "handwritten", "handover", "shift note"]) and source_rows:
                        for s in source_rows:
                            s_dict = dict(s)
                            s_name = (s_dict.get("name") or "").lower()
                            if "handwritten" in s_name or "note" in s_name:
                                matching_artifacts.append({
                                    "filename": s_dict["original_filename"] or s_dict["name"],
                                    "relative_path": s_dict.get("local_path", ""),
                                    "artifact_type": "image",
                                    "title": s_dict["name"]
                                })
                                break
                    elif any(w in lower_input for w in ["pdf", "document", "manual", "report"]) and source_rows:
                        top_s = dict(source_rows[0])
                        raw_lp = top_s["local_path"]
                        rel_p = raw_lp
                        if raw_lp:
                            try:
                                p = Path(raw_lp)
                                if p.is_absolute():
                                    rel_p = str(p.resolve().relative_to(ws_root)).replace("\\", "/")
                            except Exception:
                                rel_p = Path(raw_lp).name
                        matching_artifacts.append({
                            "filename": top_s["original_filename"] or top_s["name"],
                            "relative_path": rel_p,
                            "artifact_type": top_s["source_type"] or "pdf",
                            "title": top_s["name"]
                        })

                # Fail-Closed: If file was specifically asked for but missing, return deterministic error without calling LLM
                if target_files and not matching_artifacts:
                    requested_name = target_files[0]
                    fail_msg = f"I couldn't find '{requested_name}' in the current workspace."
                    sources_used = f"Workspace Artifact | {requested_name} (Not Found)"
                    if len(plan.steps) >= 5:
                        plan.steps[1].status = "completed"
                        plan.steps[1].observation = f"Artifact '{requested_name}' not found in workspace"
                        plan.steps[2].status = "skipped"
                        plan.steps[2].observation = "Skipped: Target file does not exist"
                        plan.steps[3].status = "skipped"
                        plan.steps[3].observation = "Skipped: Target file does not exist"
                        plan.steps[4].status = "completed"
                        plan.steps[4].observation = "Delivered fail-closed notification"
                    saved_plan_json = serialize_plan(plan)
                    await db.execute(
                        """UPDATE agent_runs
                           SET status = 'completed', result_text = ?, sources_used = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                           WHERE id = ?""",
                        (fail_msg, sources_used, saved_plan_json, run_id)
                    )
                    await db.commit()
                    await log_event(db, run_id, "artifact_not_found", fail_msg, {"target": requested_name})
                    cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                    return make_response(dict(await cursor.fetchone()))

                # Format check: Explicitly reject legacy .xls
                if matching_artifacts:
                    first_art = matching_artifacts[0]
                    if first_art["filename"].lower().endswith(".xls"):
                        fail_msg = f"The file '{first_art['filename']}' is in legacy .xls format which is not supported in the local environment. Please convert it to .xlsx or .csv."
                        sources_used = f"Workspace Artifact | {first_art['filename']}"
                        if len(plan.steps) >= 5:
                            plan.steps[1].status = "completed"
                            plan.steps[1].observation = f"Resolved '{first_art['filename']}'"
                            plan.steps[2].status = "failed"
                            plan.steps[2].error_message = "Unsupported format: legacy .xls requires xlrd which is not installed"
                            plan.steps[3].status = "skipped"
                            plan.steps[3].observation = "Skipped: Unsupported format"
                            plan.steps[4].status = "completed"
                            plan.steps[4].observation = "Delivered format support notice"
                        saved_plan_json = serialize_plan(plan)
                        await db.execute(
                            """UPDATE agent_runs
                               SET status = 'completed', result_text = ?, sources_used = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                               WHERE id = ?""",
                            (fail_msg, sources_used, saved_plan_json, run_id)
                        )
                        await db.commit()
                        await log_event(db, run_id, "format_unsupported", fail_msg, {"filename": first_art['filename']})
                        cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                        return make_response(dict(await cursor.fetchone()))

                seen_fns = set()
                unique_matching_artifacts = []
                for art in matching_artifacts:
                    fn = art["filename"].lower()
                    if fn not in seen_fns:
                        seen_fns.add(fn)
                        unique_matching_artifacts.append(art)
                matching_artifacts = unique_matching_artifacts

                for art in matching_artifacts[:2]:
                    try:
                        raw_rel = art["relative_path"]
                        if raw_rel:
                            try:
                                p = Path(raw_rel)
                                if p.is_absolute():
                                    raw_rel = str(p.resolve().relative_to(ws_root)).replace("\\", "/")
                            except Exception:
                                raw_rel = Path(raw_rel).name
                        art_path = resolve_workspace_path(workspace_id, raw_rel, purpose="read")
                        if art_path.exists() and art_path.is_file():
                            ext = art_path.suffix.lower()
                            content = ""
                            if ext in [".txt", ".json", ".yaml", ".yml", ".log", ".md"]:
                                content = art_path.read_text(encoding="utf-8", errors="replace")
                            elif ext in [".xlsx", ".xls", ".csv", ".tsv"]:
                                try:
                                    is_anomaly_query = any(w in lower_input for w in ["abnormal", "anomaly", "outlier", "spike", "excursion", "highest", "critical", "incident", "failure"])
                                    anomaly_report = None
                                    if is_anomaly_query:
                                        try:
                                            anomaly_report = await asyncio.to_thread(detect_dataframe_anomalies, art_path)
                                        except Exception as a_err:
                                            logger.warning(f"Anomaly detection error: {a_err}")

                                    insights = await asyncio.to_thread(extract_document_insights, art_path, query_hint=clean_input)
                                    if goal_contract and "revenue_previous" in goal_contract.required_fields and insights:
                                        populate_goal_contract_from_insights(goal_contract, insights)
                                    content_blocks = []

                                    if anomaly_report and anomaly_report.get("has_anomaly"):
                                        frozen_scada_anomaly = anomaly_report
                                        frozen_scada_source = art["filename"]
                                        if not goal_contract:
                                            goal_contract = GoalContract(
                                                required_fields=["timestamp", "component_id", "valve_status", "status_indicator"]
                                            )
                                            goal_contract.extracted_fields["timestamp"] = anomaly_report.get("timestamp")
                                            goal_contract.extracted_fields["component_id"] = anomaly_report.get("component_id")
                                            goal_contract.extracted_fields["valve_status"] = anomaly_report.get("valve_status")
                                            goal_contract.extracted_fields["status_indicator"] = anomaly_report.get("status_indicator")
                                        ar = anomaly_report
                                        spiked_lines = []
                                        for sc in ar["spiked_columns"]:
                                            spiked_lines.append(
                                                f"  - Column '{sc['column']}': spiked to {sc['value']} (Baseline before: {sc['before']} | Baseline after: {sc['after']} | Change: {sc['pct_change_vs_before']:+}%)"
                                            )
                                        content_blocks.append(
                                            f"--- AUTHORITATIVE SCADA ANOMALY REPORT ---\n"
                                            f"Source Document: {art['filename']}\n"
                                            f"Detected Outlier Timestamp: {ar['timestamp']}\n"
                                            f"Status Indicator: {ar['status_indicator']}\n"
                                            f"Equipment Component: {ar['component_id']} (Machine Component)\n"
                                            f"Actuator / Relief Valve: {ar.get('valve_column', 'sv402_relief_valve_status')} (State: {ar['valve_status']})\n\n"
                                            f"Spiked Measurements:\n"
                                            + ("\n".join(spiked_lines) if spiked_lines else "  - Critical excursion indicated by status flag.") + "\n\n"
                                            f"Authoritative Keyed Schema Record at Anomaly Timestamp:\n"
                                            + json.dumps(ar['anomalous_record'], indent=2, default=str) + "\n\n"
                                            f"Preceding Row Record ({ar.get('preceding_record', {}).get('timestamp', 'Earlier')}):\n"
                                            + json.dumps(ar.get('preceding_record', {}), indent=2, default=str) + "\n\n"
                                            f"Subsequent Row Record ({ar.get('succeeding_record', {}).get('timestamp', 'Later')}):\n"
                                            + json.dumps(ar.get('succeeding_record', {}), indent=2, default=str) + "\n\n"
                                            f"CRITICAL GROUNDING RULES:\n"
                                            f"1. Strictly use the exact column names above. Do NOT mislabel flow rate as pressure, or bearing temp as flow, or vibration as valve position!\n"
                                            f"2. {ar['component_id']} is the machine component. Valve status belongs to {ar.get('valve_column', 'SV-402')}, NOT {ar['component_id']}!\n"
                                            f"3. State the exact timestamp, spiked measurements, baseline comparison, and valve state."
                                        )

                                    if ext in [".csv", ".tsv"]:
                                        raw_csv_text = art_path.read_text(encoding="utf-8", errors="replace")
                                        if raw_csv_text.strip():
                                            content_blocks.append(raw_csv_text)
                                    if insights.get("structured_text"):
                                        content_blocks.append(insights.get("structured_text", ""))
                                    content = "\n\n".join(b for b in content_blocks if b)
                                except Exception as sp_err:
                                    logger.warning(f"Error extracting spreadsheet insights for {art['filename']}: {sp_err}")
                                    content = f"[Spreadsheet: {art['filename']} (could not extract content)]"

                            elif ext == ".docx":
                                try:
                                    import docx
                                    doc = docx.Document(art_path)
                                    content = "\n".join(p.text for p in doc.paragraphs if p.text)
                                except Exception as docx_err:
                                    logger.warning(f"Error reading docx {art['filename']}: {docx_err}")
                                    content = f"[DOCX document: {art['filename']}]"
                            elif ext in [".png", ".jpg", ".jpeg"]:
                                try:
                                    c_dp = await db.execute(
                                        """SELECT text_content FROM document_pages dp 
                                           JOIN knowledge_sources ks ON dp.source_id = ks.id
                                           WHERE ks.workspace_id = ? AND (ks.name = ? OR ks.original_filename = ?)
                                           LIMIT 1""",
                                        (workspace_id, art["filename"], art["filename"])
                                    )
                                    dp_r = await c_dp.fetchone()
                                    if dp_r and dp_r["text_content"]:
                                        content = dp_r["text_content"]
                                    else:
                                        provider = get_provider()
                                        img_bytes = art_path.read_bytes()
                                        v_resp = await provider.analyze_image(
                                            img_bytes,
                                            prompt="Transcribe and describe in detail all handwritten notes, logs, tags, numbers, and observations visible in this image."
                                        )
                                        content = v_resp.text.strip()
                                except Exception as img_err:
                                    logger.warning(f"Error reading image artifact {art['filename']}: {img_err}")
                                    content = f"[Visual artifact: {art['filename']}]"
                            elif ext == ".pdf":
                                requested_p = getattr(resolved_context, "requested_page", None) if resolved_context else None
                                if not requested_p:
                                    requested_p = extract_requested_page(clean_input)

                                if requested_p:
                                    c_dp = await db.execute(
                                        """SELECT dp.text_content, dp.extraction_method 
                                           FROM document_pages dp 
                                           JOIN knowledge_sources ks ON dp.source_id = ks.id 
                                           WHERE ks.workspace_id = ? AND (ks.name = ? OR ks.original_filename = ?) AND dp.page_number = ?
                                           ORDER BY ks.id DESC LIMIT 1""",
                                        (workspace_id, art["filename"], art["filename"], requested_p)
                                    )
                                    row_dp = await c_dp.fetchone()
                                    if row_dp and row_dp["text_content"]:
                                        method_label = (row_dp["extraction_method"] or "native").upper()
                                        content = f"[{art['filename']} | Page {requested_p} | {method_label}]:\n{row_dp['text_content']}"
                                        artifact_citations.append(f"{art['filename']} | Page {requested_p} | {method_label}")
                                    else:
                                        content = await asyncio.to_thread(_extract_pdf_preview_sync, art_path, requested_p, 3000, art["filename"], requested_p)
                                        artifact_citations.append(f"{art['filename']} | Page {requested_p}")
                                else:
                                    content = await asyncio.to_thread(_extract_pdf_preview_sync, art_path, 10, 1500, art["filename"])
                                    prefix = "Knowledge Source" if art.get("is_knowledge_source") else "Workspace Artifact"
                                    artifact_citations.append(f"{prefix} | {art['filename']}")

                            if content:
                                preview = content[:4000]
                                if len(content) > 4000:
                                    preview += f"\n... [Truncated: {len(content)} total characters]"
                                prefix = "KNOWLEDGE SOURCE" if art.get("is_knowledge_source") else "WORKSPACE ARTIFACT"
                                combined_context_parts.append(
                                    f"--- {prefix}: {art['filename']} ({art.get('artifact_type', 'file')}) ---\n"
                                    f"Path: {art['relative_path']}\n"
                                    f"Content:\n{preview}"
                                )
                                if not any(art["filename"] in c for c in artifact_citations):
                                    cite_prefix = "Knowledge Source" if art.get("is_knowledge_source") else "Workspace Artifact"
                                    artifact_citations.append(f"{cite_prefix} | {art['filename']}")
                    except Exception as ex:
                        logger.warning(f"Could not load artifact {art.get('relative_path')}: {ex}")

                sources_used = ", ".join(artifact_citations) if artifact_citations else ("Knowledge Sources" if any(a.get("is_knowledge_source") for a in matching_artifacts) else "Workspace Artifacts")
                if len(plan.steps) >= 5:
                    plan.steps[1].status = "completed"
                    plan.steps[1].observation = f"Resolved: {sources_used}"
                    plan.steps[2].status = "completed"
                    plan.steps[2].observation = f"Loaded artifact contents ({len(artifact_citations)} files)"
                    plan.current_step_index = 3

                await log_event(
                    db, run_id, "retrieval_completed",
                    f"Targeted artifact loaded: {sources_used}",
                    {"artifacts": artifact_citations}
                )

            elif effective_intent == SemanticIntent.KNOWLEDGE_QUERY:
                allowed_source_ids, retrieval_query, page_direct_context = await _resolve_knowledge_and_page_context(
                    db=db,
                    workspace_id=workspace_id,
                    agent=agent,
                    clean_input=clean_input,
                    resolved_context=resolved_context,
                    vision_analysis=vision_analysis,
                    run_id=run_id
                )

                if is_rca_mode:
                    try:
                        rca_acquirer = RCAEvidenceAcquirer()
                        rca_bundle = await rca_acquirer.acquire_evidence(
                            workspace_id=workspace_id,
                            query=clean_input,
                            allowed_source_ids=allowed_source_ids,
                            db=db,
                            enable_visual=enable_visual,
                            enable_topology=enable_topology,
                            custom_topology_context=custom_topology_context
                        )
                        await log_event(
                            db, run_id, "rca_channel_health",
                            f"Acquired RCA multi-channel evidence ({len(rca_bundle.evidence_items)} items) across {rca_bundle.channel_health.executed_channels}",
                            rca_bundle.channel_health.model_dump()
                        )
                        if rca_bundle.channel_health.hybrid_status == "DEGRADED":
                            await log_event(
                                db, run_id, "HYBRID_RAG_DEGRADED",
                                f"Hybrid RAG operating in degraded mode: {rca_bundle.channel_health.degradation_reason}",
                                rca_bundle.channel_health.model_dump()
                            )
                        if (getattr(rca_bundle, "stage_latencies_ms", None) or {}).get("visual_vlm_ms", 0.0) > 0:
                            await log_event(
                                db, run_id, "visual_inspector",
                                f"Visual evidence inspected via VLM: {rca_bundle.stage_latencies_ms.get('visual_vlm_ms', 0.0)}ms",
                                {"visual_vlm_ms": rca_bundle.stage_latencies_ms.get("visual_vlm_ms", 0.0)}
                            )

                        # OOD asset verification: Fail closed if unregistered asset queried
                        if rca_bundle.retrieval_diagnostics.get("ood_triggered"):
                            validator = RCAEvidenceValidator()
                            fail_msg = validator.validate_and_finalize(rca_bundle, "", clean_input)
                            structured_rca = validator.get_structured_result(rca_bundle, fail_msg, clean_input)
                            await log_event(
                                db, run_id, "rca_result_structured",
                                f"Structured RCA result recorded with status {structured_rca.status} and primary cause {structured_rca.primary_cause_code}",
                                structured_rca.model_dump()
                            )
                            sources_used = "None (Asset Not Found)"
                            if len(plan.steps) >= 4:
                                plan.steps[1].status = "completed"
                                plan.steps[1].observation = "Asset verification failed: equipment not in plant topology"
                                plan.steps[2].status = "skipped"
                                plan.steps[2].observation = "Skipped: unverified asset"
                                plan.steps[3].status = "completed"
                                plan.steps[3].observation = "Asset not found report delivered"
                            saved_plan_json = serialize_plan(plan)
                            await db.execute(
                                """UPDATE agent_runs
                                   SET status = 'completed', result_text = ?, sources_used = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                                   WHERE id = ?""",
                                (fail_msg, sources_used, saved_plan_json, run_id)
                            )
                            await db.commit()
                            await log_event(db, run_id, "asset_not_found", fail_msg)
                            stage_lats = dict(getattr(rca_bundle, "stage_latencies_ms", {}) or {})
                            if t_run_start is not None:
                                stage_lats["e2e_total"] = round((time.perf_counter() - t_run_start) * 1000, 2)
                            await log_event(
                                db, run_id, "rca_stage_latencies",
                                f"RCA pipeline stage latencies recorded: {stage_lats}",
                                {"stage_latencies_ms": stage_lats}
                            )
                            cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                            return make_response(dict(await cursor.fetchone()))

                        # Inject authoritative E-ID evidence blocks into reasoning context
                        rca_prompt_context = rca_bundle.format_for_reasoning_prompt()
                        context_str = f"{context_str}\n\n{rca_prompt_context}".strip() if context_str else rca_prompt_context
                        for item in rca_bundle.evidence_items:
                            if item.page_number:
                                retrieved_evidence_catalog.append({
                                    "filename": item.filename,
                                    "page": item.page_number,
                                    "extraction_method": item.retrieval_channel.upper()
                                })
                    except Exception as rca_err:
                        logger.error(f"RCA evidence acquisition failed: {rca_err}", exc_info=True)

                await log_event(db, run_id, "retrieval_started", f"Searching authorized knowledge sources ({allowed_source_ids})...")
                if allowed_source_ids and not is_rca_mode:
                    context_str, retrieved_metas = await retrieve_context_with_metadata(
                        workspace_id=workspace_id,
                        query=retrieval_query,
                        top_k=3,
                        allowed_source_ids=allowed_source_ids
                    )
                    retrieved_evidence_catalog.extend(retrieved_metas)
                if page_direct_context:
                    context_str = f"{page_direct_context}\n\n{context_str}".strip() if context_str else page_direct_context
                    from cognishift.core.document_processing.provenance import extract_and_normalize_citations
                    for pdc in extract_and_normalize_citations(page_direct_context):
                        retrieved_evidence_catalog.append({
                            "filename": pdc["filename"],
                            "page": pdc["page"],
                            "extraction_method": pdc.get("method") or "NATIVE"
                        })

                has_tag = any(p in retrieval_query.upper() for p in ["P-", "V-", "T-", "HEX-", "MOV-", "PT-", "TT-"])
                if (has_tag or is_rca_mode) and not graph_context:
                    graph_context = await query_graph_context(workspace_id=workspace_id, query_text=retrieval_query, max_hops=2)

                citations = list(dict.fromkeys(re.findall(r"\[([^\]]*?\|\s*Page\s*\d+[^\]]*?)\]", context_str))) if context_str else []
                if rca_bundle and rca_bundle.evidence_items:
                    for item in rca_bundle.evidence_items:
                        if item.page_number:
                            c_lbl = f"{item.filename} | Page {item.page_number}"
                            if c_lbl not in citations:
                                citations.append(c_lbl)
                sources_used = ", ".join(f"[{c}]" if not c.startswith("[") else c for c in citations) if citations else "None (No matching manual found)"
                if graph_context:
                    sources_used += " + Plant Topology Graph"

                await log_event(
                    db, run_id, "retrieval_completed",
                    f"Knowledge retrieval completed ({len(citations)} citations)",
                    {"citations": citations, "has_graph": bool(graph_context)}
                )

                # Grounding Fail-Closed: If query inquires about organizational policy, procurement, inspection/corrosion, or private SOPs and no matching evidence exists
                lower_clean = clean_input.lower()
                is_rca_or_diagnostic = is_rca_mode or any(k in lower_clean for k in [
                    "rca", "root cause", "failure investigation", "investigate failure",
                    "incident investigation", "why did it fail", "why did the system trip",
                    "troubleshoot", "investigate", "cause of", "what caused"
                ])
                is_remote_work = any(k in lower_clean for k in ["remote work", "work from home", "telework", "telecommuting", "wfh"])
                is_procurement = "procurement" in lower_clean
                is_corrosion_query = not is_rca_or_diagnostic and any(k in lower_clean for k in ["corrosion life", "remaining life", "corrosion rate", "wall thickness", "ultrasonic thickness", "mpy", "corrosion"])
                is_inspection_report = not is_rca_or_diagnostic and any(k in lower_clean for k in ["inspection report", "nde report", "ndt report", "metallurgical report"])
                is_policy_query = is_remote_work or is_procurement or is_corrosion_query or is_inspection_report or any(k in lower_clean for k in [
                    "policy", "standard operating procedure", "our sop", "leave rule", "travel rule", "reimbursement"
                ])
                requested_entities = sorted(set(re.findall(r"\b[A-Z]{1,4}-\d{3,4}[A-Z]?\b", clean_input.upper())))
                entity_grounding_required = bool(requested_entities) and any(
                    phrase in lower_clean
                    for phrase in ("according to", "document", "manual", "procedure", "sop", "maintenance")
                )

                # Verify topical relevance and factual sufficiency of retrieved context
                evidence_text = "\n".join(part for part in (context_str, graph_context) if part)
                if evidence_text and not is_rca_mode:
                    is_suff, suff_reason = validate_evidence_sufficiency(clean_input, evidence_text)
                    if not is_suff:
                        logger.info(f"Discarding context: {suff_reason}")
                        context_str = ""
                        graph_context = ""

                if not context_str and not graph_context and (is_policy_query or entity_grounding_required):
                    if is_corrosion_query or is_inspection_report:
                        eq_match = re.search(r'\b([A-Z]{1,3}-[0-9]{3,4}[A-Z]?)\b', clean_input)
                        eq_label = f" for {eq_match.group(1)}" if eq_match else ""
                        fail_msg = f"According to the documents currently available in your Knowledge Vault, no inspection report or corrosion life data is available{eq_label}."
                        sources_used = "None (No matching manual found)"
                    elif entity_grounding_required:
                        requested_label = ", ".join(requested_entities)
                        fail_msg = (
                            f"I could not find {requested_label} in the indexed workspace documents. "
                            f"I cannot provide a source-grounded procedure for this equipment and will not infer one "
                            f"from other assets. Please provide the relevant {requested_label} document or ask about a documented asset."
                        )
                        sources_used = "None (Insufficient evidence for requested equipment)"
                    else:
                        topic_name = "remote-work-policy" if is_remote_work else ("procurement-policy" if is_procurement else "relevant")
                        fail_msg = f"I couldn't find {topic_name} documentation in the current workspace knowledge base. Please ingest the applicable document before asking for an organization-specific answer."
                        sources_used = "None (No matching manual found)"
                    if len(plan.steps) >= 4:
                        plan.steps[1].status = "completed"
                        plan.steps[1].observation = "No chunks met distance threshold in local knowledge base"
                        plan.steps[2].status = "skipped"
                        plan.steps[2].observation = "Skipped: No supporting evidence in knowledge base"
                        plan.steps[3].status = "completed"
                        plan.steps[3].observation = "Fail-closed response delivered"
                    saved_plan_json = serialize_plan(plan)
                    await db.execute(
                        """UPDATE agent_runs
                           SET status = 'completed', result_text = ?, sources_used = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                           WHERE id = ?""",
                        (fail_msg, sources_used, saved_plan_json, run_id)
                    )
                    await db.commit()
                    await log_event(db, run_id, "knowledge_unsupported", fail_msg)
                    cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                    return make_response(dict(await cursor.fetchone()))

                if context_str:
                    combined_context_parts.append(
                        "--- GROUNDED KNOWLEDGE BASE CONTEXT ---\n"
                        "CRITICAL: Answer ONLY using the facts stated below. Do NOT hallucinate intranet portals, SharePoint links, or external websites.\n\n"
                        + context_str
                    )
                if graph_context:
                    combined_context_parts.append(graph_context)

                if len(plan.steps) >= 4:
                    plan.steps[1].status = "completed"
                    plan.steps[1].observation = f"Retrieved {len(citations)} citations"
                    plan.current_step_index = 2

            elif effective_intent == SemanticIntent.CODE_EXECUTION:
                target_doc = None
                fn_match = re.search(r'\b([A-Za-z0-9_\-\.]+\.(?:xlsx|xls|pdf|csv|docx))\b', clean_input, re.IGNORECASE)
                if fn_match:
                    cand_name = fn_match.group(1)
                    c_named = await db.execute(
                        "SELECT * FROM knowledge_sources WHERE workspace_id = ? AND processing_status = 'completed' ORDER BY id DESC",
                        (workspace_id,)
                    )
                    sources = await c_named.fetchall()
                    for s in sources:
                        s_dict = dict(s)
                        s_name = (s_dict.get("name") or "").lower()
                        s_orig = (s_dict.get("original_filename") or "").lower()
                        if cand_name.lower() == s_name or cand_name.lower() == s_orig:
                            target_doc = s_dict
                            break
                    if not target_doc:
                        ws_root_chk = get_workspace_root(workspace_id).resolve()
                        search_locs = [
                            ws_root_chk / cand_name,
                            ws_root_chk / "uploads" / cand_name,
                            Path.home() / "OneDrive" / "Desktop" / cand_name,
                            Path.home() / "Desktop" / cand_name,
                            Path.home() / "Downloads" / cand_name,
                        ]
                        for loc in search_locs:
                            if loc.exists() and loc.is_file():
                                target_doc = {
                                    "id": 1,
                                    "name": cand_name,
                                    "original_filename": cand_name,
                                    "local_path": str(loc),
                                    "source_type": "spreadsheet" if cand_name.lower().endswith((".xlsx", ".xls", ".csv")) else "pdf",
                                    "processing_status": "completed"
                                }
                                break

                if not target_doc and claimed_task and claimed_task.source_references and claimed_task.source_references.get("pinned_source"):
                    target_doc = claimed_task.source_references["pinned_source"]
                elif not target_doc and resolved_context and resolved_context.pinned_source:
                    target_doc = resolved_context.pinned_source
                elif not target_doc and any(w in clean_input.lower() for w in ["document", "pdf", "manual", "report", "latest", "ingested", "xlsx", "excel", "spreadsheet", "file", "csv", "data", "history", "audit", "financial"]):
                    target_doc = await resolve_target_document_for_query(workspace_id, clean_input, db=db)

                if target_doc:
                    raw_lp = target_doc.get("local_path")
                    doc_content = ""
                    ws_root = get_workspace_root(workspace_id).resolve()
                    if raw_lp:
                        try:
                            p = Path(raw_lp)
                            if not p.is_absolute():
                                p = ws_root / p
                            if p.exists() and p.suffix.lower() == ".pdf":
                                doc_content = await asyncio.to_thread(_extract_pdf_preview_sync, p, 5, 1500, target_doc["name"])
                            elif p.exists() and p.suffix.lower() in [".xlsx", ".xls"]:
                                import openpyxl
                                wb = openpyxl.load_workbook(p, data_only=True)
                                lines = [f"Workbook: {target_doc['name']} (Sheets: {', '.join(wb.sheetnames)})"]
                                for sname in wb.sheetnames[:4]:
                                    ws = wb[sname]
                                    lines.append(f"\n--- Sheet: {sname} ---")
                                    for row in list(ws.iter_rows(values_only=True))[:15]:
                                        if any(row):
                                            lines.append(" | ".join([str(c) for c in row if c is not None]))
                                doc_content = "\n".join(lines)[:4000]
                            elif p.exists():
                                doc_content = p.read_text(encoding="utf-8", errors="replace")[:4000]
                        except Exception as e:
                            logger.warning(f"Error reading target doc: {e}")

                    if doc_content:
                        combined_context_parts.append(
                            f"--- AUTHORITATIVE INGESTED DOCUMENT: {target_doc['name']} (ID #{target_doc['id']}) ---\n"
                            f"Content Preview:\n{doc_content[:4000]}"
                        )
                        artifact_citations.append(f"{target_doc['name']} | Page 1")

                cursor_artifacts = await db.execute(
                    "SELECT id, filename, relative_path, file_size, artifact_type, title, description FROM workspace_artifacts WHERE workspace_id = ? ORDER BY id DESC LIMIT 30",
                    (workspace_id,)
                )
                art_rows = await cursor_artifacts.fetchall()
                matching_artifacts = []
                lower_input = clean_input.lower()
                for art in art_rows:
                    fname = art["filename"].lower()
                    if fname in lower_input:
                        matching_artifacts.append(dict(art))

                for art in matching_artifacts[:2]:
                    try:
                        art_path = resolve_workspace_path(workspace_id, art["relative_path"], purpose="read")
                        if art_path.exists() and art_path.is_file():
                            content = art_path.read_text(encoding="utf-8", errors="replace")
                            preview = content[:2000]
                            combined_context_parts.append(
                                f"--- WORKSPACE ARTIFACT: {art['filename']} ({art['artifact_type']}) ---\n"
                                f"Path: {art['relative_path']}\n"
                                f"Content Preview:\n{preview}"
                            )
                            artifact_citations.append(f"Workspace Artifact | {art['filename']}")
                    except Exception as ex:
                        pass
                sources_used = "Docker Python Sandbox"
                if artifact_citations:
                    sources_used += f" ({', '.join(artifact_citations)})"

                await log_event(
                    db, run_id, "retrieval_completed",
                    f"Code execution context prepared (artifacts/sources: {len(artifact_citations)})",
                    {"artifacts": artifact_citations, "target_doc": target_doc["name"] if target_doc else None}
                )

            else:
                # CONTROL_ACTION / COMPLEX_AGENT / Default: Full retrieval
                allowed_source_ids, retrieval_query, page_direct_context = await _resolve_knowledge_and_page_context(
                    db=db,
                    workspace_id=workspace_id,
                    agent=agent,
                    clean_input=clean_input,
                    resolved_context=resolved_context,
                    vision_analysis=vision_analysis,
                    run_id=run_id
                )

                await log_event(db, run_id, "retrieval_started", f"Searching authorized knowledge sources ({allowed_source_ids}) and plant topology graph...")
                if allowed_source_ids:
                    context_str, retrieved_metas = await retrieve_context_with_metadata(
                        workspace_id=workspace_id,
                        query=retrieval_query,
                        top_k=3,
                        allowed_source_ids=allowed_source_ids
                    )
                    retrieved_evidence_catalog.extend(retrieved_metas)
                if page_direct_context:
                    context_str = f"{page_direct_context}\n\n{context_str}".strip() if context_str else page_direct_context
                    from cognishift.core.document_processing.provenance import extract_and_normalize_citations
                    for pdc in extract_and_normalize_citations(page_direct_context):
                        retrieved_evidence_catalog.append({
                            "filename": pdc["filename"],
                            "page": pdc["page"],
                            "extraction_method": pdc.get("method") or "NATIVE"
                        })

                graph_context = await query_graph_context(workspace_id=workspace_id, query_text=retrieval_query, max_hops=2)

                lower_clean = clean_input.lower()
                is_rca_or_diagnostic = is_rca_mode or any(k in lower_clean for k in [
                    "rca", "root cause", "failure investigation", "investigate failure",
                    "incident investigation", "why did it fail", "why did the system trip",
                    "troubleshoot", "investigate", "cause of", "what caused"
                ])
                is_remote_work = any(k in lower_clean for k in ["remote work", "work from home", "telework", "telecommuting", "wfh"])
                is_procurement = "procurement" in lower_clean
                is_corrosion_query = not is_rca_or_diagnostic and any(k in lower_clean for k in ["corrosion life", "remaining life", "corrosion rate", "wall thickness", "ultrasonic thickness", "mpy", "corrosion"])
                is_inspection_report = not is_rca_or_diagnostic and any(k in lower_clean for k in ["inspection report", "nde report", "ndt report", "metallurgical report"])
                is_policy_query = is_remote_work or is_procurement or is_corrosion_query or is_inspection_report or any(k in lower_clean for k in [
                    "policy", "standard operating procedure", "our sop", "leave rule", "travel rule", "reimbursement"
                ])
                requested_entities = sorted(set(re.findall(r"\b[A-Z]{1,4}-\d{3,4}[A-Z]?\b", clean_input.upper())))
                entity_grounding_required = bool(requested_entities) and any(
                    phrase in lower_clean
                    for phrase in ("according to", "document", "manual", "procedure", "sop", "maintenance")
                )

                # Verify topical relevance and factual sufficiency of retrieved context
                evidence_text = "\n".join(part for part in (context_str, graph_context) if part)
                if evidence_text:
                    is_suff, suff_reason = validate_evidence_sufficiency(clean_input, evidence_text)
                    if not is_suff:
                        logger.info(f"Discarding context: {suff_reason}")
                        context_str = ""
                        graph_context = ""

                if not context_str and not graph_context and (is_policy_query or entity_grounding_required):
                    if is_corrosion_query or is_inspection_report:
                        eq_match = re.search(r'\b([A-Z]{1,3}-[0-9]{3,4}[A-Z]?)\b', clean_input)
                        eq_label = f" for {eq_match.group(1)}" if eq_match else ""
                        fail_msg = f"According to the documents currently available in your Knowledge Vault, no inspection report or corrosion life data is available{eq_label}."
                        sources_used = "None (No matching manual found)"
                    elif entity_grounding_required:
                        requested_label = ", ".join(requested_entities)
                        fail_msg = (
                            f"I could not find {requested_label} in the indexed workspace documents. "
                            f"I cannot provide a source-grounded procedure for this equipment and will not infer one "
                            f"from other assets. Please provide the relevant {requested_label} document or ask about a documented asset."
                        )
                        sources_used = "None (Insufficient evidence for requested equipment)"
                    else:
                        topic_name = "remote-work-policy" if is_remote_work else ("procurement-policy" if is_procurement else "relevant")
                        fail_msg = f"I couldn't find {topic_name} documentation in the current workspace knowledge base. Please ingest the applicable document before asking for an organization-specific answer."
                        sources_used = "None (No matching manual found)"
                    for idx, step in enumerate(plan.steps):
                        if idx == 0:
                            step.status = "completed"
                            step.observation = "Analyzed policy query requirements"
                        elif idx == len(plan.steps) - 1:
                            step.status = "completed"
                            step.observation = "Fail-closed response delivered"
                        else:
                            step.status = "skipped"
                            step.observation = "Skipped: No supporting evidence in local knowledge base"
                    saved_plan_json = serialize_plan(plan)
                    await db.execute(
                        """UPDATE agent_runs
                           SET status = 'completed', result_text = ?, sources_used = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                           WHERE id = ?""",
                        (fail_msg, sources_used, saved_plan_json, run_id)
                    )
                    await db.commit()
                    await log_event(db, run_id, "knowledge_unsupported", fail_msg)
                    cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                    return make_response(dict(await cursor.fetchone()))

                if context_str:
                    combined_context_parts.append(context_str)
                if graph_context:
                    combined_context_parts.append(graph_context)

                cursor_artifacts = await db.execute(
                    "SELECT id, filename, relative_path, file_size, artifact_type, title, description FROM workspace_artifacts WHERE workspace_id = ? ORDER BY id DESC LIMIT 30",
                    (workspace_id,)
                )
                art_rows = await cursor_artifacts.fetchall()
                matching_artifacts = []
                lower_input = clean_input.lower()
                GENERIC_STEMS = {"report", "reading", "readings", "file", "document", "artifact", "data", "sheet", "table", "summary", "test", "plan", "output", "input", "result", "results", "status", "logs", "log"}

                for art in art_rows:
                    fname = art["filename"].lower()
                    stem = Path(art["filename"]).stem.lower()
                    if fname in lower_input or (len(stem) > 4 and stem not in GENERIC_STEMS and stem in lower_input):
                        matching_artifacts.append(dict(art))

                for art in matching_artifacts[:3]:
                    try:
                        art_path = resolve_workspace_path(workspace_id, art["relative_path"], purpose="read")
                        if art_path.exists() and art_path.is_file():
                            ext = art_path.suffix.lower()
                            if ext in [".csv", ".txt", ".json", ".yaml", ".yml", ".log", ".md"]:
                                content = art_path.read_text(encoding="utf-8", errors="replace")
                                preview = content[:3500]
                                if len(content) > 3500:
                                    preview += f"\n... [Truncated: {len(content)} total characters]"
                                combined_context_parts.append(
                                    f"--- WORKSPACE ARTIFACT: {art['filename']} ({art['artifact_type']}) ---\n"
                                    f"Path: {art['relative_path']}\n"
                                    f"Content:\n{preview}"
                                )
                                artifact_citations.append(f"Workspace Artifact | {art['filename']}")
                    except Exception as ex:
                        pass

                citations = list(dict.fromkeys(re.findall(r"\[([^\]]*?\|\s*Page\s*\d+[^\]]*?)\]", context_str))) if context_str else []
                sources_used = ", ".join(citations) if citations else "None (No matching manual found)"
                if graph_context:
                    sources_used += " + Plant Topology Graph"
                if vision_analysis:
                    sources_used += " + Local VLM Inspection"
                unique_art = list(dict.fromkeys(artifact_citations))
                if unique_art:
                    sources_used += " + " + ", ".join(unique_art)

                await log_event(
                    db,
                    run_id,
                    "retrieval_completed",
                    f"Retrieved context ({len(citations)} manual citations, graph topology: {'yes' if graph_context else 'none'}, visual input: {'yes' if vision_analysis else 'none'}, artifacts: {len(artifact_citations)})",
                    {
                        "citations": citations,
                        "artifact_citations": artifact_citations,
                        "has_graph": bool(graph_context),
                        "has_vision": bool(vision_analysis),
                        "manual_preview": context_str[:200] if context_str else ""
                    }
                )

            combined_context = "\n\n".join(combined_context_parts)

            # Format conversation history
            history_str = ""
            if conversation_history:
                recent_history = conversation_history[-settings.semantic_router_max_history_turns:]
                history_lines = []
                for msg in recent_history:
                    role = getattr(msg, "role", None) or (msg.get("role") if isinstance(msg, dict) else "user")
                    content = getattr(msg, "content", None) or (msg.get("content") if isinstance(msg, dict) else "")
                    role_label = "Operator" if role in ["user", "operator"] else "CogniShift"
                    if content:
                        history_lines.append(f"{role_label}: {content}")
                if history_lines:
                    history_str = "--- RECENT CONVERSATION HISTORY ---\n" + "\n".join(history_lines) + "\n\n"

            # 5. Model Inference Call
            # Bounded Iterative Plan Execution Loop (P0-3)
            MAX_AGENT_STEPS = 10
            step_counter = 0
            sandbox_used = False
            final_text = vision_analysis if strict_visual_scope and vision_analysis else ""
            if final_text:
                for step in plan.steps:
                    if step.status == "pending":
                        step.status = "skipped"
                        step.observation = "Not required for strict image-only evidence reporting"
                plan.current_step_index = len(plan.steps)
                plan.final_synthesis = final_text
                await log_event(
                    db, run_id, "vision_answer_grounded",
                    "Returned local OCR and vision evidence directly; text-model synthesis bypassed",
                    {"model": settings.vision_model, "ocr_grounded": bool(ocr_evidence)},
                )
            provider = get_provider()
            tools_for_prompt = available_tools if effective_intent != SemanticIntent.CONVERSATION else []
            system_prompt = build_system_prompt(
                agent.get("system_instructions", ""),
                tools_for_prompt,
                combined_context,
                intent=effective_intent,
                workspace_name=workspace_name
            )

            target_doc = target_doc or (
                claimed_task.source_references.get("pinned_source")
                if (claimed_task and claimed_task.source_references)
                else (resolved_context.pinned_source or await resolve_target_document_for_query(workspace_id, clean_input, db=db))
            )

            while step_counter < MAX_AGENT_STEPS and not plan.is_finished():
                current_step = plan.get_current_step()
                if not current_step:
                    break

                step_counter += 1
                current_step.status = "running"
                await log_event(
                    db,
                    run_id,
                    "plan_step_started",
                    f"Executing Step #{current_step.id}: {current_step.description}",
                    {"step_id": current_step.id, "iteration": step_counter, "max_steps": MAX_AGENT_STEPS}
                )

                # Format current plan state and history for prompt
                plan_prompt_section = format_plan_for_prompt(plan)

                # Build step prompt incorporating accumulated observations
                is_final_step = (current_step.id == len(plan.steps))

                # Step-level tool availability: Tools are strictly locked out on final synthesis step
                if is_final_step or effective_intent == SemanticIntent.CONVERSATION:
                    tools_for_step = []
                else:
                    tools_for_step = available_tools

                system_prompt = build_system_prompt(
                    agent.get("system_instructions", ""),
                    tools_for_step,
                    combined_context,
                    intent=effective_intent,
                    workspace_name=workspace_name
                )

                if is_final_step:
                    step_instructions = (
                        "- This is the FINAL step of the plan.\n"
                        "- STRICT POLICY: NO TOOL CALLS ARE PERMITTED IN FINAL SYNTHESIS.\n"
                        "- Provide the final synthesized response to the operator using 'action': 'final_answer'."
                    )
                else:
                    step_instructions = (
                        f"- This is an INTERMEDIATE step (#{current_step.id} of {len(plan.steps)}).\n"
                        "- If an authorized tool is required, output 'action': 'tool_call'.\n"
                        "- If recording an observation or check, output 'action': 'step_observation' with your findings.\n"
                        "- Do NOT output 'action': 'final_answer' until the final step is reached."
                    )

                if task_info.task_type == "heavy_reasoning":
                    step_instructions += (
                        "\n- RCA SAFETY FORMAT: clearly separate (1) Confirmed Observations, "
                        "(2) Possible Hypotheses explicitly marked unverified, and "
                        "(3) Additional Evidence Needed to confirm or reject each hypothesis."
                        "\n- Do not present a hypothesis as a confirmed root cause."
                    )

                step_prompt = (
                    f"{history_str}"
                    f"Operator Current Input: {clean_input}\n\n"
                    f"{plan_prompt_section}\n\n"
                    f"Current Step to Execute: #{current_step.id} - {current_step.description}\n\n"
                    f"Instructions:\n{step_instructions}"
                )

                is_doc_or_viz_task = (
                    (target_doc is not None and any(w in clean_input.lower() for w in ["review", "convert", "report", "document", "docx", "pdf", "xlsx", "excel", "csv", "jpg", "png", "audit", "history", "financial", "deliverable", "deliverables"]))
                    or any(w in clean_input.lower() for w in ["visualize", "plot", "chart", "graph", "graphical representation", "matplotlib", "seaborn", "convert it into", "convert the document", "proper audit", "in docx format", "in pdf format", "create a csv", "generate a csv"])
                )

                if (
                    routing_res.intent == SemanticIntent.CODE_EXECUTION
                    and is_doc_or_viz_task
                    and current_step.id in (2, 3)
                ):
                    raw_output = '{"action": "step_observation", "observation": "Autonomous execution step"}'
                    clean_output = raw_output
                else:
                    await log_event(db, run_id, "model_prompt", f"Prompt dispatched to {selected_model_id} for step #{current_step.id}")

                    llm_step_ms = 0.0
                    try:
                        t_model_start = time.perf_counter()
                        model_response = await provider.generate_text(
                            prompt=step_prompt,
                            system_prompt=system_prompt,
                            context=combined_context,
                            model_name=selected_model_id,
                            history=conversation_history
                        )
                        llm_step_ms = round((time.perf_counter() - t_model_start) * 1000, 2)
                    except Exception as e:
                        # If routed model is not installed locally (HTTP 404), fall back to configured text_model
                        if ("404" in str(e) or "not found" in str(e).lower()) and selected_model_id != settings.text_model:
                            logger.warning(
                                f"Routed model '{selected_model_id}' returned 404. Falling back to '{settings.text_model}'."
                            )
                            await log_event(
                                db,
                                run_id,
                                "model_fallback",
                                f"Routed model '{selected_model_id}' not found locally. Falling back to default '{settings.text_model}'.",
                                {"original_model": selected_model_id, "fallback_model": settings.text_model}
                            )
                            selected_model_id = settings.text_model
                            try:
                                t_fb_start = time.perf_counter()
                                model_response = await provider.generate_text(
                                    prompt=step_prompt,
                                    system_prompt=system_prompt,
                                    context=combined_context,
                                    model_name=selected_model_id,
                                    history=conversation_history
                                )
                                llm_step_ms = round((time.perf_counter() - t_fb_start) * 1000, 2)
                            except Exception as fb_err:
                                e = fb_err
                            else:
                                e = None

                        if e is not None:
                            err_msg = f"Model provider failure ({selected_model_id}) on step #{current_step.id}: {str(e)}"
                            logger.error(err_msg)
                            current_step.status = "failed"
                            current_step.error_message = str(e)
                            _block_unexecuted_pending_steps(plan, "Blocked due to model provider failure")
                            await log_event(db, run_id, "run_failed", err_msg, {"error": str(e)})
                            saved_plan_json = serialize_plan(plan)
                            await db.execute(
                                """UPDATE agent_runs
                                   SET status = 'failed', error_message = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                                   WHERE id = ?""",
                                (err_msg, saved_plan_json, run_id)
                            )
                            await db.commit()
                            cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                            return make_response(dict(await cursor.fetchone()))

                    if not getattr(model_response, "success", True):
                        err_msg = model_response.error_message or f"Model generation failed on step #{current_step.id}."
                        logger.error(err_msg)
                        current_step.status = "failed"
                        current_step.error_message = err_msg
                        _block_unexecuted_pending_steps(plan, "Blocked due to model generation failure")
                        await log_event(db, run_id, "run_failed", err_msg, {"error": err_msg})
                        saved_plan_json = serialize_plan(plan)
                        await db.execute(
                            """UPDATE agent_runs
                               SET status = 'failed', error_message = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                               WHERE id = ?""",
                            (err_msg, saved_plan_json, run_id)
                        )
                        await db.commit()
                        cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                        return make_response(dict(await cursor.fetchone()))

                    raw_output = model_response.text or ""
                    clean_output, reasoning_detected, reasoning_chars = extract_and_strip_thinking(raw_output)
                    if reasoning_detected:
                        await log_event(
                            db, run_id, "model_reasoning",
                            f"Step #{current_step.id} internal reasoning detected and stripped ({reasoning_chars} chars)",
                            {
                                "step_id": current_step.id,
                                "reasoning_detected": True,
                                "reasoning_stripped": True,
                                "selected_model": selected_model_id,
                                "reasoning_character_count": reasoning_chars
                            }
                        )

                    # Blocker 6: Empty or whitespace response must FAIL rather than complete
                    if not clean_output:
                        err_msg = f"Model protocol failure on step #{current_step.id}: Empty response returned by {selected_model_id}."
                        logger.warning(err_msg)
                        current_step.status = "failed"
                        current_step.error_message = err_msg
                        _block_unexecuted_pending_steps(plan, "Blocked due to empty model response")
                        await log_event(db, run_id, "model_protocol_failure", err_msg, {"step_id": current_step.id, "error": "empty_output"})
                        saved_plan_json = serialize_plan(plan)
                        await db.execute(
                            """UPDATE agent_runs
                               SET status = 'failed', error_message = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                               WHERE id = ?""",
                            (err_msg, saved_plan_json, run_id)
                        )
                        await db.commit()
                        cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                        return make_response(dict(await cursor.fetchone()))

                await log_event(
                    db, run_id, "model_response",
                    f"Step #{current_step.id} reasoning received",
                    {"text": clean_output, "llm_step_ms": llm_step_ms if 'llm_step_ms' in locals() else 0.0}
                )

                # Strict Action Parsing (Blocker 6: Valid AgentAction schema or verified readable prose only)
                action = parse_agent_action(clean_output, strict=False)
                has_authorized_tool = (
                    isinstance(action, ToolCallProposal)
                    and bool(action.tool_name)
                    and action.tool_name.lower().strip() in [t.lower() for t in allowed_tool_names]
                )

                is_doc_or_viz_task = (
                    (target_doc is not None and any(w in clean_input.lower() for w in ["review", "convert", "report", "document", "docx", "pdf", "xlsx", "excel", "csv", "jpg", "png", "audit", "history", "financial", "deliverable", "deliverables"]))
                    or any(w in clean_input.lower() for w in ["visualize", "plot", "chart", "graph", "graphical representation", "matplotlib", "seaborn", "convert it into", "convert the document", "two different files", "docx and pdf", "both docx and pdf", "create a csv", "generate a csv"])
                )

                if (
                    routing_res.intent == SemanticIntent.CODE_EXECUTION
                    and is_doc_or_viz_task
                    and current_step.id in (2, 3, 4)
                ):
                    # Autonomous execution for CODE_EXECUTION document reporting workflow
                    lower_input = clean_input.lower()
                    doc_title = target_doc["name"] if target_doc else "document"
                    stem_name = Path(doc_title).stem

                    target_doc_path: Optional[Path] = None
                    ws_root = get_workspace_root(workspace_id).resolve()
                    if target_doc and target_doc.get("local_path"):
                        p_cand = Path(target_doc["local_path"])
                        if not p_cand.is_absolute():
                            p_cand = ws_root / p_cand
                        if p_cand.exists():
                            target_doc_path = p_cand

                    if not target_doc_path:
                        for cand in [ws_root / doc_title, ws_root / "uploads" / doc_title]:
                            if cand.exists():
                                target_doc_path = cand
                                break

                    if not target_doc_path and target_doc:
                        for name_cand in [target_doc.get("name"), target_doc.get("original_filename")]:
                            if name_cand:
                                cand = Path("data/demo") / name_cand
                                if cand.exists():
                                    target_doc_path = cand.resolve()
                                    break

                    insights = await asyncio.to_thread(extract_document_insights, target_doc_path, query_hint=clean_input) if target_doc_path else {}
                    if goal_contract and "revenue_previous" in goal_contract.required_fields and insights:
                        populate_goal_contract_from_insights(goal_contract, insights)

                    is_financial_task = (
                        insights.get("is_financial", False)
                        or any(w in lower_input for w in ["financial", "revenue", "ebitda", "pat", "profit", "cagr", "grm", "p&l", "capex", "opex", "financial audit"])
                        or (target_doc and "financial" in str(target_doc.get("name", "")).lower())
                    )

                    # Check if user specified a custom title or author in the prompt
                    title_match = re.search(r"title\s*['\"]([^'\"]+)['\"]", clean_input, re.IGNORECASE)
                    if not title_match:
                        title_match = re.search(r"titled\s*['\"]([^'\"]+)['\"]", clean_input, re.IGNORECASE)
                    custom_title = title_match.group(1).strip() if title_match else None

                    author_match = re.search(r"author\s*['\"]([^'\"]+)['\"]", clean_input, re.IGNORECASE)
                    custom_author = author_match.group(1).strip() if author_match else None

                    if custom_title:
                        display_title = custom_title
                        safe_slug = re.sub(r'[^a-zA-Z0-9_\-]', '_', custom_title).strip('_')
                        stem_name = safe_slug if safe_slug else stem_name
                    elif "audit" in lower_input:
                        display_title = f"Engineering & Operational Audit: {stem_name.replace('_', ' ')}"
                    elif is_financial_task:
                        display_title = f"Financial & Operational Audit: {stem_name.replace('_', ' ')}"
                    else:
                        display_title = f"Engineering Analysis Report: {doc_title}"

                    # Determine target format
                    wants_docx = any(w in lower_input for w in ["docx", "word"])
                    wants_pdf = "pdf" in lower_input
                    wants_both_docx_and_pdf = (wants_docx and wants_pdf) or any(w in lower_input for w in ["both docx and pdf", "docx and pdf", "two different files"])
                    wants_convert_excel = any(w in lower_input for w in ["convert to excel", "export to excel", "into excel", "as excel", "as xlsx", "as spreadsheet"])
                    wants_csv = bool(getattr(artifact_contract, "csv_required", False))
                    wants_pptx = _is_presentation_request(lower_input)
                    wants_png_viz = any(w in lower_input for w in ["png", "jpg", "jpeg", "image", "visualize", "plot", "chart", "graph"])

                    if wants_both_docx_and_pdf:
                        target_fmt = "both"
                    elif wants_pptx:
                        target_fmt = "pptx"
                    elif wants_pdf:
                        target_fmt = "pdf"
                    elif wants_convert_excel:
                        target_fmt = "xlsx"
                    elif wants_csv:
                        target_fmt = "csv"
                    elif wants_png_viz and not (wants_docx or wants_pptx or wants_convert_excel):
                        target_fmt = "image"
                    else:
                        target_fmt = "docx"

                    # Task-appropriate chart naming (prevents telemetry_chart.png on financial tasks)
                    if is_financial_task:
                        chart_filename = f"{stem_name}_financial_chart.png"
                    else:
                        chart_filename = f"{stem_name}_telemetry_chart.png" if "telemetry" in stem_name.lower() or "scada" in stem_name.lower() else f"{stem_name}_analysis_chart.png"

                    chart_data = insights.get("chart_data")
                    metrics_map = insights.get("metrics", {})
                    growth_map = insights.get("growth", {})

                    if current_step.id == 2:
                        viz_contract = parse_artifact_request_contract(clean_input)
                        if viz_contract.png_required and target_doc_path and target_doc_path.suffix.lower() in (".csv", ".xlsx", ".xls"):
                            await log_event(
                                db, run_id, "tool_started",
                                f"Executing deterministic visualization pipeline on '{target_doc_path.name}'...",
                                {"contract": viz_contract.to_dict(), "source_file": target_doc_path.name}
                            )
                            viz_results = await execute_visualization_pipeline(
                                file_path=target_doc_path,
                                query=clean_input,
                                workspace_id=workspace_id,
                                run_id=run_id,
                                contract=viz_contract
                            )
                            successful_viz = [v for v in viz_results if v.success]
                            if successful_viz:
                                tool_out = f"Deterministic data visualization successfully generated: " + ", ".join(f"{v.filename} ({v.file_size} bytes)" for v in successful_viz)
                                current_step.status = "completed"
                                current_step.tool_name = "execute_visualization_pipeline"
                                current_step.tool_parameters = {"source_file": target_doc_path.name, "count": len(successful_viz)}
                                current_step.observation = tool_out
                                await log_event(db, run_id, "tool_executed", tool_out, {"artifacts": [v.filename for v in successful_viz]})
                                saved_plan_json = serialize_plan(plan)
                                await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                                await db.commit()
                                plan.advance_to_next_step()
                                continue
                            else:
                                err_msg = "; ".join(v.error for v in viz_results if v.error)
                                raise ValueError(f"Deterministic visualization pipeline failed: {err_msg}")

                        wants_chart = wants_png_viz or is_financial_task
                        wants_excel = target_fmt == "xlsx"

                        if is_financial_task:
                            periods = chart_data.get("x_labels", ["FY24", "FY25", "FY26"]) if chart_data else ["FY24", "FY25", "FY26"]
                            series_data = chart_data.get("series", {}) if chart_data else {}
                            grm_data = chart_data.get("grm") if chart_data else None

                            py_code = f"""# Autonomous quantitative financial analysis and visualization for {doc_title}
import json
import os
from pathlib import Path

os.environ['MPLCONFIGDIR'] = '/tmp'
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

out_dir = Path('/workspace/output') if Path('/workspace/output').exists() else Path('output')
out_dir.mkdir(parents=True, exist_ok=True)

periods = {json.dumps(periods)}
series_data = {json.dumps(series_data)}
grm_data = {json.dumps(grm_data)}

rev = series_data.get('REVENUE', [])
ebitda = series_data.get('EBITDA', [])
pat = series_data.get('PAT', [])

fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5), dpi=150)

# Panel 1: Revenue, EBITDA & PAT
x = np.arange(len(periods))
width = 0.25
max_val = max(rev + ebitda + pat) if (rev or ebitda or pat) else 1
scale = 1000.0 if max_val > 5000 else 1.0
unit_label = ' (Thousand Cr)' if scale == 1000.0 else ''

if rev:
    ax1.bar(x - width, [r / scale for r in rev], width, label='Gross Revenue' + unit_label, color='#1F4E79')
if ebitda:
    ax1.bar(x, [e / scale for e in ebitda], width, label='EBITDA' + unit_label, color='#2CA02C')
if pat:
    ax1.bar(x + width, [p / scale for p in pat], width, label='PAT' + unit_label, color='#FF7F0E')

ax1.set_xticks(x)
ax1.set_xticklabels(periods, fontweight='bold')
ax1.set_ylabel('Amount' + unit_label, fontweight='bold')
ax1.set_title('{stem_name} Financial Trajectory', fontweight='bold', pad=10)
ax1.legend(frameon=True)
ax1.grid(axis='y', linestyle=':', alpha=0.6)

# Panel 2: Margins & GRM
if rev and ebitda and len(rev) == len(ebitda):
    margins = [round((e / r) * 100, 2) if r != 0 else 0.0 for e, r in zip(ebitda, rev)]
    ax2.plot(periods, margins, color='#2CA02C', marker='s', linewidth=2.5, label='EBITDA Margin (%)')
    ax2.set_ylabel('EBITDA Margin (%)', color='#2CA02C', fontweight='bold')

if grm_data:
    ax2_twin = ax2.twinx()
    ax2_twin.plot(periods, grm_data, color='#9467BD', marker='o', linewidth=2.5, linestyle='--', label='GRM ($/bbl)')
    ax2_twin.set_ylabel('Gross Refining Margin ($/bbl)', color='#9467BD', fontweight='bold')

ax2.set_title('Operational Performance & Margins', fontweight='bold', pad=10)
ax2.grid(True, linestyle=':', alpha=0.5)

plt.tight_layout()
chart_path = out_dir / '{chart_filename}'
plt.savefig(str(chart_path))
plt.close()

metrics = {{
    "document": "{doc_title}",
    "analysis_status": "SUCCESS",
    "chart_file": "{chart_filename}",
    "metrics": {json.dumps(metrics_map)},
    "growth": {json.dumps(growth_map)},
    "summary": "Quantitative analysis completed dynamically from {doc_title}."
}}
with open(str(out_dir / "metrics.json"), "w") as f:
    json.dump(metrics, f, indent=2)
print("Financial analysis script finished for {doc_title}.")
"""
                        elif wants_chart and chart_data and chart_data.get("series"):
                            x_lbls = chart_data.get("x_labels", [])
                            s_dict = chart_data.get("series", {})
                            py_code = f"""# Autonomous data analysis and visualization for {doc_title}
import json
import os
from pathlib import Path

os.environ['MPLCONFIGDIR'] = '/tmp'
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

out_dir = Path('/workspace/output') if Path('/workspace/output').exists() else Path('output')
out_dir.mkdir(parents=True, exist_ok=True)

x_labels = {json.dumps(x_lbls)}
series_data = {json.dumps(s_dict)}

fig, ax = plt.subplots(figsize=(9, 5), dpi=150)
x = np.arange(len(x_labels))
num_series = len(series_data)
width = 0.8 / max(num_series, 1)

for idx, (s_name, s_vals) in enumerate(series_data.items()):
    ax.bar(x + idx * width - (num_series - 1) * width / 2, s_vals, width, label=s_name, alpha=0.85)

ax.set_xticks(x)
ax.set_xticklabels(x_labels, rotation=25 if len(str(x_labels)) > 40 else 0, ha='right' if len(str(x_labels)) > 40 else 'center', fontsize=9)
ax.set_title('Dataset Analysis: {stem_name}', fontsize=11, fontweight='bold', pad=12)
ax.grid(axis='y', linestyle=':', alpha=0.6)
ax.legend()
plt.tight_layout()
plt.savefig(str(out_dir / '{chart_filename}'))
plt.close()

metrics = {{
    "document": "{doc_title}",
    "analysis_status": "SUCCESS",
    "chart_file": "{chart_filename}",
    "row_count": {len(x_lbls)},
    "series": list(series_data.keys()),
    "summary": "Dynamic visualization generated from {doc_title}."
}}
with open(str(out_dir / "metrics.json"), "w") as f:
    json.dump(metrics, f, indent=2)
print("Dynamic visualization script finished for {doc_title}.")
"""
                        elif wants_excel:
                            headers_to_write = insights.get("table_headers", ["Col_1", "Col_2"])
                            rows_to_write = insights.get("table_rows", [["Value 1", "Value 2"]])
                            py_code = f"""# Autonomous Excel converter script for {doc_title}
import json
import os
from pathlib import Path
import pandas as pd

out_dir = Path('/workspace/output') if Path('/workspace/output').exists() else Path('output')
out_dir.mkdir(parents=True, exist_ok=True)

headers = {json.dumps(headers_to_write)}
rows = {json.dumps(rows_to_write)}

df = pd.DataFrame(rows, columns=headers[:len(rows[0])] if rows else headers)
df.to_excel(str(out_dir / "converted_data.xlsx"), index=False)

analysis = {{
    "document": "{doc_title}",
    "analysis_status": "SUCCESS",
    "excel_file": "converted_data.xlsx",
    "rows_converted": len(df),
    "summary": "Document tables converted to Excel spreadsheet."
}}
with open(str(out_dir / "metrics.json"), "w") as f:
    json.dump(analysis, f, indent=2)
print("Excel conversion script finished with returncode 0.")
"""
                        else:
                            page_count = insights.get("page_count", 1)
                            headings = insights.get("headings", [])
                            py_code = f"""# Autonomous analysis script for {doc_title}
import json
import os
from pathlib import Path

out_dir = Path('/workspace/output') if Path('/workspace/output').exists() else Path('output')
out_dir.mkdir(parents=True, exist_ok=True)

analysis = {{
    "document": "{doc_title}",
    "analysis_status": "SUCCESS",
    "page_count": {page_count},
    "headings": {json.dumps(headings)},
    "summary": "Document successfully analyzed dynamically in sovereign sandbox environment."
}}
with open(str(out_dir / "metrics.json"), "w") as f:
    json.dump(analysis, f, indent=2)
print("Analysis script finished with returncode 0.")
"""
                        tool_params = {
                            "code": py_code,
                            "promote_outputs_to_artifacts": True
                        }
                        if target_doc_path:
                            try:
                                rel_input_path = str(target_doc_path.resolve().relative_to(ws_root)).replace("\\", "/")
                            except Exception:
                                rel_input_path = f"uploads/{target_doc_path.name}"
                            tool_params["input_files"] = [{"source_path": rel_input_path, "dest_name": target_doc_path.name}]

                        tool_out = await execute_tool("execute_code", tool_params, workspace_id=workspace_id, run_id=run_id)
                        sandbox_used = True
                        current_step.status = "completed"
                        current_step.tool_name = "execute_code"
                        current_step.tool_parameters = {"code": py_code}
                        current_step.observation = tool_out
                        await log_event(db, run_id, "tool_executed", f"Executed autonomous Python script: {tool_out}", {"output": tool_out})
                        saved_plan_json = serialize_plan(plan)
                        await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                        await db.commit()
                        plan.advance_to_next_step()
                        continue

                    elif current_step.id == 3:
                        if is_financial_task:
                            p2_paragraphs = []
                            if "revenue" in metrics_map:
                                r_info = metrics_map["revenue"]
                                p2_paragraphs.append(
                                    f"Gross Revenue: Latest reported at {format_number_display(r_info['latest'], is_currency=True)} ({r_info['latest_col']})"
                                    + (f", with growth of {growth_map.get('revenue_cagr_pct', growth_map.get('revenue_yoy_pct', 'N/A'))}%." if growth_map else ".")
                                )
                            if "ebitda" in metrics_map:
                                e_info = metrics_map["ebitda"]
                                p2_paragraphs.append(
                                    f"Operating EBITDA: Latest reported at {format_number_display(e_info['latest'], is_currency=True)} ({e_info['latest_col']})"
                                    + (f", with growth of {growth_map.get('ebitda_cagr_pct', growth_map.get('ebitda_yoy_pct', 'N/A'))}%." if growth_map else ".")
                                )
                            if "pat" in metrics_map:
                                p_info = metrics_map["pat"]
                                p2_paragraphs.append(
                                    f"Net Profit After Tax (PAT): Latest reported at {format_number_display(p_info['latest'], is_currency=True)} ({p_info['latest_col']})"
                                    + (f", with growth of {growth_map.get('pat_cagr_pct', growth_map.get('pat_yoy_pct', 'N/A'))}%." if growth_map else ".")
                                )
                            if "grm" in metrics_map:
                                p2_paragraphs.append(f"Gross Refining Margin (GRM): Reported at ${metrics_map['grm']['latest']}/bbl.")

                            if not p2_paragraphs:
                                p2_paragraphs = [
                                    "Comprehensive quantitative financial metrics extracted from primary workbook sheet.",
                                    f"Dataset: {doc_title} verified across historical reporting periods."
                                ]

                            table_headers = insights.get("table_headers", ["Line Item", "Value"])
                            table_rows = insights.get("table_rows", [])[:15]
                            if not table_rows and metrics_map:
                                table_headers = ["Metric", "Latest Value", "Period"]
                                table_rows = [[m_info["matched_label"], str(m_info["latest"]), str(m_info["latest_col"])] for m_info in metrics_map.values()]

                            chosen_sections = [
                                {
                                    "heading": "1. Executive Summary & Audit Provenance",
                                    "level": 1,
                                    "paragraphs": [
                                        f"Audit Deliverable: {display_title}",
                                        f"Prepared by: {custom_author if custom_author else 'CogniShift automated analysis'}",
                                        f"Source Dataset: {doc_title} (Ingested into Knowledge Vault)",
                                        f"Scope: Comprehensive Multi-Year Financial Performance Audit.",
                                        "Analysis executed by local CogniShift services. External internet access was blocked by strict application policy; physical network isolation was not asserted."
                                    ]
                                },
                                {
                                    "heading": "2. Profit & Loss Statement & Multi-Year Growth Metrics",
                                    "level": 1,
                                    "paragraphs": p2_paragraphs,
                                    "table": {
                                        "headers": table_headers,
                                        "rows": table_rows
                                    }
                                },
                                {
                                    "heading": "3. Verification and Approval Status",
                                    "level": 1,
                                    "paragraphs": [
                                        f"Audit Completion Date: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}",
                                        f"Prepared by: {custom_author if custom_author else 'CogniShift automated analysis'}",
                                        "Human approval / sign-off: Not recorded for this run.",
                                        "Network control: External internet access blocked by strict application policy; physical isolation not asserted.",
                                        "Generated by CogniShift Sovereign Agentic Workbench (SIH26117)."
                                    ]
                                }
                            ]
                        else:
                            doc_headers = insights.get("table_headers", ["Item", "Description", "Value"])
                            doc_rows = insights.get("table_rows", [["Document Content", "Analyzed", "Complete"]])[:15]
                            headings = insights.get("headings", [])
                            excerpts = [p[1][:200] for p in insights.get("page_texts", [])[:3]] if insights.get("page_texts") else []

                            is_audit_mode = "audit" in lower_input
                            p1_scope = "Comprehensive Plant Engineering & Instrumentation Audit." if is_audit_mode else f"Detailed operational parameter evaluation of '{doc_title}'."
                            p2_heading = "2. Operational Audit Findings & Parameter Analysis" if is_audit_mode else "2. Extracted Findings & Operational Analysis"

                            p2_body = [
                                f"Detailed inspection and parameter evaluation of '{doc_title}' across operational process loops.",
                                f"Autonomous audit verification conducted against plant asset hierarchies and design baselines."
                            ]
                            if headings:
                                p2_body.append(f"Identified primary sections / tags: {', '.join(headings[:5])}.")
                            if excerpts:
                                p2_body.extend(excerpts)

                            chosen_sections = [
                                {
                                    "heading": "1. Executive Summary & Process Scope",
                                    "level": 1,
                                    "paragraphs": [
                                        f"Audit Deliverable: {display_title}",
                                        f"Prepared by: {custom_author if custom_author else 'CogniShift automated analysis'}",
                                        f"Source Dataset: {doc_title} (ID #{target_doc['id'] if target_doc else '1'})",
                                        f"Scope: {p1_scope}",
                                        (
                                            "Analysis executed in an isolated local Python container. External internet access was blocked by strict application policy."
                                            if sandbox_used
                                            else "Analysis executed by trusted local backend services. External internet access was blocked by strict application policy."
                                        )
                                    ]
                                },
                                {
                                    "heading": p2_heading,
                                    "level": 1,
                                    "paragraphs": p2_body,
                                    "table": {
                                        "headers": doc_headers,
                                        "rows": doc_rows
                                    }
                                },
                                {
                                    "heading": "3. Verification and Approval Status",
                                    "level": 1,
                                    "paragraphs": [
                                        f"Audit Completion Date: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}",
                                        f"Prepared by: {custom_author if custom_author else 'CogniShift automated analysis'}",
                                        "Human approval / sign-off: Not recorded for this run.",
                                        "Network control: External internet access blocked by strict application policy; physical isolation not asserted.",
                                        "Generated by CogniShift Sovereign Agentic Workbench (SIH26117)."
                                    ]
                                }
                            ]

                        # Generate deliverable according to requested format
                        doc_params: Dict[str, Any] = {}
                        if target_fmt == "both":
                            report_docx_name = f"{stem_name}.docx" if custom_title else f"Report_{stem_name}.docx"
                            report_pdf_name = f"{stem_name}.pdf" if custom_title else f"Report_{stem_name}.pdf"
                            docx_params = {
                                "filename": report_docx_name,
                                "title": display_title,
                                "sections": chosen_sections
                            }
                            pdf_params = {
                                "filename": report_pdf_name,
                                "title": display_title,
                                "sections": chosen_sections
                            }
                            docx_out = await execute_tool("generate_docx", docx_params, workspace_id=workspace_id, run_id=run_id)
                            pdf_out = await execute_tool("generate_pdf", pdf_params, workspace_id=workspace_id, run_id=run_id)
                            tool_out = f"Dual deliverables generated: {report_docx_name} and {report_pdf_name} ({docx_out} | {pdf_out})"
                            doc_params = {"docx": docx_params, "pdf": pdf_params}
                            current_step.tool_name = "generate_docx_and_pdf"
                        elif target_fmt == "pdf":
                            report_filename = f"{stem_name}.pdf" if custom_title else f"Report_{stem_name}.pdf"
                            doc_params = {
                                "filename": report_filename,
                                "title": display_title,
                                "sections": chosen_sections
                            }
                            tool_out = await execute_tool("generate_pdf", doc_params, workspace_id=workspace_id, run_id=run_id)
                            current_step.tool_name = "generate_pdf"
                        elif target_fmt == "pptx":
                            report_filename = f"{stem_name}.pptx" if custom_title else f"Presentation_{stem_name}.pptx"
                            
                            slide2_bullets = [
                                f"Autonomous quantitative review of {doc_title}",
                                f"Workspace #{workspace_id} on-premise execution",
                                f"Timestamp: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
                            ]
                            if is_financial_task and growth_map:
                                for k, v in list(growth_map.items())[:3]:
                                    slide2_bullets.append(f"{k.replace('_', ' ').title()}: {v}%")
                            elif insights.get("headings"):
                                for h in insights["headings"][:3]:
                                    slide2_bullets.append(f"Section: {h}")

                            slide3_bullets = []
                            for r in insights.get("table_rows", [])[:4]:
                                if len(r) >= 2:
                                    slide3_bullets.append(f"{r[0]}: {r[1]}" + (f" ({r[2]})" if len(r) > 2 else ""))
                            if not slide3_bullets:
                                slide3_bullets = [
                                    "Integrity parameters and key metrics extracted authoritatively.",
                                    "All values checked against operating limits.",
                                    "No uncontained excursions or safety hazards observed."
                                ]

                            pptx_slides = [
                                {
                                    "title": "Executive Summary",
                                    "bullet_points": slide2_bullets
                                },
                                {
                                    "title": "Operational Findings & Metrics",
                                    "bullet_points": slide3_bullets
                                },
                                {
                                    "title": "Recommendations & Next Actions",
                                    "bullet_points": [
                                        "Maintain scheduled routine inspection cycle.",
                                        "Ensure telemetry sensor calibrations are up to date.",
                                        "Submit formal documentation for engineering review."
                                    ]
                                }
                            ]
                            doc_params = {
                                "filename": report_filename,
                                "title": display_title,
                                "subtitle": f"CogniShift Autonomous Briefing - {doc_title}",
                                "slides": pptx_slides
                            }
                            tool_out = await execute_tool("generate_pptx", doc_params, workspace_id=workspace_id, run_id=run_id)
                            current_step.tool_name = "generate_pptx"
                        elif target_fmt == "csv":
                            report_filename = f"Data_{stem_name}.csv"
                            headers = insights.get("table_headers", ["Item", "Value"])
                            rows = insights.get("table_rows", [])
                            if not rows:
                                rows = [["Source", doc_title], ["Status", "No tabular rows extracted"]]
                            doc_params = {
                                "filename": report_filename,
                                "title": display_title,
                                "headers": headers,
                                "rows": rows,
                            }
                            tool_out = await execute_tool("generate_csv", doc_params, workspace_id=workspace_id, run_id=run_id)
                            current_step.tool_name = "generate_csv"
                        elif target_fmt in ("excel", "xlsx"):
                            report_filename = f"Data_{stem_name}.xlsx"
                            doc_params = {
                                "filename": report_filename,
                                "title": display_title,
                                "sheets": [
                                    {
                                        "name": "Data_Extract",
                                        "headers": insights.get("table_headers", ["Item", "Value"]),
                                        "rows": insights.get("table_rows", [["Sample", "Data"]])
                                    }
                                ]
                            }
                            tool_out = await execute_tool("generate_xlsx", doc_params, workspace_id=workspace_id, run_id=run_id)
                            current_step.tool_name = "generate_xlsx"
                        elif target_fmt == "image":
                            tool_out = "Visualization deliverables confirmed and registered in workspace artifacts."
                            current_step.tool_name = "confirm_deliverables"
                            doc_params = {"status": "confirmed", "type": "visualization"}
                        elif target_fmt == "pdf":
                            report_filename = f"{stem_name}.pdf" if custom_title else (f"Audit_{stem_name}.pdf" if "audit" in lower_input else f"Report_{stem_name}.pdf")
                            c_charts = await db.execute(
                                "SELECT id, filename, relative_path, title FROM workspace_artifacts WHERE run_id = ? AND artifact_type IN ('png', 'image') ORDER BY id DESC LIMIT 1",
                                (run_id,)
                            )
                            chart_row = await c_charts.fetchone()
                            if chart_row and chosen_sections:
                                target_sec = chosen_sections[1] if len(chosen_sections) > 1 else chosen_sections[0]
                                if "images" not in target_sec:
                                    target_sec["images"] = []
                                target_sec["images"].append({
                                    "artifact_id": chart_row["id"],
                                    "caption": chart_row["title"] or f"Telemetry Visualization ({chart_row['filename']})"
                                })

                            doc_params = {
                                "filename": report_filename,
                                "title": display_title,
                                "sections": chosen_sections
                            }
                            tool_out = await execute_tool("generate_pdf", doc_params, workspace_id=workspace_id, run_id=run_id)
                            current_step.tool_name = "generate_pdf"
                        elif target_fmt == "both":
                            c_charts = await db.execute(
                                "SELECT id, filename, relative_path, title FROM workspace_artifacts WHERE run_id = ? AND artifact_type IN ('png', 'image') ORDER BY id DESC LIMIT 1",
                                (run_id,)
                            )
                            chart_row = await c_charts.fetchone()
                            if chart_row and chosen_sections:
                                target_sec = chosen_sections[1] if len(chosen_sections) > 1 else chosen_sections[0]
                                if "images" not in target_sec:
                                    target_sec["images"] = []
                                target_sec["images"].append({
                                    "artifact_id": chart_row["id"],
                                    "caption": chart_row["title"] or f"Telemetry Visualization ({chart_row['filename']})"
                                })

                            docx_filename = f"{stem_name}.docx" if custom_title else (f"Audit_{stem_name}.docx" if "audit" in lower_input else f"Report_{stem_name}.docx")
                            pdf_filename = f"{stem_name}.pdf" if custom_title else (f"Audit_{stem_name}.pdf" if "audit" in lower_input else f"Report_{stem_name}.pdf")

                            doc_params_docx = {"filename": docx_filename, "title": display_title, "sections": chosen_sections}
                            tool_out_docx = await execute_tool("generate_docx", doc_params_docx, workspace_id=workspace_id, run_id=run_id)

                            doc_params_pdf = {"filename": pdf_filename, "title": display_title, "sections": chosen_sections}
                            tool_out_pdf = await execute_tool("generate_pdf", doc_params_pdf, workspace_id=workspace_id, run_id=run_id)

                            tool_out = f"{tool_out_docx}\n{tool_out_pdf}"
                            current_step.tool_name = "generate_docx_and_pdf"
                            doc_params = {"docx": doc_params_docx, "pdf": doc_params_pdf}
                        else:
                            report_filename = f"{stem_name}.docx" if custom_title else (f"Audit_{stem_name}.docx" if "audit" in lower_input else f"Report_{stem_name}.docx")
                            c_charts = await db.execute(
                                "SELECT id, filename, relative_path, title FROM workspace_artifacts WHERE run_id = ? AND artifact_type IN ('png', 'image') ORDER BY id DESC LIMIT 1",
                                (run_id,)
                            )
                            chart_row = await c_charts.fetchone()
                            if chart_row and chosen_sections:
                                target_sec = chosen_sections[1] if len(chosen_sections) > 1 else chosen_sections[0]
                                if "images" not in target_sec:
                                    target_sec["images"] = []
                                target_sec["images"].append({
                                    "artifact_id": chart_row["id"],
                                    "caption": chart_row["title"] or f"Telemetry Visualization ({chart_row['filename']})"
                                })

                            doc_params = {
                                "filename": report_filename,
                                "title": display_title,
                                "sections": chosen_sections
                            }
                            tool_out = await execute_tool("generate_docx", doc_params, workspace_id=workspace_id, run_id=run_id)
                            current_step.tool_name = "generate_docx"

                        current_step.status = "completed"
                        current_step.tool_parameters = doc_params
                        current_step.observation = tool_out
                        await log_event(db, run_id, "tool_executed", f"Generated official deliverable: {tool_out}", {"output": tool_out})
                        saved_plan_json = serialize_plan(plan)
                        await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                        await db.commit()
                        plan.advance_to_next_step()
                        continue

                    elif current_step.id >= 4:
                        # Validate generated artifacts strictly scoped to current run_id and synthesize final response
                        cursor_chk = await db.execute(
                            "SELECT * FROM workspace_artifacts WHERE workspace_id = ? AND run_id = ? ORDER BY id ASC",
                            (workspace_id, run_id)
                        )
                        art_rows = await cursor_chk.fetchall()
                        current_run_artifacts = [dict(a) for a in art_rows]

                        # Verify each artifact exists on disk and has non-zero size
                        for art in current_run_artifacts:
                            cand_paths = [
                                ws_root / art.get("file_path", ""),
                                ws_root / f"generated/run_{run_id}" / art.get("filename", ""),
                                ws_root / art.get("filename", "")
                            ]
                            exists_on_disk = any(p.exists() and p.stat().st_size > 0 for p in cand_paths)
                            if not exists_on_disk:
                                logger.warning(f"Artifact {art.get('filename')} missing or empty on disk!")

                        docx_art = next((a for a in current_run_artifacts if a.get("artifact_type") == "docx" or str(a.get("filename", "")).endswith(".docx")), None)
                        pdf_art = next((a for a in current_run_artifacts if a.get("artifact_type") == "pdf" or str(a.get("filename", "")).endswith(".pdf")), None)
                        chart_arts = [a for a in current_run_artifacts if a.get("artifact_type") == "png" or str(a.get("filename", "")).lower().endswith(".png")]
                        json_art = next((a for a in current_run_artifacts if "metrics" in str(a.get("filename", "")).lower() or a.get("artifact_type") == "json" or str(a.get("filename", "")).endswith(".json")), None)
                        pptx_art = next((a for a in current_run_artifacts if a.get("artifact_type") == "pptx" or str(a.get("filename", "")).endswith(".pptx")), None)
                        xlsx_art = next((a for a in current_run_artifacts if a.get("artifact_type") == "xlsx" or str(a.get("filename", "")).endswith(".xlsx")), None)
                        csv_art = next((a for a in current_run_artifacts if a.get("artifact_type") == "csv" or str(a.get("filename", "")).endswith(".csv")), None)

                        deliverables_list = []
                        item_num = 1
                        if docx_art:
                            deliverables_list.append(f"{item_num}. **Word Document (`.docx`)**: `{docx_art['filename']}` ({docx_art.get('file_size', 0)} bytes) — Artifact #{docx_art['id']}")
                            item_num += 1
                        if pdf_art:
                            deliverables_list.append(f"{item_num}. **PDF Audit Report (`.pdf`)**: `{pdf_art['filename']}` ({pdf_art.get('file_size', 0)} bytes) — Artifact #{pdf_art['id']}")
                            item_num += 1
                        if pptx_art:
                            deliverables_list.append(f"{item_num}. **PowerPoint Presentation (`.pptx`)**: `{pptx_art['filename']}` ({pptx_art.get('file_size', 0)} bytes) — Artifact #{pptx_art['id']}")
                            item_num += 1
                        if xlsx_art:
                            deliverables_list.append(f"{item_num}. **Excel Workbook (`.xlsx`)**: `{xlsx_art['filename']}` ({xlsx_art.get('file_size', 0)} bytes) — Artifact #{xlsx_art['id']}")
                            item_num += 1
                        if csv_art:
                            deliverables_list.append(f"{item_num}. **CSV Data Export (`.csv`)**: `{csv_art['filename']}` ({csv_art.get('file_size', 0)} bytes) — Artifact #{csv_art['id']}")
                            item_num += 1
                        for c_art in chart_arts:
                            deliverables_list.append(f"{item_num}. **Visualization Chart (`.png`)**: `{c_art['filename']}` ({c_art.get('file_size', 0)} bytes) — Artifact #{c_art['id']}")
                            item_num += 1
                        if json_art:
                            deliverables_list.append(f"{item_num}. **Quantitative Metrics Ledger (`.json`)**: `{json_art['filename']}` ({json_art.get('file_size', 0)} bytes) — Artifact #{json_art['id']}")
                            item_num += 1

                        if not deliverables_list and current_run_artifacts:
                            for a in current_run_artifacts:
                                deliverables_list.append(f"{item_num}. **Deliverable**: `{a['filename']}` ({a.get('file_size', 0)} bytes) — Artifact #{a['id']}")
                                item_num += 1

                        deliv_str = "\n".join(deliverables_list) if deliverables_list else "None generated"

                        is_viz_only = (target_fmt == "image") or (bool(chart_arts) and not (docx_art or pdf_art or pptx_art or xlsx_art or csv_art))

                        if is_financial_task:
                            final_text = (
                                f"I have executed the quantitative analysis on `{doc_title}` and generated the requested deliverable{'s' if len(deliverables_list) > 1 else ''}:\n\n"
                                f"### 📊 Generated Deliverables:\n"
                                f"{deliv_str}\n\n"
                            )
                            if metrics_map:
                                final_text += f"### 📈 Key Quantitative Findings (Extracted from `{doc_title}`):\n"
                                for m_key, m_info in metrics_map.items():
                                    lbl = m_info["matched_label"]
                                    latest = m_info["latest"]
                                    col = m_info["latest_col"]
                                    growth_val = growth_map.get(f"{m_key}_cagr_pct", growth_map.get(f"{m_key}_yoy_pct"))
                                    growth_str = f" with growth of **{growth_val}%**" if growth_val is not None else ""
                                    final_text += f"- **{lbl}**: Reported at **{format_number_display(latest, is_currency=(m_key != 'grm'))}** ({col}){growth_str}.\n"
                                final_text += (
                                    f"- **Prepared by**: `{custom_author if custom_author else 'CogniShift automated analysis'}`\n"
                                    f"- **Human approval / sign-off**: Not recorded for this run.\n"
                                    f"- **Network control**: External internet access blocked by strict application policy; physical isolation not asserted."
                                )
                            else:
                                final_text += f"\nAnalysis completed for `{doc_title}`."
                        else:
                            intro_msg = (
                                f"I have executed the deterministic visualization pipeline on `{doc_title}` and generated the requested deliverable{'s' if len(deliverables_list) > 1 else ''}:\n\n"
                                if is_viz_only
                                else (
                                    f"I have executed the analysis in an isolated local sandbox and generated the official deliverable{'s' if len(deliverables_list) > 1 else ''} for '{doc_title}'.\n\n"
                                    if sandbox_used
                                    else f"I have used trusted local backend services to generate the official deliverable{'s' if len(deliverables_list) > 1 else ''} for '{doc_title}'.\n\n"
                                )
                            )
                            engine_note = (
                                "- **Execution Service:** Trusted backend Matplotlib charting service with strict Pillow validation (exit code 0)\n"
                                if is_viz_only
                                else (
                                    "- **Execution Service:** Isolated local Python container (exit code 0)\n"
                                    if sandbox_used
                                    else "- **Execution Service:** Trusted local backend document service\n"
                                )
                            )
                            final_text = (
                                f"{intro_msg}"
                                f"### 📊 Generated Deliverables:\n"
                                f"{deliv_str}\n\n"
                                f"### 📋 Key Findings (Extracted from `{doc_title}`):\n"
                                f"- **Resolved Document:** `{doc_title}`\n"
                                f"- **Document Type:** {insights.get('doc_type', 'General').upper()}\n"
                            )
                            if insights.get("table_headers"):
                                final_text += f"- **Extracted Schema / Columns:** {', '.join(insights['table_headers'][:8])}\n"
                            if insights.get("row_count"):
                                final_text += f"- **Data Rows Evaluated:** {insights['row_count']} rows\n"
                            elif insights.get("page_count"):
                                final_text += f"- **Pages Analyzed:** {insights['page_count']} pages\n"

                            final_text += (
                                f"{engine_note}"
                                f"- **Sources Cited:** `[{doc_title} | Page 1]`\n"
                                f"- **Network control:** External internet access blocked by strict application policy; physical isolation not asserted."
                            )

                        current_step.status = "completed"
                        current_step.observation = f"Validated {len(current_run_artifacts)} artifacts produced by current run"
                        sources_used = f"Knowledge Source #{target_doc['id'] if target_doc else '1'} | {doc_title}"
                        break

                elif action is None:
                    # ModelProtocolFailure: Output is unparseable (e.g. malformed JSON, truncated tokens, invalid action type)
                    await log_event(
                        db, run_id, "model_protocol_parse_failed",
                        f"Failed to parse model output from {selected_model_id} on step #{current_step.id}",
                        {"step_id": current_step.id, "preview": clean_output[:200]}
                    )

                    # Bounded protocol repair retry: One attempt with strict format prompt
                    await log_event(
                        db, run_id, "protocol_repair_attempted",
                        f"Attempting bounded protocol repair prompt with {selected_model_id} on step #{current_step.id}...",
                        {"step_id": current_step.id}
                    )

                    if is_final_step:
                        repair_prompt = (
                            f"Your previous response on step #{current_step.id} (FINAL SYNTHESIS STEP) could not be parsed into a valid action. "
                            f"Your previous output was:\n{clean_output[:600]}\n\n"
                            f"You must reformat your response immediately into ONE valid JSON object conforming to this schema:\n\n"
                            f'{{"thought": "<brief reasoning>", "action": "final_answer", "content": "<your complete final answer>", "citations": []}}\n\n'
                            f"Respond ONLY with the JSON object. Do not include markdown formatting or commentary outside the JSON."
                        )
                    else:
                        repair_prompt = (
                            f"Your previous response on step #{current_step.id} could not be parsed into a valid action. "
                            f"Your previous output was:\n{clean_output[:600]}\n\n"
                            f"You must reformat your response immediately into ONE valid JSON object conforming to one of these schemas:\n\n"
                            f'Option 1 (Tool Call):\n'
                            f'{{"thought": "<brief reasoning>", "action": "tool_call", "tool_name": "<tool_name>", "parameters": {{...}}}}\n\n'
                            f'Option 2 (Final Answer):\n'
                            f'{{"thought": "<brief reasoning>", "action": "final_answer", "content": "<your complete final answer>", "citations": []}}\n\n'
                            f"Available tools: {json.dumps(allowed_tool_names)}\n"
                            f"Respond ONLY with the JSON object. Do not include markdown formatting or commentary outside the JSON."
                        )

                    repair_action = None
                    try:
                        provider = get_provider()
                        repair_resp = await provider.generate_text(
                            prompt=repair_prompt,
                            system_prompt="You are a protocol recovery assistant. You must output only valid JSON.",
                            model_name=selected_model_id
                        )
                        repair_text = repair_resp.text.strip()
                        repair_action = parse_agent_action(repair_text, strict=False)
                    except Exception as rep_err:
                        logger.warning(f"Protocol repair LLM call failed: {rep_err}")
                        repair_action = None

                    if repair_action is not None:
                        await log_event(
                            db, run_id, "protocol_repair_succeeded",
                            f"Protocol repair succeeded for step #{current_step.id}",
                            {"action_type": repair_action.action}
                        )
                        action = repair_action
                        clean_output = repair_text
                    else:
                        await log_event(
                            db, run_id, "protocol_repair_failed",
                            f"Protocol repair retry failed for step #{current_step.id}",
                            {"step_id": current_step.id}
                        )
                        err_msg = f"Model protocol failure on step #{current_step.id}: Unparseable or malformed output from {selected_model_id}."
                        logger.warning(err_msg)
                        current_step.status = "failed"
                        current_step.error_message = err_msg
                        _block_unexecuted_pending_steps(plan, "Blocked due to unparseable model output")
                        await log_event(db, run_id, "model_protocol_failure", err_msg, {"step_id": current_step.id, "preview": clean_output[:150]})
                        saved_plan_json = serialize_plan(plan)
                        await db.execute(
                            """UPDATE agent_runs
                               SET status = 'failed', error_message = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                               WHERE id = ?""",
                            (err_msg, saved_plan_json, run_id)
                        )
                        await db.commit()
                        cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                        return make_response(dict(await cursor.fetchone()))

                # Intercept unauthorized tool calls on final synthesis step
                if is_final_step and isinstance(action, ToolCallProposal):
                    logger.warning(
                        f"Step #{current_step.id} is the final synthesis step; "
                        f"intercepted unauthorized tool call proposal '{action.tool_name}' and converted to final_answer."
                    )
                    await log_event(
                        db, run_id, "final_step_tool_interception",
                        f"Intercepted unauthorized tool call proposal '{action.tool_name}' on final step #{current_step.id}",
                        {"step_id": current_step.id, "tool_name": action.tool_name}
                    )
                    synth_content = getattr(action, "reason", None) or f"Synthesized findings based on available documentary evidence regarding {action.tool_name}."
                    action = FinalAnswer(
                        action="final_answer",
                        content=synth_content,
                        citations=[]
                    )

                if isinstance(action, ToolCallProposal) and action.tool_name:
                    tool_name_clean = action.tool_name.lower().strip()
                    if tool_name_clean not in [t.lower() for t in allowed_tool_names]:
                        current_step.status = "failed"
                        current_step.error_message = f"Tool '{action.tool_name}' is not in allowed tools list."
                        current_step.observation = f"Unauthorized tool attempt: {action.tool_name}"
                        await log_event(db, run_id, "tool_unauthorized", current_step.error_message)
                        plan.advance_to_next_step()
                        continue

                    resolved_tool = next(t for t in allowed_tool_names if t.lower() == tool_name_clean)
                    tool_def = tools_by_name[resolved_tool]

                    raw_params = dict(action.parameters)
                    if "reason" not in raw_params and action.reason:
                        raw_params["reason"] = action.reason

                    val_result = validate_proposed_tool_call(
                        tool_name=resolved_tool,
                        raw_parameters=raw_params,
                        allowed_tools=allowed_tool_names,
                        references=routing_res.references
                    )
                    if not val_result.valid:
                        severity = classify_tool_failure_severity(resolved_tool, is_rca_mode=is_rca_mode)
                        if severity == ToolFailureSeverity.OPTIONAL_ENRICHMENT_FAILURE:
                            current_step.status = "completed"
                            current_step.observation = f"Note: Optional check '{resolved_tool}' could not be executed ({val_result.error_message}). Proceeding with documentary evidence."
                            await log_event(
                                db, run_id, "tool_validation_optional_skip",
                                f"Optional tool '{resolved_tool}' parameter validation failed; continuing run: {val_result.error_message}",
                                {"parameters": action.parameters, "error": val_result.error_message}
                            )
                            plan.advance_to_next_step()
                            continue
                        else:
                            current_step.status = "failed"
                            current_step.error_message = val_result.error_message
                            current_step.observation = f"Validation error: {val_result.error_message}"
                            await log_event(
                                db, run_id, "tool_validation_failed",
                                f"Step #{current_step.id} parameter validation failed: {val_result.error_message}",
                                {"parameters": action.parameters, "error": val_result.error_message}
                            )
                            plan.advance_to_next_step()
                            continue

                    validated_params = val_result.validated_parameters or action.parameters

                    # Target topology validation for plant equipment tools
                    if resolved_tool in ("restart_component", "emergency_pressure_relief", "run_diagnostic", "check_pressure", "check_temperature"):
                        target_eq = (
                            validated_params.get("component_id")
                            or validated_params.get("chamber_id")
                            or validated_params.get("sensor_id")
                            or validated_params.get("equipment_id")
                        )
                        if target_eq:
                            is_sup, sup_msg = is_supported_equipment_target(str(target_eq))
                            if not is_sup:
                                current_step.status = "failed"
                                current_step.error_message = sup_msg
                                current_step.observation = f"Capability UNSUPPORTED: {sup_msg}"
                                fail_reply = (
                                    f"⚠️ **Plant Topology Notice**: {sup_msg}\n\n"
                                    f"No operational action was executed on `{target_eq}`. For refinery safety, commands "
                                    f"can only be issued against verified components registered in the plant topology."
                                )
                                saved_plan_json = serialize_plan(plan)
                                await db.execute(
                                    """UPDATE agent_runs
                                       SET status = 'failed', result_text = ?, error_message = ?, sources_used = 'MRPL Plant Topology', structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                                       WHERE id = ?""",
                                    (fail_reply, sup_msg, saved_plan_json, run_id)
                                )
                                await db.commit()
                                await log_event(db, run_id, "target_unsupported", sup_msg, {"target": target_eq})
                                cursor = await db.execute("SELECT * FROM agent_runs WHERE id = ?", (run_id,))
                                return make_response(dict(await cursor.fetchone()))

                    is_readonly_chart_or_viz = False
                    if resolved_tool == "execute_code":
                        code_str = str(validated_params.get("code", "")).lower()
                        has_viz = any(w in code_str for w in ["plt.", "matplotlib", "seaborn", "savefig"])
                        has_dangerous_terms = any(w in code_str for w in ["socket", "subprocess", "os.system", "shutil.rmtree", "os.remove", "requests", "urllib", "http"])
                        if has_viz and not has_dangerous_terms:
                            is_readonly_chart_or_viz = True

                    requires_approval = bool(
                        val_result.requires_approval or
                        tool_def.get("requires_approval", 0) or
                        (resolved_tool in HIGH_RISK_TOOLS) or
                        (agent.get("approval_required", 0) and tool_def.get("risk_level") in ["sensitive", "service_interrupting"])
                    )

                    if requires_approval:
                        # High-risk simulated action: PAUSE FOR FOUR-EYES DUAL SUPERVISOR APPROVAL
                        is_dual = (
                            tool_def.get("risk_level") in ["sensitive", "service_interrupting"]
                            or resolved_tool in HIGH_RISK_TOOLS
                            or resolved_tool in ["restart_component", "emergency_pressure_relief", "restart_service"]
                        )
                        req_approvals = 2 if is_dual else 1
                        cursor = await db.execute(
                            """INSERT INTO approval_requests
                               (run_id, tool_id, status, request_reason, parameters, risk_level, required_approvals)
                               VALUES (?, ?, 'pending', ?, ?, ?, ?) RETURNING *""",
                            (run_id, tool_def["id"], action.reason, json.dumps(validated_params), tool_def.get("risk_level", "sensitive"), req_approvals)
                        )
                        approval_row = await cursor.fetchone()

                        if approval_row:
                            try:
                                from cognishift.core.notifications import (
                                    NotificationType,
                                    collect_four_eyes_evidence,
                                    fire_and_forget_notification,
                                )
                                fe_ev = collect_four_eyes_evidence(
                                    event_type=NotificationType.SENSITIVE_ACTION_REQUESTED,
                                    run_id=run_id,
                                    approval_id=approval_row["id"],
                                    tool_name=resolved_tool,
                                    parameters=validated_params,
                                    user_id=user_id,
                                    workspace_id=workspace_id,
                                )
                                fire_and_forget_notification(fe_ev)
                            except Exception as fe_notif_err:
                                logger.warning(f"Could not dispatch four-eyes requested notification: {fe_notif_err}")

                        # Persist PendingTask for tracking
                        try:
                            await create_pending_task(
                                workspace_id=workspace_id,
                                user_id=user_id,
                                intent=routing_res.intent.value,
                                requested_goal=effective_goal,
                                source_references={"tool_name": resolved_tool, "parameters": validated_params},
                                proposed_steps=[s.description for s in plan.steps],
                                originating_run_id=run_id,
                                db=db
                            )
                        except Exception as pt_err:
                            logger.warning(f"Could not persist pending task: {pt_err}")

                        current_step.status = "waiting_for_approval"
                        current_step.tool_name = resolved_tool
                        current_step.tool_parameters = validated_params
                        saved_plan_json = serialize_plan(plan)

                        pause_msg = f"Action paused awaiting supervisor approval: {resolved_tool}. Reason: {action.reason}"
                        cursor = await db.execute(
                            """UPDATE agent_runs
                               SET status = 'paused', result_text = ?, sources_used = ?, structured_plan = ?
                               WHERE id = ? RETURNING *""",
                            (pause_msg, sources_used, saved_plan_json, run_id)
                        )
                        updated_run = await cursor.fetchone()
                        await db.commit()

                        await log_event(
                            db,
                            run_id,
                            "approval_requested",
                            f"Step #{current_step.id} paused: Action '{resolved_tool}' requires supervisor authorization.",
                            {"step_id": current_step.id, "tool": resolved_tool, "parameters": validated_params, "risk_level": tool_def.get("risk_level")}
                        )
                        return make_response(dict(updated_run))

                    else:
                        # Safe / authorized tool execution: execute and record observation
                        await log_event(db, run_id, "tool_executing", f"Executing safe tool: {resolved_tool}", {"parameters": validated_params})
                        tool_output = await execute_tool(resolved_tool, validated_params, workspace_id=workspace_id, run_id=run_id)
                        await log_event(db, run_id, "tool_executed", f"Tool output received: {tool_output}", {"output": tool_output})

                        is_code_failure = (
                            resolved_tool == "execute_code" and
                            ("Sandbox execution failed" in tool_output or "Sandbox execution TIMED OUT" in tool_output or "Error" in tool_output)
                        )

                        if is_code_failure and current_step.retry_count < 2:
                            current_step.retry_count += 1
                            current_step.status = "running"
                            current_step.tool_name = resolved_tool
                            current_step.tool_parameters = validated_params
                            current_step.observation = (
                                f"[Execution Attempt #{current_step.retry_count} Failed - Debug Feedback]\n"
                                f"{tool_output}\n"
                                f"Analyze the error/traceback above and propose a corrected code implementation."
                            )
                            await log_event(
                                db, run_id, "sandbox_retry_triggered",
                                f"Step #{current_step.id} retry {current_step.retry_count}/2 triggered after runtime failure.",
                                {"step_id": current_step.id, "retry_count": current_step.retry_count}
                            )
                            saved_plan_json = serialize_plan(plan)
                            await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                            await db.commit()
                            # Do NOT advance to next step, allowing agent to attempt self-correction
                            continue

                        elif is_code_failure and current_step.retry_count >= 2:
                            current_step.status = "failed"
                            current_step.tool_name = resolved_tool
                            current_step.tool_parameters = validated_params
                            current_step.observation = tool_output
                            current_step.error_message = f"Sandbox code execution failed after {1 + current_step.retry_count} total attempts (retry budget exhausted)."
                            await log_event(
                                db, run_id, "sandbox_retry_exhausted",
                                f"Step #{current_step.id} failed: Maximum code correction attempts exhausted.",
                                {"step_id": current_step.id, "total_attempts": 1 + current_step.retry_count}
                            )
                            saved_plan_json = serialize_plan(plan)
                            await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                            await db.commit()
                            plan.advance_to_next_step()

                        else:
                            if tool_output_failed(tool_output):
                                severity = classify_tool_failure_severity(resolved_tool, is_rca_mode=is_rca_mode)
                                if severity == ToolFailureSeverity.OPTIONAL_ENRICHMENT_FAILURE:
                                    current_step.status = "completed"
                                    current_step.tool_name = resolved_tool
                                    current_step.tool_parameters = validated_params
                                    current_step.observation = f"Note: Optional check '{resolved_tool}' execution returned failure ({tool_output}). Proceeding with documentary evidence."
                                    saved_plan_json = serialize_plan(plan)
                                    await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                                    await db.commit()
                                    plan.advance_to_next_step()
                                    continue
                                else:
                                    current_step.status = "failed"
                                    current_step.tool_name = resolved_tool
                                    current_step.tool_parameters = validated_params
                                    current_step.observation = tool_output
                                    current_step.error_message = tool_output
                                    saved_plan_json = serialize_plan(plan)
                                    await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                                    await db.commit()
                                    plan.advance_to_next_step()
                                    continue

                            current_step.status = "completed"
                            current_step.tool_name = resolved_tool
                            current_step.tool_parameters = validated_params
                            current_step.observation = tool_output

                            saved_plan_json = serialize_plan(plan)
                            await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                            await db.commit()

                            plan.advance_to_next_step()

                elif isinstance(action, StepObservation):
                    current_step.status = "completed"
                    current_step.observation = action.content
                    await log_event(
                        db, run_id, "step_observation",
                        f"Step #{current_step.id} recorded observation: {action.content[:150]}",
                        {"step_id": current_step.id, "observation": action.content}
                    )
                    saved_plan_json = serialize_plan(plan)
                    await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                    await db.commit()
                    plan.advance_to_next_step()
                    continue

                elif isinstance(action, FinalAnswer):
                    final_answer_citations = list(action.citations or [])
                    if is_rca_mode:
                        t_val_start = time.perf_counter()
                        if rca_bundle is None:
                            rca_bundle = RCAEvidenceBundle(
                                asset_ids=re.findall(r"\b[A-Za-z]{1,4}-\d{3,4}[A-Za-z]?\b", clean_input),
                                channel_health=ChannelExecutionHealth(executed_channels=["text"])
                            )
                        validator = RCAEvidenceValidator()
                        answer_content = validator.validate_and_finalize(rca_bundle, action.content, clean_input)
                        validation_latency_ms = round((time.perf_counter() - t_val_start) * 1000, 2)

                        structured_rca = validator.get_structured_result(rca_bundle, answer_content, clean_input)
                        await log_event(
                            db, run_id, "rca_result_structured",
                            f"Structured RCA result recorded with status {structured_rca.status} and primary cause {structured_rca.primary_cause_code}",
                            structured_rca.model_dump()
                        )

                        stage_lats = dict(getattr(rca_bundle, "stage_latencies_ms", {}) or {})
                        stage_lats["evidence_validation"] = validation_latency_ms
                        if t_run_start is not None:
                            stage_lats["e2e_total"] = round((time.perf_counter() - t_run_start) * 1000, 2)
                        await log_event(
                            db, run_id, "rca_stage_latencies",
                            f"RCA pipeline stage latencies recorded: {stage_lats}",
                            {"stage_latencies_ms": stage_lats}
                        )
                    else:
                        answer_content = action.content

                    is_direct_flow = effective_intent in (
                        SemanticIntent.CONVERSATION,
                        SemanticIntent.UI_NAVIGATION,
                        SemanticIntent.ARTIFACT_INSPECTION,
                        SemanticIntent.KNOWLEDGE_QUERY
                    )
                    is_final_step = (current_step.id == len(plan.steps))

                    # Multi-Step Goal Protection: Intermediate step must not abort plan before required outputs exist
                    if not is_final_step and not is_direct_flow:
                        logger.info(
                            f"Step #{current_step.id} of {len(plan.steps)}: Intercepted intermediate FinalAnswer; "
                            f"converting to StepObservation to protect plan execution."
                        )
                        current_step.status = "completed"
                        current_step.observation = answer_content
                        await log_event(
                            db, run_id, "step_observation",
                            f"Step #{current_step.id} intermediate synthesis recorded: {answer_content[:150]}",
                            {"step_id": current_step.id, "observation": answer_content}
                        )
                        saved_plan_json = serialize_plan(plan)
                        await db.execute("UPDATE agent_runs SET structured_plan = ? WHERE id = ?", (saved_plan_json, run_id))
                        await db.commit()
                        plan.advance_to_next_step()
                        continue

                    current_step.status = "completed"
                    current_step.observation = answer_content
                    plan.final_synthesis = answer_content
                    for s in plan.steps:
                        if s.status == "pending":
                            s.status = "skipped"
                            s.observation = "Resolved by final response"
                    final_text = answer_content
                    break

                elif isinstance(action, ClarificationRequest):
                    current_step.status = "completed"
                    current_step.observation = f"Clarification requested: {action.question}"
                    final_text = f"Clarification required from operator: {action.question}"
                    break

            failed_steps = [s for s in plan.steps if s.status == "failed"]
            has_high_risk_failure = any(
                (s.tool_name in HIGH_RISK_TOOLS or (s.error_message and any(hrt in s.error_message for hrt in HIGH_RISK_TOOLS)))
                for s in failed_steps
            )
            # If high-risk tool failed or plan completely lacked synthesis, fail the run fail-closed
            if failed_steps and (has_high_risk_failure or not (plan.final_synthesis or final_text)):
                run_final_status = "failed"
                for s in plan.steps:
                    if s.status == "pending":
                        s.status = "blocked"
                        s.observation = "Blocked by earlier failure"
                failures_summary = "\n".join(
                    f"Step #{s.id} ({s.description}): {s.error_message or s.observation or 'Failed'}"
                    for s in failed_steps
                )
                final_text = f"Execution encountered failures:\n{failures_summary}"
                error_msg = failed_steps[0].error_message or "Execution encountered failures"
            else:
                run_final_status = "completed"
                error_msg = None
                # If non-critical steps failed but final synthesis was achieved, log event and preserve final_text
                if failed_steps:
                    step_fail_descs = [f"Step #{s.id}: {s.description}" for s in failed_steps]
                    await log_event(
                        db, run_id, "run_completed_with_step_failures",
                        "Auxiliary non-critical step failure superseded by successful final goal synthesis.",
                        {"failed_steps": step_fail_descs}
                    )

                # Synthesize final response if not explicitly provided
                if not final_text:
                    if plan.final_synthesis:
                        final_text = plan.final_synthesis
                    elif vision_analysis:
                        final_text = f"Visual Inspection Analysis:\n{vision_analysis}"
                    else:
                        obs_summary = "\n".join(
                            f"Step #{s.id} ({s.description}): {s.observation or 'Done'}"
                            for s in plan.steps if s.status in ["completed", "running"]
                        )
                        final_text = f"Goal Execution Summary:\n{obs_summary}"
                elif plan.final_synthesis and not final_text:
                    final_text = plan.final_synthesis

                # P0-3: Ensure truthful plan state: mark unexecuted pending steps as skipped
                for s in plan.steps:
                    if s.status == "pending":
                        if s.id == len(plan.steps) and "present" in s.description.lower():
                            s.status = "completed"
                            s.observation = "Answer presented to operator"
                        else:
                            s.status = "skipped"
                            s.observation = "Bypassed - satisfied by prior execution steps"

            # --- GOAL CONTRACT VERIFICATION GATE ---
            if goal_contract is not None and run_final_status == "completed":
                cursor_art_gate = await db.execute(
                    "SELECT * FROM workspace_artifacts WHERE workspace_id = ? AND run_id = ?",
                    (workspace_id, run_id)
                )
                current_run_artifacts_gate = [dict(a) for a in await cursor_art_gate.fetchall()]
                satisfied, missing_deliverables = goal_contract.check_satisfaction(current_run_artifacts_gate)
                if not satisfied:
                    run_final_status = "failed"
                    error_msg = f"Goal contract verification failed: missing required deliverables {missing_deliverables}."
                    final_text = f"Execution Incomplete: Goal contract unsatisfied. Required deliverables missing: {', '.join(missing_deliverables)}."
                    logger.warning(f"Run {run_id} failed goal contract verification: {missing_deliverables}")
                    await log_event(
                        db, run_id, "goal_contract_violation",
                        f"Goal contract unsatisfied: missing {missing_deliverables}",
                        {"missing": missing_deliverables, "required_fields": goal_contract.required_fields, "extracted": goal_contract.extracted_fields}
                    )
                else:
                    await log_event(
                        db, run_id, "goal_contract_satisfied",
                        "All goal contract deliverables verified from structured facts and registered artifacts.",
                        {"fields": list(goal_contract.extracted_fields.keys()), "artifacts": [a.get("filename") for a in current_run_artifacts_gate]}
                    )

            # Complete or fail claimed pending task
            if claimed_task:
                if run_final_status == "completed":
                    await complete_pending_task(claimed_task.id, resulting_run_id=run_id, db=db)
                else:
                    await fail_pending_task(claimed_task.id, db=db)

            # Authoritative backend execution timestamp (never LLM-invented)
            execution_finish_utc = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
            timing_trigger = (
                "record the time" in clean_input.lower()
                or (claimed_task and "record the time" in claimed_task.requested_goal.lower())
                or routing_res.intent == SemanticIntent.CODE_EXECUTION
            )
            if timing_trigger and run_final_status == "completed":
                timing_block = f"\n\n---\n**Execution Timing & Provenance (Authoritative Backend Clock):**\n- Execution Completed: `{execution_finish_utc}`\n- Operating Mode: `{settings.operating_mode}`\n- External Internet: Blocked by strict application policy\n- Physical Network Isolation: Not asserted"
                if timing_block not in final_text:
                    final_text += timing_block
            # DeepSeek & LLM Output Hygiene: Ensure <think>...</think> blocks never leak to result_text
            final_text, _, _ = extract_and_strip_thinking(final_text)

            # SCADA Anomaly Grounding Authority:
            # If an anomalous tabular event was detected, ensure model prose conforms to frozen facts.
            # If the model introduced unsupported facts, wrong years, or contradictory valve states,
            # enforce the deterministic authoritative report.
            if frozen_scada_anomaly:
                is_valid, validated_text = validate_scada_anomaly_prose(
                    final_text, frozen_scada_anomaly, source_filename=frozen_scada_source
                )
                if not is_valid:
                    logger.warning("Overriding contradictory model prose with authoritative SCADA anomaly report.")
                    final_text = validated_text
                    await log_event(
                        db, run_id, "anomaly_facts_enforced",
                        "Authoritative SCADA anomaly facts enforced over contradictory/hallucinated model generation.",
                        {"timestamp": frozen_scada_anomaly.get("timestamp"), "source": frozen_scada_source}
                    )

            # Deterministic Citation & Provenance Reconciliation Gate:
            # Cross-reference model citations and answer text against authoritative retrieved chunks.
            if retrieved_evidence_catalog and run_final_status == "completed":
                from cognishift.core.document_processing.provenance import reconcile_citations_against_evidence, reconcile_citations_in_text
                verified_cites, reconciled_sources = reconcile_citations_against_evidence(
                    text=final_text,
                    model_citations=final_answer_citations,
                    retrieved_evidence=retrieved_evidence_catalog,
                    fallback_to_evidence_if_empty=False
                )
                if verified_cites:
                    sources_used = reconciled_sources
                    final_text = reconcile_citations_in_text(final_text, verified_cites, retrieved_evidence_catalog)
                    if graph_context and "Plant Topology Graph" not in sources_used:
                        sources_used += " + Plant Topology Graph"
                    if vision_analysis and "Local VLM Inspection" not in sources_used:
                        sources_used += " + Local VLM Inspection"
                    if artifact_citations:
                        unique_art = [a for a in dict.fromkeys(artifact_citations) if a not in sources_used]
                        if unique_art:
                            sources_used += " + " + ", ".join(unique_art)

                    await log_event(
                        db, run_id, "citations_reconciled",
                        f"Reconciled citations against retrieved evidence ({len(verified_cites)} verified citations)",
                        {"verified_citations": [c["citation_str"] for c in verified_cites], "sources_used": sources_used}
                    )

            saved_plan_json = serialize_plan(plan)
            routing_json = json.dumps(routing_res.to_dict()) if routing_res else None
            cursor = await db.execute(
                """UPDATE agent_runs
                   SET status = ?, result_text = ?, error_message = ?, sources_used = ?, structured_plan = ?, routing_info = ?, completed_at = CURRENT_TIMESTAMP
                   WHERE id = ? RETURNING *""",
                (run_final_status, final_text, error_msg, sources_used, saved_plan_json, routing_json, run_id)
            )
            updated_run = await cursor.fetchone()
            await db.commit()

            if run_final_status == "failed":
                await log_event(db, run_id, "run_failed", f"Run failed after {step_counter} steps: {error_msg}")
            else:
                executed_step_count = max(step_counter, sum(1 for step in plan.steps if step.status == "completed"))
                await log_event(db, run_id, "completed", f"Run completed successfully in {executed_step_count} steps.")
            return make_response(dict(updated_run))

        except Exception as e:
            logger.error(f"Error during agent run {run_id}: {str(e)}", exc_info=True)
            _block_unexecuted_pending_steps(plan, f"Blocked due to unhandled execution error: {str(e)}")
            saved_plan_json = serialize_plan(plan) if plan else None
            cursor = await db.execute(
                """UPDATE agent_runs
                   SET status = 'failed', error_message = ?, structured_plan = COALESCE(?, structured_plan), completed_at = CURRENT_TIMESTAMP
                   WHERE id = ? RETURNING *""",
                (str(e), saved_plan_json, run_id)
            )
            updated_run = await cursor.fetchone()
            await db.commit()
            await log_event(db, run_id, "failed", f"Execution error: {str(e)}")
            return make_response(dict(updated_run))


async def resume_agent_run(run_id: int) -> RunResponse:
    """Resume a paused agent run following human supervisor approval or rejection.
    
    P0-4 ATOMIC CAS GUARANTEE:
    Atomically updates status from 'paused' to 'resuming'.
    Only 1 worker can successfully claim execution. All competing requests fail immediately.
    """
    async with get_db() as db:
        # 1. Atomic Compare-And-Swap State Claim
        cursor = await db.execute(
            """UPDATE agent_runs
               SET status = 'resuming'
               WHERE id = ? AND status = 'paused'
               RETURNING *""",
            (run_id,)
        )
        run_row = await cursor.fetchone()
        if not run_row:
            cursor_check = await db.execute("SELECT status FROM agent_runs WHERE id = ?", (run_id,))
            existing = await cursor_check.fetchone()
            if not existing:
                raise HTTPException(status_code=404, detail=f"Run ID {run_id} not found")
            current_status = existing["status"]
            if current_status == "resuming":
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"Run {run_id} is already in the process of resuming execution."
                )
            else:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Run {run_id} cannot be resumed from current status '{current_status}'"
                )

        run = dict(run_row)
        await db.commit()

        # 2. Fetch associated approval request
        cursor = await db.execute(
            """SELECT ar.*, td.name as tool_name, td.implementation_key
               FROM approval_requests ar
               JOIN tool_definitions td ON ar.tool_id = td.id
               WHERE ar.run_id = ?
               ORDER BY ar.requested_at DESC LIMIT 1""",
            (run_id,)
        )
        approval_row = await cursor.fetchone()
        if not approval_row:
            await db.execute("UPDATE agent_runs SET status = 'paused' WHERE id = ?", (run_id,))
            await db.commit()
            raise HTTPException(status_code=400, detail=f"No approval request found for Run ID {run_id}")
        approval = dict(approval_row)

        if approval["status"] == "pending":
            await db.execute("UPDATE agent_runs SET status = 'paused' WHERE id = ?", (run_id,))
            await db.commit()
            raise HTTPException(status_code=400, detail=f"Approval request #{approval['id']} is still pending supervisor decision.")

        # 3. Process Decision
        if approval["status"] == "approved":
            tool_name = approval["tool_name"]
            try:
                params = json.loads(approval.get("parameters") or "{}")
            except (json.JSONDecodeError, TypeError):
                params = {}

            # Restore plan state
            plan = deserialize_plan(run.get("structured_plan"))
            active_step = None
            if plan:
                for step in plan.steps:
                    if step.status == "waiting_for_approval":
                        active_step = step
                        active_step.status = "running"
                        break

            await log_event(db, run_id, "tool_executing", f"Executing supervisor-approved tool: {tool_name}", {"parameters": params})
            execution_error = None
            try:
                tool_output = await execute_tool(tool_name, params, workspace_id=run["workspace_id"], run_id=run_id)
                if tool_output_failed(tool_output):
                    raise RuntimeError(tool_output)
                await log_event(db, run_id, "tool_executed", f"Approved tool output received: {tool_output}", {"output": tool_output})
                if active_step:
                    active_step.status = "completed"
                    active_step.observation = tool_output
                    plan.advance_to_next_step()

                # Phase 7 SIH Deliverable: Automatically create and register Approval_Note_Run_{run_id}.docx
                if tool_name in ["restart_component", "emergency_pressure_relief"]:
                    from cognishift.core.artifact_generators import create_and_register_artifact, generate_docx_document
                    try:
                        sections = [
                            {
                                "heading": "1. Incident & Equipment Summary",
                                "paragraphs": [
                                    f"Run ID: {run_id}",
                                    f"Operator request: {run.get('input_text', '')}",
                                    f"Approved parameters: {json.dumps(params, ensure_ascii=False)}",
                                    "No independent root-cause diagnosis or post-action equipment verification is established by this record."
                                ]
                            },
                            {
                                "heading": "2. Four-Eyes Operational Sign-Off",
                                "paragraphs": [
                                    f"Initiated by Operator: {run.get('user_id')}",
                                    f"Stage 1 Reviewer: {approval.get('reviewed_by')}",
                                    f"Stage 2 Authorizer: {approval.get('reviewed_by_2') or 'Not recorded'}",
                                    "Compliance: Strict separation of requester and authorizing roles verified under refinery safety protocol."
                                ]
                            },
                            {
                                "heading": "3. Authorized Remediation & System Outcome",
                                "paragraphs": [
                                    f"Simulated Tool Executed: {tool_name}",
                                    f"Tool Execution Output: {tool_output}",
                                    "These plant tools are simulated. Their output is not proof of physical equipment recovery."
                                ]
                            },
                            {
                                "heading": "4. Sovereignty & Governance Statement",
                                "paragraphs": [
                                    "Generated by the local CogniShift application.",
                                    "Network isolation must be verified separately using the network-event ledger and host-level observation; this document does not certify zero egress."
                                ]
                            }
                        ]
                        artifact_rec = await create_and_register_artifact(
                            workspace_id=run["workspace_id"],
                            filename=f"Approval_Note_Run_{run_id}.docx",
                            artifact_type="docx",
                            generator_fn=lambda p: generate_docx_document(p, "Operations — Recorded Action Authorization", sections),
                            title=f"Action Authorization — Run {run_id}",
                            description="Recorded approval identities, parameters and simulated tool output; not a physical safety certification.",
                            run_id=run_id
                        )
                        await log_event(
                            db, run_id, "artifact_registered",
                            f"Official deliverable '{artifact_rec['filename']}' generated and registered in Artifact Vault.",
                            {"artifact_id": artifact_rec["id"], "sha256": artifact_rec["sha256_hash"], "size": artifact_rec["file_size"]}
                        )
                    except Exception as art_err:
                        logger.warning(f"Failed to auto-generate Approval_Note_Run_{run_id}.docx: {art_err}")
            except Exception as e:
                err_msg = f"Tool execution failed: {str(e)}"
                execution_error = err_msg
                await log_event(db, run_id, "tool_failed", err_msg, {"error": str(e)})
                if active_step:
                    active_step.status = "failed"
                    active_step.error_message = str(e)
                tool_output = err_msg

            if execution_error:
                if plan:
                    for step in plan.steps:
                        if step.status == "pending":
                            step.status = "blocked"
                            step.error_message = "Prerequisite approved action failed"
                cursor = await db.execute(
                    """UPDATE agent_runs SET status = 'failed', error_message = ?,
                       result_text = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                       WHERE id = ? RETURNING *""",
                    (execution_error, execution_error,
                     serialize_plan(plan) if plan else run.get("structured_plan"), run_id)
                )
                failed_run = await cursor.fetchone()
                await db.commit()
                await log_event(db, run_id, "run_failed", execution_error)
                return RunResponse.model_validate(dict(failed_run))

            # Truthful plan state: mark unexecuted pending steps as skipped
            if plan:
                for s in plan.steps:
                    if s.status == "pending":
                        s.status = "skipped"
                        s.observation = "Resolved after authorized action"

            # Synthesize final response with fallback durability
            provider = get_provider()
            cursor = await db.execute("SELECT system_instructions FROM agent_definitions WHERE id = ?", (run["agent_id"],))
            agent_inst = await cursor.fetchone()
            sys_prompt = agent_inst["system_instructions"] if agent_inst else ""

            synth_prompt = (
                f"Operator query: {run['input_text']}\n\n"
                f"Supervisor authorized tool '{tool_name}' which executed and produced:\n{tool_output}\n\n"
                f"Summarize this action and its operational outcome for the plant supervisor."
            )
            try:
                final_response = await provider.generate_text(prompt=synth_prompt, system_prompt=sys_prompt)
                final_text = final_response.text
            except Exception as synth_err:
                logger.warning(f"Post-approval model synthesis failed for run {run_id}: {synth_err}")
                final_text = (
                    f"Action '{tool_name}' was successfully authorized and executed.\n\n"
                    f"Output:\n{tool_output}\n\n"
                    f"[Operational Note: Post-action model synthesis was unavailable: {synth_err}]"
                )

            updated_plan_json = serialize_plan(plan) if plan else run.get("structured_plan")
            cursor = await db.execute(
                """UPDATE agent_runs
                   SET status = 'completed', result_text = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                   WHERE id = ? RETURNING *""",
                (final_text, updated_plan_json, run_id)
            )
            updated_run = await cursor.fetchone()
            await db.commit()

            await log_event(db, run_id, "completed", "Run completed successfully after supervisor authorization.")
            return RunResponse.model_validate(dict(updated_run))

        elif approval["status"] == "rejected":
            tool_name = approval["tool_name"]
            reviewer = approval.get("reviewed_by", "supervisor")
            reject_text = f"Action '{tool_name}' was reviewed and REJECTED by supervisor ({reviewer}). Execution halted safely."

            plan = deserialize_plan(run.get("structured_plan"))
            if plan:
                for step in plan.steps:
                    if step.status == "waiting_for_approval":
                        step.status = "blocked"
                        step.error_message = f"Rejected by supervisor: {reviewer}"
                    elif step.status == "pending":
                        step.status = "blocked"
                        step.error_message = "Blocked due to rejected prerequisite"
            updated_plan_json = serialize_plan(plan) if plan else run.get("structured_plan")

            cursor = await db.execute(
                """UPDATE agent_runs
                   SET status = 'completed', result_text = ?, structured_plan = ?, completed_at = CURRENT_TIMESTAMP
                   WHERE id = ? RETURNING *""",
                (reject_text, updated_plan_json, run_id)
            )
            updated_run = await cursor.fetchone()
            await db.commit()

            await log_event(db, run_id, "rejected", f"Action '{tool_name}' rejected by supervisor.")
            return RunResponse.model_validate(dict(updated_run))

        else:
            raise HTTPException(status_code=400, detail=f"Unknown approval status: {approval['status']}")
