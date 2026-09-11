import pytest
from pathlib import Path
import openpyxl

from cognishift.core.document_processing.schemas import DocumentType, UnsupportedFileError
from cognishift.core.document_processing.inspector import detect_file_type, inspect_document
from cognishift.core.document_processing.office_renderer import (
    extract_xlsx_structured_content,
    verify_xlsx_cell,
    render_table_to_png_bytes,
)


def test_detect_xlsx_and_reject_legacy_xls(tmp_path: Path):
    xlsx_path = tmp_path / 'telemetry_log.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'SensorReadings'
    ws.append(['Timestamp', 'Sensor_Tag', 'Value', 'Unit'])
    ws.append(['2026-03-30 10:00:00', 'PT-101', 4.25, 'bar'])
    wb.save(str(xlsx_path))

    doc_type = detect_file_type(xlsx_path)
    assert doc_type == DocumentType.XLSX

    inspection = inspect_document(xlsx_path)
    assert inspection.is_valid is True
    assert inspection.document_type == DocumentType.XLSX
    assert inspection.mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert inspection.page_count >= 1

    # Legacy .xls rejection (OLE2 header)
    xls_legacy_path = tmp_path / 'legacy_telemetry.xls'
    with open(xls_legacy_path, 'wb') as f:
        f.write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 64)

    with pytest.raises(UnsupportedFileError) as exc_info:
        detect_file_type(xls_legacy_path)
    assert 'Legacy binary office format' in str(exc_info.value)


def test_extract_xlsx_and_visual_tiles(tmp_path: Path):
    xlsx_path = tmp_path / 'trip_events.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'EventHistory'
    ws.append(['Event_ID', 'Timestamp', 'Equipment', 'Description', 'Severity'])
    for i in range(1, 15):
        ws.append([f'EV-{i:03d}', f'10:{i:02d}:00', 'P-101A', f'Pressure drop {i}', 'HIGH' if i > 10 else 'INFO'])
    wb.save(str(xlsx_path))

    text_chunks, visual_tiles = extract_xlsx_structured_content(xlsx_path, chunk_row_size=10)
    assert len(text_chunks) >= 2
    assert 'EventHistory' in text_chunks[0]['text']
    assert 'P-101A' in text_chunks[0]['text']

    # Verify visual tile generation
    assert len(visual_tiles) >= 2
    png_bytes, tile_meta = visual_tiles[0]
    assert png_bytes.startswith(b'\x89PNG\r\n\x1a\n')
    assert tile_meta['sheet_name'] == 'EventHistory'
    assert tile_meta['row_start'] == 2


def test_verify_xlsx_cell(tmp_path: Path):
    xlsx_path = tmp_path / 'alarm_limits.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Thresholds'
    ws['A1'] = 'Parameter'
    ws['B1'] = 'TripSetpoint'
    ws['A2'] = 'PT-101_Max'
    ws['B2'] = 14.85
    wb.save(str(xlsx_path))

    res = verify_xlsx_cell(xlsx_path, 'Thresholds', 'B2')
    assert res['found'] is True
    assert res['value'] == 14.85

    res_missing = verify_xlsx_cell(xlsx_path, 'NonExistentSheet', 'B2')
    assert res_missing['found'] is False
